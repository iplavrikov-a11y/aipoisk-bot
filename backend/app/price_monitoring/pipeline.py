from __future__ import annotations

import asyncio
from dataclasses import asdict
import io
import json
import logging
import os
import re
import statistics
from datetime import datetime, timezone
from typing import Any, Callable, List, Optional
from urllib.parse import urlparse

import httpx
import openpyxl

from ..ai import call_llm
from ..exact_product.yandex_search import YandexSearchEngine, YandexSearchResult
from ..models import SystemSettings, parse_json_list
from ..supplier_locations import annotate_price_quotes_radius_status, sort_quotes_by_radius_priority
from .cleaner import (
    build_commercial_search_queries,
    build_fallback_broad_queries,
    clean_nomenclature_name,
    extract_generic_nomenclature_codes,
    is_drawing_or_part_code,
    normalize_unit,
)
from .excel_report import generate_price_monitoring_excel
from .extractor import (
    extract_supplier_quote,
    fetch_page_content_and_meta,
    find_price_grounding_match,
    is_foreign_source,
    is_valid_domain,
)
from .models import BatchMonitoringResult, ItemPriceResult, SupplierQuote

logger = logging.getLogger(__name__)

# Реестр федеральных ритейлеров с официальными юрлицами и ИНН для гарантированной точности реквизитов
KNOWN_MAJOR_RETAILERS: dict[str, dict[str, str]] = {
    "vseinstrumenti.ru": {
        "company_name": 'ООО "Всеинструменты.ру"',
        "inn": "7722753969",
        "contacts": "+7 (800) 550-37-70",
    },
    "etm.ru": {
        "company_name": 'ООО "ЭТМ"',
        "inn": "7804028214",
        "contacts": "+7 (800) 775-17-71",
    },
    "chipdip.ru": {
        "company_name": 'АО "Чип и Дип"',
        "inn": "7709022648",
        "contacts": "+7 (800) 555-62-70",
    },
    "rs24.ru": {
        "company_name": 'ООО "Русский Свет"',
        "inn": "6901032540",
        "contacts": "+7 (800) 777-70-04",
    },
    "lemanapro.ru": {
        "company_name": 'ООО "Лемана ПРО"',
        "inn": "5029069967",
        "contacts": "+7 (800) 700-00-99",
    },
    "dns-shop.ru": {
        "company_name": 'ООО "ДНС Ритейл"',
        "inn": "2540167061",
        "contacts": "+7 (800) 770-79-99",
    },
    "komus.ru": {
        "company_name": 'ООО "Комус"',
        "inn": "7721793895",
        "contacts": "+7 (800) 200-33-83",
    },
    "citilink.ru": {
        "company_name": 'ООО "Ситилинк"',
        "inn": "7718979307",
        "contacts": "+7 (800) 250-40-10",
    },
    "ozon.ru": {
        "company_name": 'ООО "Интернет Решения"',
        "inn": "7704217370",
        "contacts": "+7 (800) 234-77-11",
    },
    "wildberries.ru": {
        "company_name": 'ООО "РВБ"',
        "inn": "9714053621",
        "contacts": "+7 (800) 100-75-05",
    },
    "avito.ru": {
        "company_name": 'ООО "КЕХ еКоммерц"',
        "inn": "7710668349",
        "contacts": "+7 (800) 600-00-01",
    },
    "market.yandex.ru": {
        "company_name": 'ООО "Яндекс"',
        "inn": "7736207543",
        "contacts": "+7 (800) 234-27-12",
    },
    "megamarket.ru": {
        "company_name": 'ООО "Маркетплейс"',
        "inn": "9701048328",
        "contacts": "+7 (800) 600-08-88",
    },
}


def _domain_matches(candidate_domain: str, registered_domain: str) -> bool:
    """Match only the registered retailer domain itself or one of its subdomains."""
    candidate = str(candidate_domain or "").lower().strip(".")
    registered = str(registered_domain or "").lower().strip(".")
    return bool(candidate and registered and (candidate == registered or candidate.endswith("." + registered)))


def _is_measurement_like_identifier(token: str) -> bool:
    """Exclude quantities such as 1.2V and 1000mAh from exact product identifiers."""
    compact = re.sub(r"\s+", "", str(token or "").lower())
    return bool(
        re.fullmatch(
            r"\d+(?:[.,]\d+)?(?:г|кг|мг|т|л|мл|м|см|мм|км|шт|в|вт|квт|а|ма(?:/ч)?|ач|гц|v|w|kw|hz|mah|ah)",
            compact,
            re.I,
        )
    )


def _price_evidence_excerpt(text: str, price_match: Optional[re.Match[str]], limit: int = 500) -> str:
    """Return a bounded source window around the exact grounding match."""
    source = str(text or "")
    if limit <= 0 or price_match is None:
        return ""
    start = max(0, price_match.start() - (limit // 2))
    end = min(len(source), start + limit)
    if end - start < limit:
        start = max(0, end - limit)
    return " ".join(source[start:end].split())[:limit]


def parse_input_excel(file_bytes: bytes) -> List[tuple[str, str]]:
    """Читает загруженный Excel-файл и извлекает кортежи (номенклатура, единица_измерения)."""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    except Exception as exc:
        logger.warning("parse_input_excel: not a valid excel workbook (%s), falling back to text parsing", exc)
        try:
            from .cleaner import parse_items_from_text
            text = file_bytes.decode("utf-8", errors="ignore")
            return parse_items_from_text(text)
        except Exception:
            return []

    ws = wb.active
    items: List[tuple[str, str]] = []

    # Определяем индексы колонок
    nom_col = 1
    unit_col = 2
    model_col = None
    start_row = 1

    # Ищем заголовки в первых 15 строках
    for r in range(1, min(15, ws.max_row + 1)):
        for c in range(1, min(15, ws.max_column + 1)):
            val = str(ws.cell(row=r, column=c).value or "").strip().lower()
            if "номенклатур" in val or "наименован" in val or "товар" in val:
                nom_col = c
                start_row = r + 1
            elif any(k in val for k in ("модель", "марка", "артикул", "характеристик")):
                model_col = c
            elif "ед." in val or "ед.изм" in val or "единиц" in val:
                unit_col = c

    # Считываем строки данных
    for r in range(start_row, ws.max_row + 1):
        nom_val = str(ws.cell(row=r, column=nom_col).value or "").strip()
        unit_val = str(ws.cell(row=r, column=unit_col).value or "").strip()
        if not nom_val:
            continue
        # Пропускаем повторные заголовки
        if nom_val.lower().startswith("номенклатур") or nom_val.lower().startswith("наименован"):
            continue
        if model_col and model_col != nom_col:
            model_val = str(ws.cell(row=r, column=model_col).value or "").strip()
            if model_val and model_val.lower() not in nom_val.lower():
                nom_val = f"{nom_val} {model_val}".strip()
        items.append((nom_val, unit_val or "шт"))

    return items


def calculate_variation_coefficient(prices: list[float]) -> float:
    """
    Расчет коэффициента вариации по Приказу Минэкономразвития №567 (44-ФЗ):
    V = (sigma / mean) * 100%
    где sigma = sqrt(sum((p_i - mean)^2) / (n - 1))
    """
    if len(prices) < 2:
        return 0.0
    mean_p = statistics.mean(prices)
    if mean_p <= 0:
        return 0.0
    stdev_p = statistics.stdev(prices)
    return (stdev_p / mean_p) * 100.0


def cluster_and_reconcile_item_quotes(
    item_res: ItemPriceResult,
    cluster_ratio: float = 1.4,
    max_suppliers: int = 5,
    max_variation_pct: float = 33.0,
) -> None:
    """
    1. Авто-приведение кратностей (баллоны 50л -> кг, м3 -> штуки, мешки 50кг -> т, упаковки 100 шт).
    2. Кластеризация цен по Приказу №567 (44-ФЗ) с коэффициентом вариации V <= max_variation_pct (по умолчанию 33%):
       - Ищет комбинаторное подмножество поставщиков максимального размера k (до max_suppliers, затем k-1, ..., 2),
         удовлетворяющее критерию однородности V <= max_variation_pct.
       - Если найдено подмножество с V <= max_variation_pct, котировки внутри него считаются подтвержденными.
       - Все котировки вне однородного подмножества маркируются is_outlier = True.
       - Если ни одной однородной группы из >= 2 поставщиков нет и разброс превышает max_variation_pct,
         все котировки маркируются как выбросы (is_outlier = True), чтобы честно отразить отсутствие товара.
    """
    valid_quotes = [q for q in item_res.quotes if q.price_rub > 0]
    if not valid_quotes:
        return

    # 1. Приведение кратностей фасовок
    if len(valid_quotes) >= 2:
        prices = [q.price_rub for q in valid_quotes]
        min_p = min(prices)
        max_p = max(prices)
        ratio = max_p / min_p if min_p > 0 else 1.0

        if ratio >= 2.5:
            item_lower = item_res.clean_name.lower()
            norm_u = normalize_unit(item_res.unit)

            for q in valid_quotes:
                q_ratio = q.price_rub / min_p

                # 1. Пропан (баллон 50л vs кг, коэффициент 21)
                if "пропан" in item_lower and norm_u in ("кг", "к-г"):
                    if 14.0 <= q_ratio <= 35.0:
                        converted = round(q.price_rub / 21.0, 2)
                        q.price_rub = converted
                        q.conversion_comment = f"Авто-пересчет кратности: {q.raw_price_rub:,.0f} ₽ за баллон 50л -> {converted:,.2f} ₽/кг"

                # 2. Кислород (баллон 40л vs м3, коэффициент 6.3)
                elif "кислород" in item_lower and norm_u in ("м3", "куб"):
                    if 4.5 <= q_ratio <= 14.0:
                        converted = round(q.price_rub / 6.3, 2)
                        q.price_rub = converted
                        q.conversion_comment = f"Авто-пересчет кратности: {q.raw_price_rub:,.0f} ₽ за баллон 40л -> {converted:,.2f} ₽/м³"

                # 3. Вес и сыпучие: Тонна <-> Килограмм
                if norm_u in ("кг", "к-г") and q_ratio >= 400.0 and q.price_rub > 20000:
                    converted = round(q.price_rub / 1000.0, 2)
                    q.price_rub = converted
                    q.conversion_comment = f"Авто-пересчет кратности: {q.raw_price_rub:,.0f} ₽ за тонну -> {converted:,.2f} ₽/кг"
                elif norm_u in ("т", "тонна") and (max_p / q.price_rub) >= 300.0 and q.price_rub < 3000:
                    converted = round(q.price_rub * 1000.0, 2)
                    q.price_rub = converted
                    q.conversion_comment = f"Авто-пересчет кратности: {q.raw_price_rub:,.2f} ₽/кг -> {converted:,.0f} ₽/т"

                # 4. Цемент в мешках: 50 кг vs тонна (20 мешков в тонне)
                elif norm_u in ("т", "тонна") and "цемент" in item_lower and (max_p / q.price_rub) >= 15.0 and q.price_rub < 1500:
                    converted = round(q.price_rub * 20.0, 2)
                    q.price_rub = converted
                    q.conversion_comment = f"Авто-пересчет кратности: {q.raw_price_rub:,.0f} ₽ за мешок 50 кг -> {converted:,.0f} ₽/т"

                # 5. ГСМ и жидкости: Бочка 200л vs Литр
                elif norm_u in ("л", "литр") and (q_ratio >= 100.0 or (max_p / q.price_rub) >= 100.0) and q.price_rub > 15000:
                    converted = round(q.price_rub / 200.0, 2)
                    q.price_rub = converted
                    q.conversion_comment = f"Авто-пересчет кратности: {q.raw_price_rub:,.0f} ₽ за бочку 200л -> {converted:,.2f} ₽/л"

                # 6. Пиломатериалы: штука vs м³
                elif norm_u in ("м3", "куб") and any(w in item_lower for w in ("доск", "брус", "пиломатериал")) and (max_p / q.price_rub) >= 10.0 and q.price_rub < 3000:
                    q.is_outlier = True
                    q.comment = f"⚠️ Цена {q.price_rub:,.0f} ₽ указана за штуку, а не за м³."

                # 7. Упаковки по 100 шт vs штука
                elif norm_u == "шт" and 70.0 <= q_ratio <= 130.0 and (
                    bool(re.search(r'\b100\s*(?:шт|штук|ед|ш)\b', item_lower))
                    or bool(re.search(r'\b(?:уп|упак|пач|пачка|короб|кор)\S*\s*(?:по\s*)?100\b', item_lower))
                ):
                    converted = round(q.price_rub / 100.0, 2)
                    q.price_rub = converted
                    q.conversion_comment = f"Авто-пересчет кратности: {q.raw_price_rub:,.0f} ₽ за уп (100 шт) -> {converted:,.2f} ₽/шт"

    # 2. Кластеризация цен по Приказу №567 (V <= max_variation_pct)
    for q in item_res.quotes:
        if not (q.is_outlier and "штуку, а не за м³" in (q.comment or "")):
            q.is_outlier = False
            q.comment = ""

    valid_quotes = [q for q in item_res.quotes if q.price_rub > 0 and not q.is_outlier]
    if len(valid_quotes) <= 1:
        return

    import itertools

    valid_quotes.sort(key=lambda q: q.price_rub)

    # Ищем однородное подмножество поставщиков максимального размера k (до max_suppliers)
    best_homogeneous_subset: list[SupplierQuote] = []
    best_v = 999.0

    target_k_max = min(len(valid_quotes), max_suppliers)
    for k in range(target_k_max, 1, -1):
        candidates_at_k: list[tuple[float, list[SupplierQuote]]] = []
        for combo in itertools.combinations(valid_quotes, k):
            p_list = [q.price_rub for q in combo]
            v = calculate_variation_coefficient(p_list)
            if v <= max_variation_pct:
                candidates_at_k.append((v, list(combo)))
        if candidates_at_k:
            # Выбираем группу с наименьшим коэффициентом вариации
            candidates_at_k.sort(key=lambda c: c[0])
            best_v, best_homogeneous_subset = candidates_at_k[0]
            break

    if best_homogeneous_subset:
        best_ids = {id(q) for q in best_homogeneous_subset}
        min_cp = min(q.price_rub for q in best_homogeneous_subset)
        max_cp = max(q.price_rub for q in best_homogeneous_subset)
        for q in item_res.quotes:
            if id(q) not in best_ids:
                q.is_outlier = True
                q.comment = f"⚠️ Исключена: цена {q.price_rub:,.2f} ₽ не вошла в однородную группу V={best_v:.1f}% ({min_cp:,.2f} - {max_cp:,.2f} ₽)"
    else:
        # Цены не образуют однородной совокупности (разброс > max_variation_pct)
        all_v = calculate_variation_coefficient([q.price_rub for q in valid_quotes])
        for q in item_res.quotes:
            q.is_outlier = True
            q.comment = f"⚠️ Неоднородная цена: вариация {all_v:.1f}% превышает лимит {max_variation_pct:.0f}% (Приказ №567)"

    # Сортируем: сначала подтвержденный кластер по возрастанию цены, затем исключенные
    item_res.quotes.sort(key=lambda q: (q.is_outlier, q.price_rub if q.price_rub > 0 else 999_999_999))


# Псевдоним для обратной совместимости
validate_and_reconcile_item_quotes = cluster_and_reconcile_item_quotes


def _attach_source_position(result: ItemPriceResult, source: dict[str, Any]) -> None:
    result.position_id = str(source.get("position_id") or source.get("id") or "")
    result.source_position_no = source.get("source_position_no") or source.get("position_no")
    result.source_file = str(source.get("source_file") or "")
    result.source_sheet = str(source.get("source_sheet") or "")
    result.source_row = source.get("source_row")
    result.source_excerpt = str(source.get("source_excerpt") or "")
    result.quantity = str(source.get("quantity") or "")
    result.requirements = [dict(requirement) for requirement in source.get("requirements") or [] if isinstance(requirement, dict)]


def _save_checkpoint(checkpoint_path: str, results: list[ItemPriceResult], total: int, yandex_requests: int = 0) -> None:
    """Атомарно сохраняет промежуточный прогресс для защиты от сбоев и перезапусков."""
    try:
        tmp_path = checkpoint_path + ".tmp"
        payload = {
            "total": total,
            "completed": len(results),
            "yandex_requests": yandex_requests,
            "items": [asdict(result) for result in results],
        }
        with open(tmp_path, "w", encoding="utf-8") as f:
            json.dump(payload, f, ensure_ascii=False, indent=2)
        os.replace(tmp_path, checkpoint_path)
    except Exception as exc:
        logger.warning("Checkpoint save error: %s", exc)


def _load_checkpoint(checkpoint_path: str) -> tuple[list[ItemPriceResult], int]:
    """Загружает сохраненные результаты, если задача была прервана."""
    results: list[ItemPriceResult] = []
    yandex_requests: int = 0
    if not os.path.exists(checkpoint_path):
        return results, yandex_requests
    try:
        with open(checkpoint_path, "r", encoding="utf-8") as f:
            data = json.load(f)
            yandex_requests = int(data.get("yandex_requests", 0) or 0)
            for it in data.get("items", []):
                quotes = [SupplierQuote(**q) for q in it.get("quotes", [])]
                item_obj = ItemPriceResult(
                    row_number=it["row_number"],
                    raw_name=it["raw_name"],
                    clean_name=it["clean_name"],
                    unit=it["unit"],
                    quotes=quotes,
                    status=it.get("status", "found"),
                    not_found_reason=str(it.get("not_found_reason", "") or ""),
                    search_diagnostics=dict(it.get("search_diagnostics", {}) or {}),
                )
                _attach_source_position(item_obj, it)
                results.append(item_obj)
        logger.info("Loaded %d completed items from checkpoint %s (yandex reqs=%d)", len(results), checkpoint_path, yandex_requests)
    except Exception as exc:
        logger.warning("Checkpoint load error %s: %s", checkpoint_path, exc)
    return results, yandex_requests


async def generate_monitoring_queries_ai(
    clean_name: str,
    unit: str,
    settings: SystemSettings,
    wave: int = 1,
    city: str = "",
) -> list[str]:
    """
    Генерирует поисковые запросы через ИИ (Правило 14: семантический анализ без жестких шаблонов).
    ИИ самостоятельно анализирует класс изделия, марку, артикул, технические параметры и коммерческий контекст.
    """
    if not settings or not getattr(settings, "has_active_ai_provider", False):
        return []

    wave_descriptions = {
        1: "первичный поиск прямых коммерческих предложений с ценами",
        2: "адаптивный добор альтернативных поставщиков, складов и дилеров",
        3: "широкий поиск по каталогам производителей и оптовым прайс-листам",
    }
    wave_desc = wave_descriptions.get(wave, "поиск коммерческих цен")
    clean_city = str(city or "").strip()
    city_context = f'\nРегион/город объекта поставки: "{clean_city}" (приоритет поставщикам, дилерам и складам с наличием или доставкой в этом регионе)' if clean_city else ""

    nom_codes = extract_generic_nomenclature_codes(clean_name)
    exact_codes = [
        c for c in nom_codes
        if is_drawing_or_part_code(c, clean_name) and not _is_measurement_like_identifier(c)
    ]
    drawing_instruction = ""
    if exact_codes:
        primary_drawing = exact_codes[0]
        drawing_instruction = f"""
ВНИМАНИЕ (ТОЧНЫЕ ИДЕНТИФИКАТОРЫ):
Позиция содержит точный артикул, модель или номер чертежа: {exact_codes}.
Каждый поисковый запрос ОБЯЗАН содержать точный идентификатор в строгих кавычках: '"{primary_drawing}"'.
КАТЕГОРИЧЕСКИ ЗАПРЕЩЕНО исключать идентификатор из запросов во избежание омонимов и поиска по общим словам!"""

    prompt = f"""Ты — ведущий эксперт по корпоративным закупкам и поисковой оптимизации (ПРАВИЛО 14: СЕМАНТИЧЕСКИЙ АНАЛИЗ И ФОРМИРОВАНИЕ ПОИСКОВЫХ ЗАПРОСОВ БЕЗ ЖЕСТКИХ СЛОВАРЕЙ).
Твоя задача — составить 3-5 естественных, высокоточных поисковых запросов для поисковой системы (Яндекс), чтобы найти актуальные коммерческие цены поставщиков и прайс-листы.

Наименование позиции: "{clean_name}"
Требуемая единица измерения: "{unit}"
Цель запросов: {wave_desc} (волна #{wave}){city_context}{drawing_instruction}

ТРЕБОВАНИЯ (ПРАВИЛО 14):
1. Семантически определи физическую суть товара, класс оборудования/материала, точную марку, артикул, ГОСТ или чертеж.
2. Составь запросы естественной длины (3-6 слов) для нахождения страниц с реальными коммерческими ценами:
   - Прямой коммерческий запрос по целевому названию и марке со словами 'купить цена'
   - Запрос по точной марке/артикулу у поставщиков и дилеров со склада
   - Запрос по производителю, каталогу или оптовому прайс-листу
3. КАТЕГОРИЧЕСКИ ЗАПРЕЩЕНО включать стоп-слова ('не менее', 'в соответствии с', 'требуется'), мусор 1С и лишние служебные знаки.
4. Запросы должны приводить на коммерческие страницы реальных компаний, складов и производителей в РФ.

Ответь СТРОГО JSON-массивом строк:
["запрос 1", "запрос 2", "запрос 3"]
"""
    try:
        raw_response = await call_llm(settings, prompt, tier="light", json_mode=True)
        parsed = parse_json_list(raw_response)
        if isinstance(parsed, list):
            res: list[str] = []
            seen: set[str] = set()
            for q in parsed:
                if isinstance(q, str):
                    q_str = " ".join(q.strip(" \"'\t\r\n").split())
                    if q_str and len(q_str) >= 5 and q_str.lower() not in seen:
                        # Если позиция содержит точный артикул, модель или чертеж, запрос обязан сохранить его.
                        if exact_codes:
                            has_draw = all(re.search(re.escape(dc), q_str, re.I) for dc in exact_codes)
                            if not has_draw:
                                continue
                        seen.add(q_str.lower())
                        res.append(q_str)
            if res:
                return res[:5]
    except Exception as exc:
        logger.debug("generate_monitoring_queries_ai error for '%s': %s", clean_name, exc)

    return []


async def process_single_item(
    item_tuple: tuple[int, str, str],
    engine: YandexSearchEngine,
    settings: SystemSettings,
    http_client: httpx.AsyncClient,
    search_semaphore: asyncio.Semaphore,
    max_suppliers: int = 5,
    min_target_suppliers: int = 3,
    cluster_ratio: float = 1.4,
    filter_marketplaces: bool = True,
    filter_foreign: bool = True,
    max_variation_pct: float = 33.0,
    city: str = "",
    geo_radius_km: Optional[int] = None,
    source_position: Optional[dict[str, Any]] = None,
) -> ItemPriceResult:
    row_idx, raw_name, unit = item_tuple
    clean = clean_nomenclature_name(raw_name)
    item_res = ItemPriceResult(
        row_number=row_idx,
        raw_name=raw_name,
        clean_name=clean,
        unit=unit,
    )
    if source_position is not None:
        _attach_source_position(item_res, source_position)
        # Only this selected row's explicit requirements influence query and offer matching.
        requirements_text = "; ".join(
            f"{requirement.get('param_name')}: {requirement.get('tz_requirement')}"
            for requirement in item_res.requirements
            if requirement.get("param_name") and requirement.get("tz_requirement")
        )
        if requirements_text:
            clean = f"{clean}; {requirements_text}"

    nom_codes = extract_generic_nomenclature_codes(clean)
    exact_identifier_codes = [
        c for c in nom_codes
        if is_drawing_or_part_code(c, clean) and not _is_measurement_like_identifier(c)
    ]

    diagnostics: dict[str, Any] = {
        "search_queries": 0,
        "search_errors": 0,
        "search_results": 0,
        "unique_candidate_urls": 0,
        "distinct_domains": 0,
        "fetch_attempts": 0,
        "fetch_unavailable": 0,
        "snippet_fallbacks": 0,
        "foreign_rejections": 0,
        "extraction_attempts": 0,
        "extraction_rejected": 0,
        "quotes_accepted": 0,
        "duplicate_suppliers": 0,
        "domain_attempt_limit": 0,
        "budget_exhausted": False,
        "waves_run": [],
    }
    item_res.search_diagnostics = diagnostics

    def _preserve_exact_codes(search_queries: list[str]) -> list[str]:
        """Keep exact product, model, and factory identifiers quoted in every wave."""
        if not exact_identifier_codes:
            return search_queries
        strict_prefix = " ".join(f'"{code}"' for code in exact_identifier_codes)
        protected: list[str] = []
        for query in search_queries:
            query_without_codes = str(query or "")
            for code in exact_identifier_codes:
                query_without_codes = re.sub(re.escape(code), " ", query_without_codes, flags=re.I)
            query_without_codes = " ".join(query_without_codes.replace('"', " ").split())
            protected.append(f"{strict_prefix} {query_without_codes}".strip())
        return protected

    candidate_urls: list[str] = []
    url_to_meta: dict[str, YandexSearchResult] = {}

    # ВОЛНА 1: Первичные коммерческие запросы (ИИ по Правилу 14 с алгоритмическим фоллбэком)
    queries: list[str] = []
    if settings.has_active_ai_provider:
        queries = await generate_monitoring_queries_ai(clean, unit, settings, wave=1, city=city)
    if not queries:
        queries = build_commercial_search_queries(clean, unit, city=city)
    queries = _preserve_exact_codes(queries)
    diagnostics["waves_run"].append(1)

    async with search_semaphore:
        for q in queries[:4]:
            diagnostics["search_queries"] += 1
            try:
                results = await engine.search(
                    q, max_results=20, enrich_query=False, filter_marketplaces=filter_marketplaces
                )
                diagnostics["search_results"] += len(results)
                for r in results:
                    if r.url and is_valid_domain(r.url, target_item_name=clean, filter_marketplaces=filter_marketplaces, filter_foreign=filter_foreign):
                        url_to_meta[r.url] = r
                        if r.url not in candidate_urls:
                            candidate_urls.append(r.url)
            except Exception as e:
                diagnostics["search_errors"] += 1
                logger.warning("Search query error for %s: %s", q, e)
            if len(candidate_urls) >= 30:
                break

    seen_domains: set[str] = set()
    seen_inns: set[str] = set()
    attempted_urls: set[str] = set()
    base_domain_attempts: dict[str, int] = {}
    target_suppliers = min_target_suppliers or max_suppliers
    MAX_LLM_EXTRACTIONS_PER_ITEM = 20
    WAVE_1_EXTRACTION_LIMIT = 12
    WAVE_2_EXTRACTION_LIMIT = 18
    MAX_ATTEMPTS_PER_BASE_DOMAIN = 2
    llm_extractions_done = 0

    def _get_base_domain(d: str) -> str:
        parts = d.split(".")
        if len(parts) >= 3 and ".".join(parts[-2:]) in ("spb.ru", "msk.ru", "nov.ru", "com.ru", "net.ru", "org.ru"):
            return ".".join(parts[-3:])
        if len(parts) >= 2:
            return ".".join(parts[-2:])
        return d

    def _get_consensus_count() -> int:
        valid_qs = [q for q in item_res.quotes if q.is_price_found and q.price_rub > 0]
        if not valid_qs:
            return 0
        cluster_and_reconcile_item_quotes(
            item_res,
            cluster_ratio=cluster_ratio,
            max_suppliers=max_suppliers,
            max_variation_pct=max_variation_pct,
        )
        return len([q for q in item_res.quotes if q.price_rub > 0 and not q.is_outlier])

    async def _evaluate_candidates(urls: list[str], extraction_limit: int) -> None:
        nonlocal llm_extractions_done
        for url in urls:
            if llm_extractions_done >= min(extraction_limit, MAX_LLM_EXTRACTIONS_PER_ITEM):
                logger.debug("Wave extraction budget reached (%d) for '%s'", extraction_limit, clean)
                break
            if _get_consensus_count() >= target_suppliers:
                logger.debug("Target consensus reached (%d suppliers) for '%s'", target_suppliers, clean)
                break
            if len([q for q in item_res.quotes if q.price_rub > 0]) >= target_suppliers * 2 + 4:
                break
            try:
                domain = urlparse(url).netloc.lower().replace("www.", "")
                base_domain = _get_base_domain(domain)
                if domain in seen_domains or base_domain in seen_domains:
                    continue
                if url in attempted_urls:
                    continue
                if base_domain_attempts.get(base_domain, 0) >= MAX_ATTEMPTS_PER_BASE_DOMAIN:
                    diagnostics["domain_attempt_limit"] += 1
                    continue

                attempted_urls.add(url)
                base_domain_attempts[base_domain] = base_domain_attempts.get(base_domain, 0) + 1
                diagnostics["fetch_attempts"] += 1

                page_text, inn, kpp, contacts = await fetch_page_content_and_meta(http_client, url)
                raw_page_text = str(page_text or "")
                meta = url_to_meta.get(url)

                # Объединяем сниппет поисковой выдачи со страницей (критично при WAF, динамическом JS или капче)
                snippet_block = ""
                if meta and (meta.snippet or meta.title):
                    snippet_block = f"[ПОИСКОВЫЙ СНИППЕТ ЯНДЕКСА]: {meta.title} | {meta.snippet}\n\n"

                if not page_text or len(page_text) < 60:
                    diagnostics["fetch_unavailable"] += 1
                    if snippet_block:
                        diagnostics["snippet_fallbacks"] += 1
                        page_text = snippet_block + f"Источник: {url}"
                    else:
                        continue
                else:
                    page_text = snippet_block + page_text

                if not page_text or len(page_text) < 30:
                    continue

                # Дополнительная проверка на иностранный язык контента
                if filter_foreign and is_foreign_source(url, page_text):
                    diagnostics["foreign_rejections"] += 1
                    logger.debug("Skip foreign source by content for %s", url)
                    continue

                llm_extractions_done += 1
                diagnostics["extraction_attempts"] = llm_extractions_done
                # Реестр федеральных сетей
                for kd, kdata in KNOWN_MAJOR_RETAILERS.items():
                    if _domain_matches(domain, kd):
                        if not inn:
                            inn = kdata["inn"]
                        if not contacts:
                            contacts = kdata["contacts"]
                        break

                quote = await extract_supplier_quote(
                    settings=settings,
                    target_item_name=clean,
                    target_unit=unit,
                    url=url,
                    page_text=page_text,
                    pre_inn=inn,
                    pre_kpp=kpp,
                    pre_contacts=contacts,
                )
                if quote and quote.is_price_found and quote.price_rub > 0:
                    grounding_price = quote.raw_price_rub or quote.price_rub
                    page_price_match = find_price_grounding_match(grounding_price, raw_page_text)
                    snippet_price_match = find_price_grounding_match(grounding_price, snippet_block)
                    if page_price_match is not None:
                        quote.evidence_level = "page_content"
                        evidence_text = raw_page_text
                        evidence_match = page_price_match
                    elif snippet_price_match is not None:
                        quote.evidence_level = "search_snippet"
                        evidence_text = snippet_block
                        evidence_match = snippet_price_match
                    else:
                        quote.evidence_level = "unknown"
                        evidence_text = ""
                        evidence_match = None
                    quote.retrieved_at = datetime.now(timezone.utc).isoformat()
                    quote.source_excerpt = _price_evidence_excerpt(evidence_text, evidence_match)
                    for kd, kdata in KNOWN_MAJOR_RETAILERS.items():
                        if _domain_matches(domain, kd):
                            quote.company_name = kdata["company_name"]
                            quote.inn = kdata["inn"]
                            if not quote.contacts:
                                quote.contacts = kdata["contacts"]
                            break
                    if quote.inn and quote.inn in seen_inns:
                        diagnostics["duplicate_suppliers"] += 1
                        logger.debug("Skip duplicate supplier INN %s for %s", quote.inn, url)
                        continue
                    if quote.inn:
                        seen_inns.add(quote.inn)
                    seen_domains.add(domain)
                    seen_domains.add(base_domain)
                    item_res.quotes.append(quote)
                    diagnostics["quotes_accepted"] += 1
                else:
                    diagnostics["extraction_rejected"] += 1
            except Exception as exc:
                logger.debug("Error processing url %s: %s", url, exc)

    await _evaluate_candidates(candidate_urls, WAVE_1_EXTRACTION_LIMIT)

    # ВОЛНА 2 (АДАПТИВНЫЙ ПОИСК ДЛЯ ДОСТИЖЕНИЯ ЦЕЛЕВОГО КОЛИЧЕСТВА ПОСТАВЩИКОВ):
    consensus_count = _get_consensus_count()
    wave2_urls: list[str] = []
    if consensus_count < target_suppliers and llm_extractions_done < MAX_LLM_EXTRACTIONS_PER_ITEM:
        fallback_queries: list[str] = []
        if settings.has_active_ai_provider:
            fallback_queries = await generate_monitoring_queries_ai(clean, unit, settings, wave=2, city=city)
        if not fallback_queries:
            fallback_queries = build_fallback_broad_queries(clean, unit, city=city)
        fallback_queries = _preserve_exact_codes(fallback_queries)
        diagnostics["waves_run"].append(2)
        async with search_semaphore:
            for bq in fallback_queries[:4]:
                diagnostics["search_queries"] += 1
                try:
                    b_results = await engine.search(
                        bq, max_results=20, enrich_query=False, filter_marketplaces=filter_marketplaces
                    )
                    diagnostics["search_results"] += len(b_results)
                    for r in b_results:
                        if r.url and is_valid_domain(r.url, target_item_name=clean, filter_marketplaces=filter_marketplaces, filter_foreign=filter_foreign):
                            url_to_meta[r.url] = r
                            if r.url not in candidate_urls and r.url not in wave2_urls:
                                wave2_urls.append(r.url)
                except Exception as b_exc:
                    diagnostics["search_errors"] += 1
                    logger.warning("Broad search query error for %s: %s", bq, b_exc)
                if len(candidate_urls) + len(wave2_urls) >= 45:
                    break

        if wave2_urls:
            await _evaluate_candidates(wave2_urls, WAVE_2_EXTRACTION_LIMIT)
            consensus_count = _get_consensus_count()

    # ВОЛНА 3 (СВЕРХШИРОКИЙ ДОБОР ДЛЯ ДОСТИЖЕНИЯ ЦЕЛЕВОГО КОЛИЧЕСТВА ПОСТАВЩИКОВ):
    should_run_wave3 = (
        consensus_count < target_suppliers
        and llm_extractions_done < MAX_LLM_EXTRACTIONS_PER_ITEM
    )
    wave3_urls: list[str] = []
    if should_run_wave3:
        wave3_queries: list[str] = []
        if settings.has_active_ai_provider:
            wave3_queries = await generate_monitoring_queries_ai(clean, unit, settings, wave=3, city=city)
        if not wave3_queries:
            tokens = clean.split()
            wave3_queries = [
                f"{clean} цена поставщики",
                f"{clean} прайс каталог склад",
            ]
            if len(tokens) >= 2:
                short_clean = " ".join(dict.fromkeys(tokens[:3]))
                wave3_queries.append(f"{short_clean} купить цена")
        wave3_queries = _preserve_exact_codes(wave3_queries)
        diagnostics["waves_run"].append(3)

        async with search_semaphore:
            for wq in wave3_queries[:4]:
                diagnostics["search_queries"] += 1
                try:
                    w_results = await engine.search(
                        wq, max_results=20, enrich_query=False, filter_marketplaces=filter_marketplaces
                    )
                    diagnostics["search_results"] += len(w_results)
                    for r in w_results:
                        if r.url and is_valid_domain(r.url, target_item_name=clean, filter_marketplaces=filter_marketplaces, filter_foreign=filter_foreign):
                            url_to_meta[r.url] = r
                            if r.url not in candidate_urls and r.url not in wave2_urls and r.url not in wave3_urls:
                                wave3_urls.append(r.url)
                except Exception as w_exc:
                    diagnostics["search_errors"] += 1
                    logger.warning("Wave 3 search query error for %s: %s", wq, w_exc)
                if len(candidate_urls) + len(wave2_urls) + len(wave3_urls) >= 55:
                    break

        if wave3_urls:
            await _evaluate_candidates(wave3_urls, MAX_LLM_EXTRACTIONS_PER_ITEM)

    # If broad waves produced no usable candidates, use the remaining fixed budget on
    # previously unattempted primary URLs instead of silently discarding their tail.
    if consensus_count < target_suppliers and llm_extractions_done < MAX_LLM_EXTRACTIONS_PER_ITEM:
        await _evaluate_candidates(candidate_urls, MAX_LLM_EXTRACTIONS_PER_ITEM)
    if consensus_count < target_suppliers and llm_extractions_done < MAX_LLM_EXTRACTIONS_PER_ITEM:
        await _evaluate_candidates(wave2_urls, MAX_LLM_EXTRACTIONS_PER_ITEM)
    if consensus_count < target_suppliers and llm_extractions_done < MAX_LLM_EXTRACTIONS_PER_ITEM:
        await _evaluate_candidates(wave3_urls, MAX_LLM_EXTRACTIONS_PER_ITEM)

    # Итоговая кластеризация цен и жесткая отсечка выбросов по 44-ФЗ (V <= max_variation_pct)
    cluster_and_reconcile_item_quotes(
        item_res,
        cluster_ratio=cluster_ratio,
        max_suppliers=max_suppliers,
        max_variation_pct=max_variation_pct,
    )
    if city and geo_radius_km:
        geo_stats = annotate_price_quotes_radius_status(item_res.quotes, city, geo_radius_km)
        sort_quotes_by_radius_priority(item_res.quotes)
        diagnostics["geo_radius"] = {
            "delivery_city": city,
            "radius_km": geo_radius_km,
            **geo_stats,
            "quotes": [
                {
                    "company_name": quote.company_name,
                    "supplier_city": quote.supplier_city,
                    "distance_km": quote.supplier_distance_km,
                    "status": quote.supplier_distance_radius_status,
                    "location_source": quote.supplier_location_source,
                    "location_confidence": quote.supplier_location_confidence,
                }
                for quote in item_res.quotes
            ],
        }

    confirmed_count = len([q for q in item_res.quotes if q.price_rub > 0 and not q.is_outlier])
    budget_exhausted = (
        llm_extractions_done >= MAX_LLM_EXTRACTIONS_PER_ITEM
        and confirmed_count < target_suppliers
    )
    diagnostics["budget_exhausted"] = budget_exhausted
    if confirmed_count >= 1:
        item_res.status = "found"
    elif any(q.price_rub > 0 for q in item_res.quotes):
        item_res.status = "divergent"
    else:
        item_res.status = "not_found"
        if budget_exhausted:
            item_res.not_found_reason = "budget_exhausted"
        elif diagnostics["search_results"] == 0 or not url_to_meta:
            item_res.not_found_reason = "no_results"
        elif diagnostics["extraction_attempts"] == 0 and diagnostics["fetch_unavailable"] > 0:
            item_res.not_found_reason = "network"
        else:
            item_res.not_found_reason = "all_rejected"
    all_candidate_urls = candidate_urls + wave2_urls + wave3_urls
    diagnostics["unique_candidate_urls"] = len(dict.fromkeys(all_candidate_urls))
    diagnostics["distinct_domains"] = len(
        {
            _get_base_domain(urlparse(url).netloc.lower().replace("www.", ""))
            for url in all_candidate_urls
            if urlparse(url).netloc
        }
    )
    return item_res


async def run_price_monitoring_pipeline(
    file_bytes: bytes,
    settings: SystemSettings,
    output_excel_path: str,
    price_per_position_rub: float = 50.0,
    progress_callback: Optional[Callable[[int, str], Any]] = None,
    max_suppliers: int = 5,
    max_parallel_items: int = 4,
    checkpoint_path: Optional[str] = None,
    cluster_ratio: float = 1.4,
    filter_marketplaces: bool = True,
    filter_foreign: bool = True,
    max_variation_pct: float = 33.0,
    items: Optional[List[tuple[str, str]]] = None,
    delivery_city: str = "",
    geo_radius_km: Optional[int] = None,
    source_positions: Optional[List[dict[str, Any]]] = None,
) -> BatchMonitoringResult:
    """Полный пайплайн мониторинга цен с защитой от сбоев (чекпоинты) и кластеризацией цен."""
    if items is not None and len(items) > 0:
        parsed_items = items
    else:
        parsed_items = parse_input_excel(file_bytes)
    total = len(parsed_items)
    if source_positions is not None:
        if len(source_positions) != total:
            raise ValueError("Количество исходных позиций не совпадает с выбранным списком")
        for source, (name, unit) in zip(source_positions, parsed_items):
            if not isinstance(source, dict) or str(source.get("name") or source.get("name_in_tz") or "").strip() != str(name).strip():
                raise ValueError("Исходные позиции не соответствуют выбранным наименованиям")
            if source.get("unit") and normalize_unit(str(source["unit"])) != normalize_unit(unit):
                raise ValueError("Единицы исходных позиций не совпадают с выбранным списком")
    if total == 0:
        raise ValueError("В загруженном файле не найдено строк с номенклатурой.")

    # Проверка лимита строк из настроек
    max_positions_limit = int(getattr(settings, "price_monitoring_max_positions", 1000) or 1000)
    if total > max_positions_limit:
        raise ValueError(f"В файле {total} позиций. Лимит в настройках системы — {max_positions_limit} строк. Разбейте файл на части.")

    if progress_callback:
        progress_callback(5, f"Загружено {total} позиций из файла")

    engine = YandexSearchEngine(
        folder_id=settings.yandex_search_folder_id,
        api_key=settings.yandex_search_api_key,
    )
    search_sem = asyncio.Semaphore(max_parallel_items)

    results: List[ItemPriceResult] = []
    completed_rows: set[int] = set()

    # Восстановление из чекпоинта при наличии
    if checkpoint_path and os.path.exists(checkpoint_path):
        results, prev_yandex_reqs = _load_checkpoint(checkpoint_path)
        engine.total_requests += prev_yandex_reqs
        # Failed rows remain explicit outcomes; an interrupted run retries them fresh.
        results = [r for r in results if r.status != "error"]
        if source_positions is not None:
            results = [r for r in results if 1 <= r.row_number <= total
                       and r.position_id == str(source_positions[r.row_number - 1].get("position_id") or source_positions[r.row_number - 1].get("id") or "")
                       and r.raw_name == parsed_items[r.row_number - 1][0]
                       and r.requirements == (source_positions[r.row_number - 1].get("requirements") or [])]
        completed_rows = {r.row_number for r in results}
        if completed_rows and progress_callback:
            progress_callback(
                int((len(completed_rows) / total) * 80) + 10,
                f"Возобновлено из контрольной точки: {len(completed_rows)}/{total} позиций уже готовы",
            )

    proxy_url = os.getenv("PRICE_MONITORING_PROXY") or os.getenv("HTTPS_PROXY")
    async with httpx.AsyncClient(timeout=10.0, follow_redirects=True, proxy=proxy_url) as http_client:
        batch_size = max_parallel_items
        for i in range(0, total, batch_size):
            chunk = [(idx + 1, nom, unit) for idx, (nom, unit) in enumerate(parsed_items[i : i + batch_size], start=i)]
            chunk_to_process = [it for it in chunk if it[0] not in completed_rows]

            if chunk_to_process:
                tasks = [
                    process_single_item(
                        item,
                        engine,
                        settings,
                        http_client,
                        search_sem,
                        max_suppliers=max_suppliers,
                        min_target_suppliers=max_suppliers,
                        cluster_ratio=cluster_ratio,
                        filter_marketplaces=filter_marketplaces,
                        filter_foreign=filter_foreign,
                        max_variation_pct=max_variation_pct,
                        city=delivery_city,
                        geo_radius_km=geo_radius_km,
                        **({"source_position": source_positions[item[0] - 1]} if source_positions is not None else {}),
                    )
                    for item in chunk_to_process
                ]
                chunk_results = await asyncio.gather(*tasks, return_exceptions=True)
                for input_item, r in zip(chunk_to_process, chunk_results):
                    if isinstance(r, ItemPriceResult):
                        if source_positions is not None:
                            _attach_source_position(r, source_positions[input_item[0] - 1])
                        results.append(r)
                        completed_rows.add(r.row_number)
                    else:
                        logger.error("Item processing exception: %s", r)
                        number, name, unit = input_item
                        results.append(ItemPriceResult(
                            row_number=number, raw_name=name, clean_name=name, unit=unit,
                            status="error", not_found_reason="processing_error",
                            search_diagnostics={"error_type": type(r).__name__},
                        ))
                        if source_positions is not None:
                            _attach_source_position(results[-1], source_positions[number - 1])
                        completed_rows.add(number)

                # Сохраняем промежуточный чекпоинт каждые batch_size строк
                if checkpoint_path:
                    _save_checkpoint(checkpoint_path, results, total, getattr(engine, "total_requests", 0))

            pct = 10 + int((len(completed_rows) / total) * 80)
            if progress_callback:
                progress_callback(pct, f"Обработано {len(completed_rows)} из {total} позиций...")

    await engine.close()

    # Сортируем результаты строго по порядку строк входного файла
    results.sort(key=lambda it: it.row_number)

    # Генерация результирующего Excel-файла
    generate_price_monitoring_excel(
        results,
        output_excel_path,
        max_suppliers=max_suppliers,
        max_variation_pct=max_variation_pct,
        delivery_city=delivery_city,
        geo_radius_km=geo_radius_km,
    )

    # Очистка файла чекпоинта после успешного завершения
    if checkpoint_path and os.path.exists(checkpoint_path):
        try:
            os.remove(checkpoint_path)
        except Exception:
            pass

    found_count = sum(1 for it in results if it.has_prices)
    not_found_count = total - found_count
    total_charged_rub = round(found_count * price_per_position_rub, 2)

    unit_price = float(getattr(settings, "yandex_search_price_per_request", 0.04) or 0.04)
    yandex_requests = int(getattr(engine, "total_requests", 0) or 0)
    yandex_cost = round(yandex_requests * unit_price, 2)

    if progress_callback:
        progress_callback(100, f"Готово! Найдено подтвержденных цен для {found_count} из {total} позиций")

    return BatchMonitoringResult(
        job_id="",
        total_positions=total,
        found_positions=found_count,
        not_found_positions=not_found_count,
        price_per_position_rub=price_per_position_rub,
        total_charged_rub=total_charged_rub,
        items=results,
        output_excel_path=output_excel_path,
        yandex_requests_count=yandex_requests,
        yandex_cost_rub=yandex_cost,
    )

