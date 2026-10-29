"""A price count or SKU shape must not become an unsupported verification claim."""
import openpyxl
import pytest

from app.price_monitoring.cleaner import (
    build_commercial_search_queries,
    build_fallback_broad_queries,
    clean_nomenclature_name,
    extract_generic_nomenclature_codes,
    is_closed_factory_drawing,
    is_drawing_or_part_code,
)
from app.price_monitoring.excel_report import generate_price_monitoring_excel
from app.price_monitoring.models import ItemPriceResult, SupplierQuote


@pytest.mark.parametrize("measurement", ["1.2В", "1,5V", "1000мА/ч", "1500мА·ч"])
def test_electrical_measurement_is_not_a_product_identifier(measurement):
    assert not is_drawing_or_part_code(measurement)
    assert is_drawing_or_part_code("HR03")


@pytest.mark.parametrize("code", [
    "SN61.XXAGHAMAX/EP", "ИРТ-5922-МВ/В/t1050", "385RF-NG-SPT63B",
    "2.607.019.435", "УОНИ-13/55", "ЭА-400/10У", "SPGA4270-200-01",
])
def test_catalog_identifier_is_not_proof_of_a_closed_factory_drawing(code):
    name = f"Изделие {code}"
    assert not is_closed_factory_drawing(code, name)
    assert all(code in query for query in build_commercial_search_queries(name, "шт")[:4])
    assert all(code in query for query in build_fallback_broad_queries(name, "шт")[:4])


@pytest.mark.parametrize("code,name", [
    ("5ВК.568.007.2", "Сальник 5ВК.568.007.2"),
    ("ЭГ.059.00.00-02", "Прокладка ЭГ.059.00.00-02"),
    ("1085.05.02", "Деталь по чертежу 1085.05.02"),
])
def test_drawing_search_stays_exact(code, name):
    assert is_closed_factory_drawing(code, name)
    assert all(f'"{code}"' in query for query in build_commercial_search_queries(name, "шт"))


@pytest.mark.parametrize("token", ["50um", "4-рядн", "19ммх0.15ммх20м"])
def test_measurement_and_descriptive_tokens_are_not_search_skus(token):
    name = f"Товар {token}"
    assert not is_drawing_or_part_code(token, name)
    assert token not in extract_generic_nomenclature_codes(name)
    queries = build_commercial_search_queries(name, "шт")
    assert not any(f'"{token}"' in query for query in queries)


def test_for_prefix_is_normalized_without_losing_real_model_code():
    name = clean_nomenclature_name("Шайба прямоуг.8ВК.951.075 д/ТВФ-100/120")
    assert "д/ТВФ-100/120" not in name
    assert "ТВФ-100/120" in name
    assert "ТВФ-100/120" in extract_generic_nomenclature_codes(name)
    assert is_drawing_or_part_code("ТВФ-100/120", name)


@pytest.mark.parametrize("token", ["HR03", "WD-40", "СТП1-2/6000", "5ВК.568.007.2"])
def test_real_model_and_drawing_codes_remain_identifiers(token):
    assert is_drawing_or_part_code(token, token)


@pytest.mark.parametrize("count", [1, 2, 3])
def test_report_discloses_number_of_sources_without_blanket_confirmation(tmp_path, count):
    item = ItemPriceResult(1, "Товар", "Товар", "шт", status="found")
    item.quotes = [SupplierQuote(company_name=f"Источник {i}", price_rub=100 + i,
                                 url=f"https://supplier{i}.example/item") for i in range(count)]
    path = tmp_path / "report.xlsx"
    generate_price_monitoring_excel([item], str(path))
    wb = openpyxl.load_workbook(path, read_only=True)
    status = wb["Детализированный реестр"].cell(2, 4).value
    assert str(count) in status
    assert "Подтверждено" not in status
    wb.close()


def test_exhausted_search_is_not_reported_as_no_market_offers(tmp_path):
    item = ItemPriceResult(1, "Прибор SN61.XXAGHAMAX/EP", "Прибор SN61.XXAGHAMAX/EP", "шт", status="not_found")
    item.not_found_reason = "budget_exhausted"
    path = tmp_path / "report.xlsx"
    generate_price_monitoring_excel([item], str(path))
    wb = openpyxl.load_workbook(path, read_only=True)
    text = " ".join(str(c.value or "") for row in wb["Детализированный реестр"] for c in row)
    assert "лимит" in text.lower()
    assert "Заказная деталь" not in text
    wb.close()


def test_report_discloses_snippet_evidence_and_retrieval_time(tmp_path):
    quote = SupplierQuote(company_name="Магазин", price_rub=100, url="https://supplier.example/item",
                          evidence_level="search_snippet", retrieved_at="2026-10-21T06:00:00+00:00",
                          source_excerpt="Цена 100 руб. за штуку")
    item = ItemPriceResult(1, "Товар", "Товар", "шт", quotes=[quote], status="found")
    path = tmp_path / "provenance.xlsx"
    generate_price_monitoring_excel([item], str(path))
    wb = openpyxl.load_workbook(path)
    notes = wb["Детализированный реестр"]["N2"].value
    assert "сниппет" in notes.lower()
    assert "2026-10-21" in notes
    comments = [cell.comment.text for row in wb["Пример отчета"] for cell in row if cell.comment]
    assert any("Цена 100 руб. за штуку" in comment for comment in comments)
    wb.close()
