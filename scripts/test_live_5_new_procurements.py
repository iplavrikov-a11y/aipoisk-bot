import asyncio
import logging
import os
import sys
import openpyxl

# Add backend to sys.path
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "backend")))

from app.db import SessionLocal
from app.models import SystemSettings
from app.exact_product.yandex_search import YandexSearchEngine
from app.price_monitoring.pipeline import process_single_item
from app.price_monitoring.excel_report import generate_price_monitoring_excel
import httpx

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
logger = logging.getLogger("test_live_5_new_procurements")

TEST_ITEMS = [
    (1, "Кабель силовой ВВГнг(А)-LS 3х2.5", "М"),
    (2, "Светильник светодиодный 36Вт 595х595 Армстронг", "ШТ"),
    (3, "Краска водно-дисперсионная интерьерная белая 14кг", "ШТ"),
    (4, "Труба профильная 40х20х2 мм", "М"),
    (5, "Цемент М500 Д20 мешок 50 кг", "ШТ"),
]

async def run_live_tests():
    db = SessionLocal()
    settings = db.query(SystemSettings).first()
    if not settings:
        print("ERROR: SystemSettings not found in DB")
        return None

    logger.info("Initializing search engine (folder: %s)...", settings.yandex_search_folder_id)
    engine = YandexSearchEngine(
        folder_id=settings.yandex_search_folder_id,
        api_key=settings.yandex_search_api_key,
    )
    search_sem = asyncio.Semaphore(2)

    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36",
        "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
        "Accept-Language": "ru-RU,ru;q=0.9,en-US;q=0.8,en;q=0.7",
    }
    limits = httpx.Limits(max_keepalive_connections=10, max_connections=20)
    timeout = httpx.Timeout(20.0, connect=10.0)

    results = []

    async with httpx.AsyncClient(headers=headers, limits=limits, timeout=timeout, follow_redirects=True, verify=False) as http_client:
        for it in TEST_ITEMS:
            logger.info(">>> Processing item %d: %s [%s]...", it[0], it[1], it[2])
            res = await process_single_item(
                item_tuple=it,
                engine=engine,
                settings=settings,
                http_client=http_client,
                search_semaphore=search_sem,
                max_suppliers=3,
                min_target_suppliers=3,
                cluster_ratio=1.4,
                filter_marketplaces=False,
            )
            results.append(res)
            disp_qs = [q for q in res.quotes if q.price_rub > 0 and not q.is_outlier]
            logger.info("Item %d complete: %d total quotes, %d in consensus cluster", it[0], len(res.quotes), len(disp_qs))

    await engine.close()
    db.close()

    out_path = "/root/projects/tenderlex/storage/live_test_5_procurements.xlsx"
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    generate_price_monitoring_excel(results, out_path, max_suppliers=3)
    logger.info("Generated Excel report at %s", out_path)

    # Perform automated Excel audit
    wb = openpyxl.load_workbook(out_path, data_only=False)
    ws = wb["Пример отчета"]

    print("\n" + "="*80)
    print("LIVE AUDIT RESULTS (5 NEW PROCUREMENTS):")
    print("="*80)

    circular_found = False
    for r in range(5, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            val = str(ws.cell(row=r, column=c).value or "")
            col_letter = openpyxl.utils.get_column_letter(c)
            cell_ref = f"{col_letter}{r}"
            if val.startswith("="):
                # Check circular reference
                if cell_ref in val:
                    print(f"CRITICAL ERROR: Circular reference in cell {cell_ref}: {val}")
                    circular_found = True

    if not circular_found:
        print("[AUDIT PASS] ZERO circular references detected in any formula!")
    else:
        print("[AUDIT FAIL] Circular reference detected!")

    print("\n" + "-"*80)
    print("ITEM-BY-ITEM DETAILED ANALYSIS:")
    print("-" * 80)

    for idx, res in enumerate(results, start=5):
        disp_quotes = [q for q in res.quotes if q.price_rub > 0 and not q.is_outlier]
        outliers = [q for q in res.quotes if q.is_outlier]
        prices = [q.price_rub for q in disp_quotes]
        ratio = (max(prices) / min(prices)) if (prices and min(prices) > 0) else 0.0

        min_val = ws[f"I{idx}"].value
        avg_val = ws[f"J{idx}"].value
        row_h = ws.row_dimensions[idx].height

        print(f"\nТовар #{res.row_number}: {res.clean_name} [{res.unit}]")
        print(f"  Количество источников в консенсусе: {len(disp_quotes)} (отсечено выбросов: {len(outliers)})")
        print(f"  Цены консенсуса: {prices}")
        print(f"  Разброс (Max/Min): {ratio:.2f}x (порог <= 2.2x)")
        print(f"  Формула Мин. цена (I{idx}): {min_val}")
        print(f"  Формула Средняя цена (J{idx}): {avg_val}")
        print(f"  Высота строки #{idx}: {row_h} pt")

        # Check contact styling in first supplier cell
        d_val = str(ws[f"D{idx}"].value or "")
        d_lines = d_val.split("\n")
        print(f"  Карточка поставщика 1 (строк: {len(d_lines)}):")
        for line in d_lines:
            print(f"    | {line}")

        # Check address absence
        has_address = any(word in d_val.lower() for word in ["ул.", "улица", "д.", "просп.", "шоссе", "область", "край", "поселок"])
        if has_address:
            print("  [WARN] Physical address detected in cell!")
        else:
            print("  [AUDIT PASS] Чисто: физические адреса отсутствуют.")

        if len(disp_quotes) >= 3:
            print("  [AUDIT PASS] Кворум >= 3 подтвержденных поставщиков выполнен!")
        else:
            print(f"  [AUDIT WARN] Найдено {len(disp_quotes)} поставщика (целевой кворум: 3).")

        if ratio <= 2.2 and ratio > 0:
            print("  [AUDIT PASS] Разброс цен в пределах нормы (<= 2.2x).")
        else:
            print(f"  [AUDIT FAIL] Разброс цен {ratio:.2f}x превышает 2.2x.")

if __name__ == "__main__":
    asyncio.run(run_live_tests())
