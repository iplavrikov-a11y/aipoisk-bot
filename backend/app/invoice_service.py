from __future__ import annotations

import base64
import os
import subprocess
import tempfile
from datetime import datetime, timezone
from email.mime.application import MIMEApplication
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from pathlib import Path
from typing import Any

import httpx
import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, Side
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from .config import config
from .models import Invoice, now_utc

SELLER_NAME = "Индивидуальный предприниматель Груздев Игорь Вячеславович"
SELLER_SHORT_NAME = "ИП Груздев И.В."
SELLER_INN = "352516048881"
SELLER_OGRNIP = "323352500038991"
SELLER_ADDRESS = "160541, РОССИЯ, Вологодская область, Вологодский р-н, д. Тарасово, дом 3"
SELLER_ACCOUNT = "40802810217430000342"
SELLER_BANK_NAME = 'АО "АЛЬФА-БАНК"'
SELLER_BANK_BIK = "044525593"
SELLER_BANK_CORR = "30101810200000000593"
SELLER_BANK_INN = "7728168971"
SELLER_BANK_ADDRESS = "162130, Вологодская область, г. Сокол, ул. Советская, д. 87"

DATA_DIR = Path(__file__).resolve().parent / "data" / "invoice"
STAMP_PATH = DATA_DIR / "stamp.png"
SIGNATURE_PATH = DATA_DIR / "signature.png"

RU_MONTHS = [
    "",
    "января",
    "февраля",
    "марта",
    "апреля",
    "мая",
    "июня",
    "июля",
    "августа",
    "сентября",
    "октября",
    "ноября",
    "декабря",
]


def format_date_ru(dt: datetime | None) -> str:
    if not dt:
        dt = now_utc()
    return f"{dt.day} {RU_MONTHS[dt.month]} {dt.year}"


def num_to_words_ru(amount_rub: int | float) -> str:
    """Convert rubles into official Russian legal text representation with 'Без НДС.'"""
    units_m = ["", "один", "два", "три", "четыре", "пять", "шесть", "семь", "восемь", "девять"]
    units_f = ["", "одна", "две", "три", "четыре", "пять", "шесть", "семь", "восемь", "девять"]
    teens = [
        "десять",
        "одиннадцать",
        "двенадцать",
        "тринадцать",
        "четырнадцать",
        "пятнадцать",
        "шестнадцать",
        "семнадцать",
        "восемнадцать",
        "девятнадцать",
    ]
    tens = [
        "",
        "",
        "двадцать",
        "тридцать",
        "сорок",
        "пятьдесят",
        "шестьдесят",
        "семьдесят",
        "восемьдесят",
        "девяносто",
    ]
    hundreds = [
        "",
        "сто",
        "двести",
        "триста",
        "четыреста",
        "пятьсот",
        "шестьсот",
        "семьсот",
        "восемьсот",
        "девятьсот",
    ]

    def triplet(n: int, feminine: bool = False) -> str:
        words = []
        h = n // 100
        t = (n % 100) // 10
        u = n % 10
        if h > 0:
            words.append(hundreds[h])
        if t == 1:
            words.append(teens[u])
        else:
            if t > 1:
                words.append(tens[t])
            if u > 0:
                words.append(units_f[u] if feminine else units_m[u])
        return " ".join(words)

    def plural(n: int, form1: str, form2: str, form5: str) -> str:
        n10 = n % 10
        n100 = n % 100
        if 11 <= n100 <= 19:
            return form5
        if n10 == 1:
            return form1
        if 2 <= n10 <= 4:
            return form2
        return form5

    rub = int(amount_rub)
    kop = int(round((amount_rub - rub) * 100))

    if rub == 0:
        rub_str = "ноль рублей"
    else:
        parts = []
        millions = (rub // 1_000_000) % 1000
        if millions > 0:
            parts.append(triplet(millions, False) + " " + plural(millions, "миллион", "миллиона", "миллионов"))
        thousands = (rub // 1_000) % 1000
        if thousands > 0:
            parts.append(triplet(thousands, True) + " " + plural(thousands, "тысяча", "тысячи", "тысяч"))
        remainder = rub % 1000
        if remainder > 0 or not parts:
            parts.append(triplet(remainder, False))
        rub_str = " ".join(p for p in parts if p).strip() + " " + plural(rub, "рубль", "рубля", "рублей")

    kop_str = f"{kop:02d} " + plural(kop, "копейка", "копейки", "копеек")
    full = f"{rub_str.strip()} {kop_str}. Без НДС."
    return full[0].upper() + full[1:]


def format_rubles_num(amount: int | float) -> str:
    """Format rubles with space thousand separator and two decimals: '10 000,00'"""
    val = float(amount)
    rub = int(val)
    kop = int(round((val - rub) * 100))
    rub_formatted = f"{rub:,}".replace(",", " ")
    return f"{rub_formatted},{kop:02d}"


def next_invoice_number(db: Session) -> str:
    """Generate sequential invoice number: TL-2026-0001"""
    year = now_utc().year
    prefix = f"TL-{year}-"
    stmt = (
        select(Invoice.invoice_number)
        .where(Invoice.invoice_number.like(f"{prefix}%"))
        .order_by(Invoice.invoice_number.desc())
        .limit(1)
    )
    last_num_str = db.scalar(stmt)
    if last_num_str and last_num_str.startswith(prefix):
        try:
            current_seq = int(last_num_str[len(prefix) :])
            next_seq = current_seq + 1
        except Exception:
            next_seq = 1
    else:
        next_seq = 1
    return f"{prefix}{next_seq:04d}"


def generate_invoice_xlsx(invoice: Invoice, output_path: str) -> str:
    """Generate authentic 1C-styled Excel (.xlsx) file with openpyxl, overlaying stamp and signature."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Счет на оплату"
    ws.views.sheetView[0].showGridLines = True

    # Page setup
    ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # Column widths matching standard 1C form
    col_widths = {
        "A": 4.5,
        "B": 16.0,
        "C": 12.0,
        "D": 12.0,
        "E": 10.0,
        "F": 12.0,
        "G": 8.0,
        "H": 8.0,
        "I": 10.0,
        "J": 6.0,
        "K": 6.0,
        "L": 14.0,
        "M": 18.0,
    }
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    font_family = "Arial"
    f_tiny = Font(name=font_family, size=8)
    f_small = Font(name=font_family, size=9)
    f_small_bold = Font(name=font_family, size=9, bold=True)
    f_norm = Font(name=font_family, size=10)
    f_norm_bold = Font(name=font_family, size=10, bold=True)
    f_title = Font(name=font_family, size=14, bold=True)

    thin_border_side = Side(border_style="thin", color="000000")
    thick_border_side = Side(border_style="medium", color="000000")

    # Row 1: Warning Notice
    ws.merge_cells("A1:M1")
    ws["A1"] = (
        "Внимание! Оплата данного счета означает согласие с условиями оказания услуг сервиса TenderLex "
        "(оферта tenderlex.ru/terms). Услуги оказываются по факту зачисления денежных средств на р/с Исполнителя."
    )
    ws["A1"].font = f_tiny
    ws["A1"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[1].height = 28

    # Rows 2..7: Bank details box
    ws.merge_cells("A2:F3")
    ws["A2"] = SELLER_BANK_NAME
    ws["A2"].font = f_norm
    ws["A2"].alignment = Alignment(vertical="top")

    ws["G2"] = "БИК"
    ws["G2"].font = f_small
    ws["G2"].alignment = Alignment(vertical="center")

    ws.merge_cells("I2:M2")
    ws["I2"] = SELLER_BANK_BIK
    ws["I2"].font = f_norm
    ws["I2"].alignment = Alignment(vertical="center")

    ws["G3"] = "Сч. №"
    ws["G3"].font = f_small
    ws["G3"].alignment = Alignment(vertical="center")

    ws.merge_cells("I3:M3")
    ws["I3"] = SELLER_BANK_CORR
    ws["I3"].font = f_norm
    ws["I3"].alignment = Alignment(vertical="center")

    ws.merge_cells("A4:F4")
    ws["A4"] = "Банк получателя"
    ws["A4"].font = f_tiny
    ws["A4"].alignment = Alignment(vertical="top")

    ws.merge_cells("A5:D5")
    ws["A5"] = f"ИНН {SELLER_INN}"
    ws["A5"].font = f_small
    ws["A5"].alignment = Alignment(vertical="center")

    ws.merge_cells("E5:F5")
    ws["E5"] = "КПП"
    ws["E5"].font = f_small
    ws["E5"].alignment = Alignment(vertical="center")

    ws["G5"] = "Сч. №"
    ws["G5"].font = f_small
    ws["G5"].alignment = Alignment(vertical="center")

    ws.merge_cells("I5:M5")
    ws["I5"] = SELLER_ACCOUNT
    ws["I5"].font = f_norm
    ws["I5"].alignment = Alignment(vertical="center")

    ws.merge_cells("A6:F6")
    ws["A6"] = SELLER_NAME
    ws["A6"].font = f_norm
    ws["A6"].alignment = Alignment(vertical="center")

    ws.merge_cells("A7:F7")
    ws["A7"] = "Получатель"
    ws["A7"].font = f_tiny
    ws["A7"].alignment = Alignment(vertical="top")

    # Apply borders to bank details
    for r in range(2, 8):
        for c in range(1, 14):
            cell = ws.cell(row=r, column=c)
            top = thin_border_side if r in (2, 4, 5, 7) else None
            bottom = thin_border_side if r in (3, 4, 6, 7) else None
            left = thin_border_side if c in (1, 7, 9) else None
            right = thin_border_side if c in (6, 8, 13) else None
            cell.border = Border(top=top, bottom=bottom, left=left, right=right)

    # Row 9: Title
    date_str = format_date_ru(invoice.created_at)
    ws.merge_cells("A9:M9")
    ws["A9"] = f"Счет на оплату № {invoice.invoice_number} от {date_str} г."
    ws["A9"].font = f_title
    ws["A9"].alignment = Alignment(vertical="bottom")
    for c in range(1, 14):
        ws.cell(row=9, column=c).border = Border(bottom=thick_border_side)
    ws.row_dimensions[9].height = 28

    # Row 11: Supplier
    ws.merge_cells("A11:B11")
    ws["A11"] = "Поставщик:"
    ws["A11"].font = f_small
    ws.merge_cells("C11:M11")
    ws["C11"] = f"{SELLER_NAME}, ИНН {SELLER_INN}, ОГРНИП {SELLER_OGRNIP}, {SELLER_ADDRESS}"
    ws["C11"].font = f_norm_bold
    ws["C11"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[11].height = 28

    # Row 13: Buyer
    buyer_parts = [invoice.buyer_name]
    if invoice.buyer_inn:
        buyer_parts.append(f"ИНН {invoice.buyer_inn}")
    if invoice.buyer_kpp:
        buyer_parts.append(f"КПП {invoice.buyer_kpp}")
    if invoice.buyer_address:
        buyer_parts.append(invoice.buyer_address)
    buyer_full = ", ".join(p for p in buyer_parts if p)

    ws.merge_cells("A13:B13")
    ws["A13"] = "Покупатель:"
    ws["A13"].font = f_small
    ws.merge_cells("C13:M13")
    ws["C13"] = buyer_full
    ws["C13"].font = f_norm_bold
    ws["C13"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[13].height = 26

    # Row 15: Base
    base_text = f"Предоставление доступа к сервису TenderLex (пополнение баланса аккаунта {invoice.buyer_email or ''})"
    ws.merge_cells("A15:B15")
    ws["A15"] = "Основание:"
    ws["A15"].font = f_small
    ws.merge_cells("C15:M15")
    ws["C15"] = base_text
    ws["C15"].font = f_small_bold
    ws["C15"].alignment = Alignment(wrap_text=True, vertical="top")

    # Row 17: Items Table Header
    headers = [
        ("A17", "№"),
        ("B17:G17", "Товары (работы, услуги)"),
        ("H17:I17", "Кол-во"),
        ("J17:K17", "Ед. изм."),
        ("L17", "Цена"),
        ("M17", "Сумма"),
    ]
    for ref, text in headers:
        if ":" in ref:
            ws.merge_cells(ref)
            top_left = ref.split(":")[0]
        else:
            top_left = ref
        ws[top_left] = text
        ws[top_left].font = f_small_bold
        ws[top_left].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for c in range(1, 14):
        ws.cell(row=17, column=c).border = Border(
            top=thick_border_side,
            bottom=thick_border_side,
            left=thin_border_side,
            right=thin_border_side,
        )
    ws.row_dimensions[17].height = 24

    # Row 18: Service Item
    amount_str = format_rubles_num(invoice.amount_rub)
    tariff_label = invoice.tariff_name or "Пополнение баланса"
    item_desc = f"{invoice.service_name} ({tariff_label}, {invoice.amount_rub:,} руб.)".replace(",", " ")

    ws["A18"] = "1"
    ws["A18"].alignment = Alignment(horizontal="center")
    ws["A18"].font = f_small

    ws.merge_cells("B18:G18")
    ws["B18"] = item_desc
    ws["B18"].font = f_small
    ws["B18"].alignment = Alignment(vertical="center", wrap_text=True)

    ws.merge_cells("H18:I18")
    ws["H18"] = "1"
    ws["H18"].font = f_small
    ws["H18"].alignment = Alignment(horizontal="right", vertical="center")

    ws.merge_cells("J18:K18")
    ws["J18"] = "усл"
    ws["J18"].font = f_small
    ws["J18"].alignment = Alignment(horizontal="center", vertical="center")

    ws["L18"] = amount_str
    ws["L18"].font = f_small
    ws["L18"].alignment = Alignment(horizontal="right", vertical="center")

    ws["M18"] = amount_str
    ws["M18"].font = f_small
    ws["M18"].alignment = Alignment(horizontal="right", vertical="center")

    for c in range(1, 14):
        ws.cell(row=18, column=c).border = Border(
            top=thin_border_side,
            bottom=thick_border_side,
            left=thin_border_side,
            right=thin_border_side,
        )
    ws.row_dimensions[18].height = 24

    # Row 19: Totals
    ws.merge_cells("A19:L19")
    ws["A19"] = "Итого:"
    ws["A19"].font = f_norm_bold
    ws["A19"].alignment = Alignment(horizontal="right")
    ws["M19"] = amount_str
    ws["M19"].font = f_norm_bold
    ws["M19"].alignment = Alignment(horizontal="right")

    ws.merge_cells("A20:L20")
    ws["A20"] = "Без налога (НДС)"
    ws["A20"].font = f_norm_bold
    ws["A20"].alignment = Alignment(horizontal="right")
    ws["M20"] = "-"
    ws["M20"].font = f_norm_bold
    ws["M20"].alignment = Alignment(horizontal="center")

    ws.merge_cells("A21:M21")
    ws["A21"] = f"Всего наименований 1, на сумму {amount_str} руб."
    ws["A21"].font = f_norm
    ws.row_dimensions[21].height = 20

    words_str = num_to_words_ru(invoice.amount_rub)
    ws.merge_cells("A22:M22")
    ws["A22"] = words_str
    ws["A22"].font = f_norm_bold
    ws.row_dimensions[22].height = 20

    for c in range(1, 14):
        ws.cell(row=23, column=c).border = Border(bottom=thick_border_side)
    ws.row_dimensions[23].height = 8

    # Row 25: Signatures
    ws.merge_cells("A25:B25")
    ws["A25"] = "Руководитель"
    ws["A25"].font = f_norm_bold

    ws.merge_cells("C25:E25")
    for c in range(3, 6):
        ws.cell(row=25, column=c).border = Border(bottom=thin_border_side)

    ws.merge_cells("F25:G25")
    ws["F25"] = "Груздев И.В."
    ws["F25"].font = f_norm_bold

    ws.merge_cells("H25:I25")
    ws["H25"] = "Бухгалтер"
    ws["H25"].font = f_norm_bold

    ws.merge_cells("J25:K25")
    for c in range(10, 12):
        ws.cell(row=25, column=c).border = Border(bottom=thin_border_side)

    ws.merge_cells("L25:M25")
    ws["L25"] = "Груздев И.В."
    ws["L25"].font = f_norm_bold
    ws.row_dimensions[25].height = 36

    # Insert signature and stamp
    if SIGNATURE_PATH.exists():
        img_sig = XLImage(str(SIGNATURE_PATH))
        img_sig.width = 100
        img_sig.height = 70
        ws.add_image(img_sig, "C24")

        img_sig_acc = XLImage(str(SIGNATURE_PATH))
        img_sig_acc.width = 100
        img_sig_acc.height = 70
        ws.add_image(img_sig_acc, "J24")

    if STAMP_PATH.exists():
        img_stamp = XLImage(str(STAMP_PATH))
        img_stamp.width = 135
        img_stamp.height = 135
        ws.add_image(img_stamp, "D24")

    out_file = Path(output_path)
    out_file.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out_file))
    return str(out_file)


def generate_invoice_pdf(invoice: Invoice, output_pdf_path: str) -> str:
    """Generate high-resolution A4 PDF invoice using headless Chromium, falling back to LibreOffice."""
    out_pdf = Path(output_pdf_path)
    out_pdf.parent.mkdir(parents=True, exist_ok=True)

    html_content = render_invoice_html(invoice, embed_images=True)
    with tempfile.TemporaryDirectory() as tmp_dir:
        tmp_html = Path(tmp_dir) / f"{invoice.invoice_number}.html"
        tmp_html.write_text(html_content, encoding="utf-8")

        # 1. Try headless Google Chrome (clean vector A4 PDF without browser headers/footers)
        chrome_bin = "/usr/bin/google-chrome"
        if os.path.exists(chrome_bin):
            cmd = [
                chrome_bin,
                "--headless=new",
                "--no-sandbox",
                "--disable-gpu",
                "--no-pdf-header-footer",
                f"--print-to-pdf={out_pdf}",
                str(tmp_html),
            ]
            try:
                res = subprocess.run(cmd, capture_output=True, text=True, timeout=15)
                if res.returncode == 0 and out_pdf.exists() and out_pdf.stat().st_size > 1000:
                    return str(out_pdf)
            except Exception:
                pass

        # 2. Fallback to LibreOffice Calc from XLSX
        tmp_xlsx = Path(tmp_dir) / f"{invoice.invoice_number}.xlsx"
        generate_invoice_xlsx(invoice, str(tmp_xlsx))
        cmd_soffice = [
            "soffice",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            tmp_dir,
            str(tmp_xlsx),
        ]
        res_soffice = subprocess.run(cmd_soffice, capture_output=True, text=True, timeout=30)
        if res_soffice.returncode != 0:
            raise RuntimeError(f"PDF conversion failed: {res_soffice.stderr or res_soffice.stdout}")
        generated_pdf = Path(tmp_dir) / f"{invoice.invoice_number}.pdf"
        if not generated_pdf.exists():
            raise FileNotFoundError(f"Generated PDF not found: {generated_pdf}")
        generated_pdf.replace(out_pdf)
        return str(out_pdf)


def render_invoice_html(invoice: Invoice, *, embed_images: bool = True) -> str:
    """Generate authentic A4 HTML representation for web viewing and browser printing (@media print)."""
    date_str = format_date_ru(invoice.created_at)
    amount_str = format_rubles_num(invoice.amount_rub)
    words_str = num_to_words_ru(invoice.amount_rub)
    tariff_label = invoice.tariff_name or "Пополнение баланса"
    item_desc = f"{invoice.service_name} ({tariff_label}, {invoice.amount_rub:,} руб.)".replace(",", " ")

    buyer_parts = [invoice.buyer_name]
    if invoice.buyer_inn:
        buyer_parts.append(f"ИНН {invoice.buyer_inn}")
    if invoice.buyer_kpp:
        buyer_parts.append(f"КПП {invoice.buyer_kpp}")
    if invoice.buyer_address:
        buyer_parts.append(invoice.buyer_address)
    buyer_full = ", ".join(p for p in buyer_parts if p)

    stamp_src = "/invoice/stamp.png"
    sig_src = "/invoice/signature.png"

    if embed_images:
        if STAMP_PATH.exists():
            b64_stamp = base64.b64encode(STAMP_PATH.read_bytes()).decode("ascii")
            stamp_src = f"data:image/png;base64,{b64_stamp}"
        if SIGNATURE_PATH.exists():
            b64_sig = base64.b64encode(SIGNATURE_PATH.read_bytes()).decode("ascii")
            sig_src = f"data:image/png;base64,{b64_sig}"

    return f"""<!DOCTYPE html>
<html lang="ru">
<head>
  <meta charset="UTF-8">
  <title>Счет на оплату № {invoice.invoice_number} от {date_str} г.</title>
  <style>
    * {{ box-sizing: border-box; }}
    body {{
      font-family: Arial, "Helvetica Neue", Helvetica, sans-serif;
      font-size: 10pt;
      line-height: 1.3;
      color: #000;
      background: #f1f5f9;
      margin: 0;
      padding: 24px 12px 60px 12px;
      display: flex;
      flex-direction: column;
      align-items: center;
      -webkit-print-color-adjust: exact;
      print-color-adjust: exact;
    }}
    .no-print-toolbar {{
      width: 210mm;
      max-width: 100%;
      margin: 0 auto 16px auto;
      display: flex;
      gap: 10px;
      align-items: center;
      justify-content: space-between;
      flex-wrap: wrap;
    }}
    .btn {{
      padding: 9px 16px;
      border-radius: 8px;
      font-size: 13px;
      font-weight: 700;
      cursor: pointer;
      text-decoration: none;
      display: inline-flex;
      align-items: center;
      gap: 6px;
      border: 1px solid transparent;
      transition: background-color 0.15s ease;
    }}
    .btn-primary {{
      background: #0d9488;
      color: #fff;
    }}
    .btn-primary:hover {{ background: #0f766e; }}
    .btn-emerald {{
      background: #059669;
      color: #fff;
    }}
    .btn-emerald:hover {{ background: #047857; }}
    .btn-outline {{
      background: #fff;
      color: #334155;
      border-color: #cbd5e1;
    }}
    .btn-outline:hover {{ background: #f8fafc; border-color: #94a3b8; }}

    .invoice-card {{
      width: 210mm;
      min-height: 297mm;
      max-width: 100%;
      margin: 0 auto;
      background: #ffffff;
      padding: 16mm 18mm 20mm 18mm;
      box-shadow: 0 4px 20px rgba(0, 0, 0, 0.12);
      border: 1px solid #cbd5e1;
      position: relative;
    }}

    .notice {{
      font-size: 8pt;
      color: #475569;
      margin-bottom: 14px;
      line-height: 1.25;
    }}
    table.bank-box {{
      width: 100%;
      border-collapse: collapse;
      margin-bottom: 22px;
      font-size: 9.5pt;
    }}
    table.bank-box td {{
      border: 1px solid #000;
      padding: 4px 7px;
      vertical-align: middle;
    }}
    .subtext {{
      font-size: 7.5pt;
      color: #444;
      display: block;
    }}
    h1.title {{
      font-size: 15pt;
      font-weight: bold;
      border-bottom: 2px solid #000;
      padding-bottom: 8px;
      margin: 22px 0 16px 0;
    }}
    .info-row {{
      display: flex;
      margin-bottom: 10px;
      font-size: 9.5pt;
      line-height: 1.35;
    }}
    .info-label {{
      width: 115px;
      flex-shrink: 0;
      color: #334155;
    }}
    .info-value {{
      font-weight: bold;
      flex: 1;
    }}
    table.items-table {{
      width: 100%;
      border-collapse: collapse;
      margin-top: 18px;
      margin-bottom: 14px;
      font-size: 9.5pt;
    }}
    table.items-table th, table.items-table td {{
      border: 1px solid #000;
      padding: 6px 8px;
    }}
    table.items-table th {{
      border-top: 2px solid #000;
      border-bottom: 2px solid #000;
      background: #fafafa;
      text-align: center;
      font-weight: bold;
    }}
    .totals-area {{
      margin-top: 10px;
      font-size: 10pt;
    }}
    .totals-row {{
      display: flex;
      justify-content: flex-end;
      gap: 24px;
      font-weight: bold;
      padding: 3px 0;
    }}
    .words-block {{
      margin-top: 14px;
      font-weight: bold;
      font-size: 10pt;
    }}
    .divider {{
      border-bottom: 2px solid #000;
      margin: 16px 0 28px 0;
    }}
    .signatures-block {{
      position: relative;
      display: flex;
      justify-content: space-between;
      margin-top: 40px;
      font-size: 10pt;
      font-weight: bold;
    }}
    .sig-col {{
      display: flex;
      align-items: flex-end;
      gap: 10px;
      position: relative;
    }}
    .sig-line {{
      display: inline-block;
      width: 130px;
      border-bottom: 1px solid #000;
      height: 18px;
    }}
    .stamp-overlay {{
      position: absolute;
      left: 110px;
      top: -45px;
      width: 135px;
      height: 135px;
      pointer-events: none;
      z-index: 10;
    }}
    .signature-overlay {{
      position: absolute;
      left: 95px;
      top: -20px;
      width: 100px;
      height: 70px;
      pointer-events: none;
      z-index: 5;
    }}

    @media print {{
      body {{
        background: #fff !important;
        padding: 0 !important;
        margin: 0 !important;
      }}
      .no-print-toolbar {{
        display: none !important;
      }}
      .invoice-card {{
        width: 100% !important;
        max-width: 100% !important;
        min-height: auto !important;
        margin: 0 !important;
        padding: 0 !important;
        border: none !important;
        box-shadow: none !important;
      }}
      @page {{
        size: A4 portrait;
        margin: 12mm 15mm 12mm 15mm;
      }}
    }}
  </style>
</head>
<body>

  <div class="no-print-toolbar">
    <div style="font-size: 13px; color: #475569; font-weight: 500;">
      Счёт <b>№ {invoice.invoice_number}</b> от {date_str} г.
    </div>
    <div style="display: flex; gap: 8px; flex-wrap: wrap;">
      <button class="btn btn-primary" onclick="window.print()">
        🖨️ Распечатать
      </button>
      <a class="btn btn-emerald" href="/api/customer/invoices/{invoice.id}/download?format=pdf" download>
        📥 Скачать PDF
      </a>
      <a class="btn btn-outline" href="/api/customer/invoices/{invoice.id}/download?format=xlsx" download>
        📊 Скачать Excel (.xlsx)
      </a>
    </div>
  </div>

  <div class="invoice-card">
    <div class="notice">
      Внимание! Оплата данного счета означает согласие с условиями оказания услуг сервиса TenderLex (оферта tenderlex.ru/terms).<br>
      Услуги оказываются по факту зачисления денежных средств на р/с Исполнителя.
    </div>

    <table class="bank-box">
      <tr>
        <td colspan="2" rowspan="2" style="width: 55%; vertical-align: top;">
          {SELLER_BANK_NAME}<br>
          <span class="subtext">Банк получателя</span>
        </td>
        <td style="width: 12%;">БИК</td>
        <td style="width: 33%;">{SELLER_BANK_BIK}</td>
      </tr>
      <tr>
        <td>Сч. №</td>
        <td>{SELLER_BANK_CORR}</td>
      </tr>
      <tr>
        <td style="width: 35%;">ИНН {SELLER_INN}</td>
        <td style="width: 20%;">КПП —</td>
        <td rowspan="2">Сч. №</td>
        <td rowspan="2">{SELLER_ACCOUNT}</td>
      </tr>
      <tr>
        <td colspan="2">
          {SELLER_NAME}<br>
          <span class="subtext">Получатель</span>
        </td>
      </tr>
    </table>

    <h1 class="title">Счет на оплату № {invoice.invoice_number} от {date_str} г.</h1>

    <div class="info-row">
      <div class="info-label">Поставщик:</div>
      <div class="info-value">{SELLER_NAME}, ИНН {SELLER_INN}, ОГРНИП {SELLER_OGRNIP}, {SELLER_ADDRESS}</div>
    </div>

    <div class="info-row">
      <div class="info-label">Покупатель:</div>
      <div class="info-value">{buyer_full}</div>
    </div>

    <div class="info-row">
      <div class="info-label">Основание:</div>
      <div class="info-value">Предоставление доступа к сервису TenderLex (пополнение баланса аккаунта {invoice.buyer_email or ''})</div>
    </div>

    <table class="items-table">
      <thead>
        <tr>
          <th style="width: 35px;">№</th>
          <th>Товары (работы, услуги)</th>
          <th style="width: 60px;">Кол-во</th>
          <th style="width: 50px;">Ед.</th>
          <th style="width: 100px;">Цена</th>
          <th style="width: 110px;">Сумма</th>
        </tr>
      </thead>
      <tbody>
        <tr>
          <td style="text-align: center;">1</td>
          <td>{item_desc}</td>
          <td style="text-align: right;">1</td>
          <td style="text-align: center;">усл</td>
          <td style="text-align: right;">{amount_str}</td>
          <td style="text-align: right;">{amount_str}</td>
        </tr>
      </tbody>
    </table>

    <div class="totals-area">
      <div class="totals-row">
        <span>Итого:</span>
        <span style="width: 110px; text-align: right;">{amount_str}</span>
      </div>
      <div class="totals-row">
        <span>Без налога (НДС):</span>
        <span style="width: 110px; text-align: right;">—</span>
      </div>
    </div>

    <div style="margin-top: 10px; font-size: 9.5pt;">
      Всего наименований 1, на сумму {amount_str} руб.
    </div>

    <div class="words-block">
      {words_str}
    </div>

    <div class="divider"></div>

    <div class="signatures-block">
      <div class="sig-col">
        <span>Руководитель</span>
        <span class="sig-line"></span>
        <span>Груздев И.В.</span>
        <img class="signature-overlay" src="{sig_src}" alt="Подпись">
        <img class="stamp-overlay" src="{stamp_src}" alt="Печать">
      </div>

      <div class="sig-col">
        <span>Бухгалтер</span>
        <span class="sig-line"></span>
        <span>Груздев И.В.</span>
        <img class="signature-overlay" src="{sig_src}" alt="Подпись" style="left: 80px;">
      </div>
    </div>
  </div>

</body>
</html>"""


def send_invoice_email(invoice: Invoice, recipient_email: str) -> tuple[bool, str]:
    """Send generated invoice with attached Excel and PDF files to recipient."""
    recipient = str(recipient_email or invoice.buyer_email or "").strip()
    if not recipient or "@" not in recipient:
        return False, "Не указан корректный email получателя"

    date_str = format_date_ru(invoice.created_at)
    amount_str = format_rubles_num(invoice.amount_rub)
    subject = f"Счёт на оплату № {invoice.invoice_number} от {date_str} - TenderLex"

    body_html = f"""<html>
<body style="font-family: Arial, sans-serif; font-size: 14px; color: #1e293b; line-height: 1.5;">
  <p>Здравствуйте!</p>
  <p>Вы сформировали счёт на оплату в сервисе <b>TenderLex</b>.</p>
  <div style="background: #f8fafc; border: 1px solid #e2e8f0; border-radius: 10px; padding: 16px; margin: 16px 0;">
    <p style="margin: 0 0 8px 0;"><b>Номер счёта:</b> {invoice.invoice_number}</p>
    <p style="margin: 0 0 8px 0;"><b>Дата:</b> {date_str} г.</p>
    <p style="margin: 0 0 8px 0;"><b>Сумма к оплате:</b> {amount_str} ₽</p>
    <p style="margin: 0 0 8px 0;"><b>К зачислению на баланс:</b> {invoice.total_credit_rub:,} ₽ (бонус: +{invoice.bonus_rub:,} ₽)</p>
    <p style="margin: 0;"><b>Покупатель:</b> {invoice.buyer_name} (ИНН {invoice.buyer_inn})</p>
  </div>

  <p><b>Реквизиты для оплаты через банк-клиент:</b></p>
  <table style="font-size: 13px; border-collapse: collapse; margin-bottom: 16px;">
    <tr><td style="padding: 4px 8px 4px 0; color: #64748b;">Получатель:</td><td><b>ИП Груздев Игорь Вячеславович</b></td></tr>
    <tr><td style="padding: 4px 8px 4px 0; color: #64748b;">ИНН:</td><td><b>352516048881</b></td></tr>
    <tr><td style="padding: 4px 8px 4px 0; color: #64748b;">Расчётный счёт:</td><td><b>40802810217430000342</b></td></tr>
    <tr><td style="padding: 4px 8px 4px 0; color: #64748b;">Банк:</td><td>АО "АЛЬФА-БАНК"</td></tr>
    <tr><td style="padding: 4px 8px 4px 0; color: #64748b;">БИК:</td><td><b>044525593</b></td></tr>
    <tr><td style="padding: 4px 8px 4px 0; color: #64748b;">Корр. счёт:</td><td>30101810200000000593</td></tr>
    <tr><td style="padding: 4px 8px 4px 0; color: #64748b;">Назначение платежа:</td><td><b>Оплата по счету № {invoice.invoice_number} от {date_str} г. за услуги сервиса TenderLex. Без НДС.</b></td></tr>
  </table>

  <p>Оригинал счёта в формате Excel (.xlsx) и PDF с факсимиле подписи и печати прикреплены к этому письму.</p>
  <p>Также вы можете открыть или распечатать счёт онлайн по ссылке:<br>
  <a href="{config.public_base_url}/cabinet" style="color: #0d9488; font-weight: bold;">Перейти в личный кабинет TenderLex</a></p>

  <p style="color: #64748b; font-size: 12px; margin-top: 24px;">Баланс аккаунта пополняется автоматически после зачисления средств на расчётный счёт (обычно в течение 1 рабочего дня).</p>
</body>
</html>"""

    try:
        with tempfile.TemporaryDirectory() as tmp_dir:
            xlsx_path = Path(tmp_dir) / f"schet_{invoice.invoice_number}.xlsx"
            generate_invoice_xlsx(invoice, str(xlsx_path))
            xlsx_bytes = xlsx_path.read_bytes()

            pdf_path = Path(tmp_dir) / f"schet_{invoice.invoice_number}.pdf"
            generate_invoice_pdf(invoice, str(pdf_path))
            pdf_bytes = pdf_path.read_bytes()

            # 1. Try sending via Email Relay (Postfix+DKIM with high deliverability)
            relay_url = str(config.email_relay_url or "").strip()
            relay_api_key = str(config.email_relay_api_key or "").strip()
            if relay_url and relay_api_key:
                attachments = [
                    {
                        "filename": f"schet_{invoice.invoice_number}.xlsx",
                        "content": base64.b64encode(xlsx_bytes).decode("ascii"),
                        "contentType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    },
                    {
                        "filename": f"schet_{invoice.invoice_number}.pdf",
                        "content": base64.b64encode(pdf_bytes).decode("ascii"),
                        "contentType": "application/pdf",
                    },
                ]
                payload = {
                    "to": recipient,
                    "subject": subject,
                    "html": body_html,
                    "from_name": config.email_from_name or "TenderLex",
                    "from_email": config.email_from_email or config.smtp_from or "info@tenderlex.ru",
                    "attachments": attachments,
                }
                with httpx.Client(timeout=25.0) as client:
                    resp = client.post(
                        f"{relay_url.rstrip('/')}/send",
                        headers={
                            "Content-Type": "application/json",
                            "Authorization": f"Bearer {relay_api_key}",
                        },
                        json=payload,
                    )
                    if resp.status_code == 200:
                        return True, "Счёт успешно отправлен на email (прикреплены Excel и PDF)"

            # 2. Fallback to direct SMTP
            smtp_host = config.smtp_host or "smtp.jino.ru"
            smtp_port = config.smtp_port or 587
            smtp_user = config.smtp_username or "info@tenderlex.ru"
            smtp_pass = config.smtp_password
            smtp_from = config.smtp_from or "info@tenderlex.ru"

            msg = MIMEMultipart()
            msg["Subject"] = subject
            msg["From"] = f"TenderLex <{smtp_from}>"
            msg["To"] = recipient
            msg.attach(MIMEText(body_html, "html", "utf-8"))

            attach_xl = MIMEApplication(xlsx_bytes, _subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            attach_xl.add_header("Content-Disposition", "attachment", filename=f"schet_{invoice.invoice_number}.xlsx")
            msg.attach(attach_xl)

            attach_p = MIMEApplication(pdf_bytes, _subtype="pdf")
            attach_p.add_header("Content-Disposition", "attachment", filename=f"schet_{invoice.invoice_number}.pdf")
            msg.attach(attach_p)

            import smtplib
            if config.smtp_use_ssl:
                server = smtplib.SMTP_SSL(smtp_host, smtp_port, timeout=config.smtp_timeout_seconds)
            else:
                server = smtplib.SMTP(smtp_host, smtp_port, timeout=config.smtp_timeout_seconds)
                if config.smtp_use_tls:
                    server.starttls()

            if smtp_user and smtp_pass:
                server.login(smtp_user, smtp_pass)

            server.send_message(msg)
            server.quit()
            return True, "Счёт успешно отправлен на email"
    except Exception as e:
        return False, f"Ошибка отправки email: {e}"


async def notify_admin_telegram(invoice: Invoice) -> None:
    """Send instant Telegram alert to admin when a new B2B invoice is generated."""
    bot_token = config.bot_token
    owner_id = config.owner_telegram_id
    if not bot_token or not owner_id:
        return

    date_str = format_date_ru(invoice.created_at)
    amount_str = format_rubles_num(invoice.amount_rub)
    bonus_str = format_rubles_num(invoice.bonus_rub)
    total_str = format_rubles_num(invoice.total_credit_rub)

    text = (
        f"🧾 <b>Выставлен счёт на оплату юридическому лицу!</b>\n\n"
        f"<b>Номер:</b> <code>{invoice.invoice_number}</code> от {date_str}\n"
        f"<b>Сумма к оплате:</b> <b>{amount_str} ₽</b>\n"
        f"<b>Поступит на баланс:</b> {total_str} ₽ (бонус: +{bonus_str} ₽)\n"
        f"<b>Тариф:</b> {invoice.tariff_name}\n\n"
        f"<b>Покупатель:</b> {invoice.buyer_name}\n"
        f"<b>ИНН:</b> <code>{invoice.buyer_inn}</code>"
        f"{f' | КПП: {invoice.buyer_kpp}' if invoice.buyer_kpp else ''}\n"
        f"<b>Адрес:</b> {invoice.buyer_address}\n"
        f"<b>Email клиента:</b> {invoice.buyer_email}\n\n"
        f"<i>После поступления средств на р/с в Альфа-Банке подтвердите зачисление в админке или пополните баланс клиента.</i>"
    )

    url = f"https://api.telegram.org/bot{bot_token}/sendMessage"
    payload = {
        "chat_id": owner_id,
        "text": text,
        "parse_mode": "HTML",
        "disable_web_page_preview": True,
    }

    try:
        async with httpx.AsyncClient(timeout=8.0) as client:
            await client.post(url, json=payload)
    except Exception:
        pass
