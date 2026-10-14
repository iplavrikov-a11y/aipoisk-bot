import unittest
from app.supplier_locations import (
    calculate_city_distance_fast,
    haversine_distance_km,
    nearest_location_distance,
    normalize_city_name,
    normalize_location_records,
    primary_city,
    serialize_locations,
    annotate_contacts_radius_status,
)


class SupplierLocationsTests(unittest.TestCase):
    def test_normalize_city_name(self) -> None:
        self.assertEqual(normalize_city_name("  Москва  "), "москва")
        self.assertEqual(normalize_city_name("Санкт-Петербург"), "санкт-петербург")
        self.assertEqual(normalize_city_name("г. Ростов-на-Дону"), "ростов-на-дону")
        self.assertEqual(normalize_city_name("Королёв"), "королев")

    def test_calculate_city_distance_fast(self) -> None:
        # Same city is 0
        self.assertEqual(calculate_city_distance_fast("Москва", "Москва"), 0)
        self.assertEqual(calculate_city_distance_fast("Ростов-на-Дону", "ростов-на-дону"), 0)

        # Rostov to Kamensk-Shakhtinsky (~140 km road distance)
        dist_kamensk = calculate_city_distance_fast("Ростов-на-Дону", "Каменск-Шахтинский")
        self.assertIsNotNone(dist_kamensk)
        assert dist_kamensk is not None
        self.assertTrue(110 < dist_kamensk < 180, f"Expected 110-180 km, got {dist_kamensk}")

        # Moscow to St. Petersburg (~700 km road distance)
        dist_spb = calculate_city_distance_fast("Москва", "Санкт-Петербург")
        self.assertIsNotNone(dist_spb)
        assert dist_spb is not None
        self.assertTrue(650 < dist_spb < 850, f"Expected 650-850 km, got {dist_spb}")

    def test_normalize_location_records(self) -> None:
        raw = [
            {"city": "Ростов-на-Дону", "confidence": 0.9},
            {"city": " ростов-на-дону ", "confidence": 0.3},
            {"city": "Шахты", "is_primary": False, "confidence": 0.8},
        ]
        norm = normalize_location_records(raw)
        self.assertEqual(len(norm), 2)
        self.assertEqual(norm[0]["city"], "Ростов-на-Дону")
        self.assertTrue(norm[0]["is_primary"])
        self.assertEqual(norm[1]["city"], "Шахты")
        self.assertEqual(primary_city(norm), "Ростов-на-Дону")

    def test_nearest_location_distance(self) -> None:
        locations = [
            {"city": "Москва", "confidence": 0.9},
            {"city": "Шахты", "confidence": 0.8},
            {"city": "Каменск-Шахтинский", "confidence": 0.95},
        ]
        nearest = nearest_location_distance(
            delivery_city="Каменск-Шахтинский",
            locations=locations,
        )
        self.assertEqual(nearest["city"], "Каменск-Шахтинский")
        self.assertEqual(nearest["distance_km"], 0)

        # Nearest to Rostov
        nearest_rnd = nearest_location_distance(
            delivery_city="Ростов-на-Дону",
            locations=locations,
        )
        self.assertEqual(nearest_rnd["city"], "Шахты")
        self.assertTrue(60 < nearest_rnd["distance_km"] < 100)

    def test_annotate_contacts_radius_status(self) -> None:
        contacts = [
            {"company_name": "ООО Близкий", "supplier_city": "Шахты"},
            {"company_name": "ООО Далекий", "supplier_city": "Москва"},
            {"company_name": "ООО Без Города", "supplier_city": ""},
        ]
        stats = annotate_contacts_radius_status(
            contacts,
            delivery_city="Каменск-Шахтинский",
            radius_km=150,
        )
        self.assertEqual(stats["within_radius"], 1)  # Шахты is ~75 km from Kamensk
        self.assertEqual(stats["outside_radius"], 1)  # Москва is ~900 km from Kamensk
        self.assertEqual(stats["unknown_distance"], 1)

        self.assertTrue(contacts[0]["supplier_distance_within_radius"])
        self.assertFalse(contacts[1]["supplier_distance_within_radius"])
        self.assertIsNone(contacts[2]["supplier_distance_within_radius"])

    def test_extract_city_from_text(self) -> None:
        from app.supplier_locations import extract_city_from_text
        self.assertEqual(extract_city_from_text("Адрес: г. Шахты, ул. Ленина, д. 5"), "шахты")
        self.assertEqual(extract_city_from_text("Наш склад расположен в городе Ростов-на-Дону"), "ростов-на-дону")
        self.assertEqual(extract_city_from_text("Производство: Каменск-Шахтинский"), "каменск-шахтинский")
        self.assertEqual(extract_city_from_text("Компания ООО ТрубПром поставки по РФ"), "")

    def test_get_cities_within_radius(self) -> None:
        from app.supplier_locations import get_cities_within_radius
        cities = get_cities_within_radius("Каменск-Шахтинский", 300)
        self.assertIn("каменск-шахтинский", cities)
        self.assertIn("шахты", cities)
        self.assertIn("ростов-на-дону", cities)
        self.assertNotIn("москва", cities)

    def test_russian_cities_database_count(self) -> None:
        from app.supplier_locations import CITY_COORDINATES
        self.assertGreater(len(CITY_COORDINATES), 1000, "Must contain over 1000 Russian cities")
        # Check smaller towns that were previously missing
        for town in ["зверево", "гуково", "белая калитва", "миллерово", "новошахтинск"]:
            self.assertIn(town, CITY_COORDINATES, f"Town {town} must be in database")

    def test_write_supplier_xlsx_with_delivery_city_and_radius(self) -> None:
        import tempfile
        from pathlib import Path
        import openpyxl
        from app.report_builder import write_supplier_xlsx

        rows = [
            {
                "company_name": "ООО ДонТрубПоставка",
                "site": "dontrub.ru",
                "phone": "+7 863 200-00-00",
                "email": "sale@dontrub.ru",
                "supplier_city": "Шахты",
                "supplier_distance_km": 78,
                "comments": "Трубы водоводные Ду 150 мм в наличии",
                "product_fit": "exact",
            }
        ]
        with tempfile.TemporaryDirectory() as td:
            out_file = Path(td) / "test_report.xlsx"
            write_supplier_xlsx(
                out_file,
                rows,
                title="Реконструкция водовода",
                delivery_city="Каменск-Шахтинский",
                geo_radius_km=300,
            )
            wb = openpyxl.load_workbook(out_file)
            ws = wb.active
            subtitle_val = str(ws.cell(row=2, column=1).value or "")
            summary_val = str(ws.cell(row=3, column=1).value or "")
            comment_val = str(ws.cell(row=6, column=5).value or "")

            self.assertIn("Каменск-Шахтинский", subtitle_val)
            self.assertIn("300", subtitle_val)
            self.assertIn("Логистический фильтр", summary_val)
            self.assertIn("Шахты", comment_val)
            self.assertIn("78 км", comment_val)


if __name__ == "__main__":
    unittest.main()
