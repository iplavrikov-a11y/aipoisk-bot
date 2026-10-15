import io
import tempfile
from pathlib import Path
import openpyxl
import pytest

from app.price_monitoring.cleaner import (
    build_commercial_search_queries,
    clean_nomenclature_name,
    extract_items_from_files_or_text,
    extract_package_hints,
    normalize_unit,
    parse_items_from_text,
)
from app.price_monitoring.excel_report import generate_price_monitoring_excel
from app.price_monitoring.models import (
    BatchMonitoringResult,
    ItemPriceResult,
    SupplierQuote,
)
from app.price_monitoring.pipeline import parse_input_excel


TEST_POSITIONS_12 = [
    ("Кислород газообразный технический", "м3"),
    ("Подшипник 13630", "шт"),
    ("Печь для бани Атмосфера М", "шт"),
    ("Электрод ОЗС-12 4мм", "кг"),
    ("Пропан сжиженный бал.50л", "кг"),
    ("Набивка ВАТИ-202 22х22мм", "кг"),
    ("Лист 10 ВСт3сп ГОСТ 19903", "т"),
    ("Круг шлифовальный ПП 80х20х20", "шт"),
    ("Ветошь хлопчатобумажная", "кг"),
    ("Смазка 1-13 ТУ 38.401-58-142", "кг"),
    ("Кабель ВВГнг-LS 3х2.5", "м"),
    ("Труба профильная 40х20х2", "м"),
]


def test_text_parsing_diverse_formats():
    """Verify parsing of various text formats including single lines, numbered lists, trailing model letters."""
    # 1. Single line with model letter 'М' (Task 903 case)
    res1 = parse_items_from_text("Печь для бани Атмосфера М")
    assert len(res1) == 1
    assert res1[0][0] == "Печь для бани Атмосфера М"
    assert res1[0][1] == "шт"

    # 2. Numbered list with various delimiters
    raw_text = """
    1. Кислород газообразный технический, м3
    2 - Подшипник 13630 - шт
    3. Электрод ОЗС-12 4мм (кг)
    4. Пропан сжиженный бал.50л, кг
    5) Набивка ВАТИ-202 22х22мм \t кг
    6. Лист 10 ВСт3сп ГОСТ 19903, т
    7. Круг шлифовальный ПП 80х20х20, шт
    8. Ветошь хлопчатобумажная, кг
    9. Смазка 1-13 ТУ 38.401-58-142, кг
    10. Кабель ВВГнг-LS 3х2.5, м
    11. Труба профильная 40х20х2, м
    """
    res2 = parse_items_from_text(raw_text)
    assert len(res2) == 11
    names = [r[0] for r in res2]
    units = [r[1] for r in res2]
    assert "Кислород газообразный технический" in names
    assert "Подшипник 13630" in names
    assert "Электрод ОЗС-12 4мм" in names
    assert "Лист 10 ВСт3сп ГОСТ 19903" in names
    assert "Кабель ВВГнг-LS 3х2.5" in names
    assert "м3" in units
    assert "кг" in units
    assert "т" in units
    assert "м" in units


def test_multi_file_and_text_extraction():
    """Verify extraction across 1, 2, 3 files and text input."""
    with tempfile.TemporaryDirectory() as tmpdir:
        tmp_path = Path(tmpdir)

        # File 1: Excel (.xlsx) with 4 items
        wb1 = openpyxl.Workbook()
        ws1 = wb1.active
        ws1.title = "Лист1"
        ws1.append(["№", "Номенклатура", "Ед. изм."])
        ws1.append([1, "Кислород газообразный технический", "м3"])
        ws1.append([2, "Подшипник 13630", "шт"])
        ws1.append([3, "Печь для бани Атмосфера М", "шт"])
        ws1.append([4, "Электрод ОЗС-12 4мм", "кг"])
        f1 = tmp_path / "spec1.xlsx"
        wb1.save(f1)

        # File 2: Excel (.xlsx) with 4 different items
        wb2 = openpyxl.Workbook()
        ws2 = wb2.active
        ws2.title = "Лист1"
        ws2.append(["Наименование", "Кол-во", "Единица"])
        ws2.append(["Пропан сжиженный бал.50л", 10, "кг"])
        ws2.append(["Набивка ВАТИ-202 22х22мм", 50, "кг"])
        ws2.append(["Лист 10 ВСт3сп ГОСТ 19903", 2, "т"])
        ws2.append(["Круг шлифовальный ПП 80х20х20", 100, "шт"])
        f2 = tmp_path / "spec2.xlsx"
        wb2.save(f2)

        # File 3: Text file (.txt) with remaining 4 items
        f3 = tmp_path / "spec3.txt"
        f3.write_text(
            "9. Ветошь хлопчатобумажная, кг\n"
            "10. Смазка 1-13 ТУ 38.401-58-142, кг\n"
            "11. Кабель ВВГнг-LS 3х2.5, м\n"
            "12. Труба профильная 40х20х2, м\n",
            encoding="utf-8",
        )

        # Scenario A: Single Excel file upload
        items_1 = extract_items_from_files_or_text([f1])
        assert len(items_1) == 4
        assert items_1[0][0] == "Кислород газообразный технический"

        # Scenario B: Two files upload (XLSX + XLSX)
        items_2 = extract_items_from_files_or_text([f1, f2])
        assert len(items_2) == 8

        # Scenario C: Three files upload (XLSX + XLSX + TXT)
        items_3 = extract_items_from_files_or_text([f1, f2, f3])
        assert len(items_3) == 12

        # Scenario D: Pure text input without any files (Job #903 reproduction)
        items_text = extract_items_from_files_or_text(
            [],
            context="Печь для бани Атмосфера М",
        )
        assert len(items_text) == 1
        assert items_text[0] == ("Печь для бани Атмосфера М", "шт")

        # Scenario E: Excel synthesis from items for openpyxl safety
        wb_synth = openpyxl.Workbook()
        ws_synth = wb_synth.active
        ws_synth.append(["Номенклатура", "Ед. изм."])
        for nom, unit in items_3:
            ws_synth.append([nom, unit])
        buf = io.BytesIO()
        wb_synth.save(buf)
        buf.seek(0)
        re_parsed = parse_input_excel(buf.getvalue())
        assert len(re_parsed) == 12


def test_12_positions_normalization_and_hints():
    """Verify clean name, normalized units, and package hints for all 12 positions."""
    for raw_nom, expected_unit in TEST_POSITIONS_12:
        clean_name = clean_nomenclature_name(raw_nom)
        assert clean_name, f"Clean name should not be empty for {raw_nom}"
        norm_u = normalize_unit(expected_unit)
        assert norm_u == expected_unit or (norm_u in {"шт", "м3", "кг", "т", "м"})

        queries = build_commercial_search_queries(clean_name, norm_u)
        assert queries, f"Search queries must be generated for {clean_name}"
        assert len(queries[0]) >= 3

        hints = extract_package_hints(clean_name, norm_u)
        assert isinstance(hints, dict)

    # Specific hints checks
    propan_hints = extract_package_hints("Пропан сжиженный бал.50л", "кг")
    assert propan_hints.get("cylinder_volume_l") == 50.0
    assert propan_hints.get("gas_type") == "propane"

    electrodes_hints = extract_package_hints("Электрод ОЗС-12 4мм пачка 5 кг", "кг")
    assert electrodes_hints.get("pack_weight_kg") == 5.0


def test_12_positions_e2e_pipeline_and_report():
    """Run simulated price monitoring pipeline for all 12 positions and verify Excel output."""
    item_results = []

    for i, (nom, unit) in enumerate(TEST_POSITIONS_12, start=1):
        base_price = (i * 125.50) + 50.0
        # 3 diverse supplier quotes per item
        q1 = SupplierQuote(
            company_name=f"ООО ПромПоставка {i}",
            price_rub=round(base_price, 2),
            raw_price_rub=round(base_price, 2),
            unit=unit,
            site_unit=unit,
            inn="7701234567",
            kpp="770101001",
            contacts="+7 (495) 123-45-67, sales@prompostavka.ru",
            url=f"https://prompostavka{i}.ru/catalog/item",
        )
        q2 = SupplierQuote(
            company_name=f"ЗАО ЗаводСнаб {i}",
            price_rub=round(base_price * 1.15, 2),
            raw_price_rub=round(base_price * 1.15, 2),
            unit=unit,
            site_unit=unit,
            inn="7801987654",
            kpp="780101001",
            contacts="+7 (812) 987-65-43, zakaz@zavodsnab.ru",
            url=f"https://zavodsnab{i}.ru/product",
        )
        q3 = SupplierQuote(
            company_name=f"ТД ОптТорг {i}",
            price_rub=round(base_price * 1.30, 2),
            raw_price_rub=round(base_price * 1.30, 2),
            unit=unit,
            site_unit=unit,
            inn="5001112233",
            kpp="500101001",
            contacts="+7 (499) 555-44-33, info@opttorg.ru",
            url=f"https://opttorg{i}.ru/item",
        )

        item_res = ItemPriceResult(
            row_number=i,
            raw_name=nom,
            clean_name=nom,
            unit=unit,
            quotes=[q1, q2, q3],
            status="found",
        )
        item_results.append(item_res)

    with tempfile.TemporaryDirectory() as tmpdir:
        output_xlsx = Path(tmpdir) / "output_test.xlsx"
        generate_price_monitoring_excel(item_results, str(output_xlsx))

        assert output_xlsx.exists()
        assert output_xlsx.stat().st_size > 2000

        # Verify generated workbook structure
        wb = openpyxl.load_workbook(output_xlsx)
        sheet_names = wb.sheetnames
        assert "Пример отчета" in sheet_names
        assert "Детализированный реестр" in sheet_names

        ws_summary = wb["Пример отчета"]
        # Rows 5+ contain the items
        cell_values = [str(ws_summary.cell(row=r, column=1).value or "") for r in range(5, 5 + len(TEST_POSITIONS_12))]
        for nom, _ in TEST_POSITIONS_12:
            assert any(nom in val for val in cell_values), f"Position {nom} must be in Summary sheet"

        ws_details = wb["Детализированный реестр"]
        # All 12 items present in rows 2 to 13
        detail_items = [ws_details.cell(row=r, column=2).value for r in range(2, 2 + len(TEST_POSITIONS_12))]
        assert len(detail_items) == 12
        for nom, _ in TEST_POSITIONS_12:
            assert nom in detail_items, f"Position {nom} must be in Details sheet"

        # Check suppliers columns populated in row 2
        assert ws_details.cell(row=2, column=6).value == "ООО ПромПоставка 1"
        assert ws_details.cell(row=2, column=15).value == "ЗАО ЗаводСнаб 1"
        assert ws_details.cell(row=2, column=24).value == "ТД ОптТорг 1"


def test_city_aware_price_monitoring():
    """Verify city-aware query generation and Excel report annotation for price monitoring."""
    # 1. Query generation with city
    queries = build_commercial_search_queries("Печь для бани Атмосфера М", "шт", city="Вологда")
    assert any("Вологда" in q for q in queries)
    assert any("купить цена" in q for q in queries)

    # 2. Excel report generation with delivery_city
    item_res = ItemPriceResult(
        row_number=1,
        raw_name="Печь для бани Атмосфера М",
        clean_name="Печь для бани Атмосфера М",
        unit="шт",
        quotes=[
            SupplierQuote(
                company_name="ООО Печи Вологда",
                price_rub=45000.0,
                raw_price_rub=45000.0,
                unit="шт",
                site_unit="шт",
                inn="3525123456",
                contacts="+7 (8172) 12-34-56, sales@vologda-pechi.ru",
                url="https://vologda-pechi.ru/atmosfera-m",
            )
        ],
        status="found",
    )

    with tempfile.TemporaryDirectory() as tmpdir:
        output_xlsx = Path(tmpdir) / "output_geo_test.xlsx"
        generate_price_monitoring_excel(
            [item_res],
            str(output_xlsx),
            max_suppliers=3,
            delivery_city="Вологда",
            geo_radius_km=110,
        )
        assert output_xlsx.exists()
        wb = openpyxl.load_workbook(output_xlsx)
        ws = wb["Пример отчета"]
        assert ws["A1"].comment is not None
        assert "Вологда" in ws["A1"].comment.text
        assert "110" in ws["A1"].comment.text


