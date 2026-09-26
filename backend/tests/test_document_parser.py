from __future__ import annotations

import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

import app.document_parser as document_parser


class DocumentParserTests(unittest.TestCase):
    def test_doc_extraction_falls_back_to_libreoffice_when_antiword_is_broken(self) -> None:
        def fake_which(name: str) -> str | None:
            if name == "antiword":
                return "/usr/bin/antiword"
            return None

        with (
            patch.object(document_parser.shutil, "which", side_effect=fake_which),
            patch.object(document_parser.subprocess, "run", side_effect=OSError("Exec format error")),
            patch.object(document_parser, "_extract_via_libreoffice", return_value="Текст технического задания") as libreoffice,
        ):
            text = document_parser._extract_doc(Path("broken.doc"))

        self.assertEqual(text, "Текст технического задания")
        libreoffice.assert_called_once()

    def test_zip_archive_extracts_member_documents_as_one_context(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            archive_path = Path(tmp) / "parts.zip"
            with zipfile.ZipFile(archive_path, "w") as archive:
                archive.writestr("part-1.txt", "Первая часть ТЗ: насос ЦНС 60-330")
                archive.writestr("nested/part-2.txt", "Вторая часть ТЗ: количество 3 шт.")

            text, status = document_parser.extract_text(archive_path)

        self.assertEqual(status, "archive_ok")
        self.assertIn("=== ARCHIVE FILE:", text)
        self.assertIn("Первая часть ТЗ", text)
        self.assertIn("Вторая часть ТЗ", text)

    def test_docx_extension_archive_falls_back_to_archive_extraction(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            archive_path = Path(tmp) / "documentation.docx"
            with zipfile.ZipFile(archive_path, "w") as archive:
                archive.writestr("Техническое задание.txt", "Огнезащитный материал, количество 2106 кг")

            text, status = document_parser.extract_text(archive_path)

        self.assertEqual(status, "docx_archive_archive_ok")
        self.assertIn("Огнезащитный материал", text)

    def test_short_docx_extraction_uses_libreoffice_text_fallback(self) -> None:
        with (
            patch.object(document_parser, "_extract_docx", return_value="\n=== TABLE 1 ==="),
            patch.object(document_parser, "_extract_via_libreoffice", return_value="Передвижная экологическая лаборатория") as libreoffice,
        ):
            text, status = document_parser.extract_text(Path("short.docx"))

        self.assertEqual(status, "docx_libreoffice_ok")
        self.assertEqual(text, "Передвижная экологическая лаборатория")
        libreoffice.assert_called_once()

    def test_corrupted_docx_relationships_fall_back_to_libreoffice_or_xml(self) -> None:
        with (
            patch.object(document_parser, "_extract_docx", side_effect=KeyError("There is no item named 'word/NULL' in the archive")),
            patch.object(document_parser, "_extract_via_libreoffice", return_value="ТЗ на модульные дома"),
        ):
            text, status = document_parser.extract_text(Path("corrupted.docx"))

        self.assertEqual(status, "docx_libreoffice_ok")
        self.assertEqual(text, "ТЗ на модульные дома")

    def test_corrupted_docx_falls_back_to_direct_xml_when_libreoffice_empty(self) -> None:
        with (
            patch.object(document_parser, "_extract_docx", side_effect=KeyError("There is no item named 'word/NULL' in the archive")),
            patch.object(document_parser, "_extract_via_libreoffice", return_value=""),
            patch.object(document_parser, "_extract_docx_xml", return_value="Прямой текст из XML документа"),
        ):
            text, status = document_parser.extract_text(Path("corrupted.docx"))

        self.assertEqual(status, "docx_xml_ok")
        self.assertEqual(text, "Прямой текст из XML документа")

    def test_broken_pandoc_does_not_block_libreoffice_fallback(self) -> None:
        with (
            patch.object(document_parser.shutil, "which", return_value="/usr/bin/pandoc"),
            patch.object(document_parser.subprocess, "run", side_effect=OSError("Exec format error")),
            patch.object(document_parser, "_extract_via_libreoffice", return_value="Поставка устройств пробоотборных") as libreoffice,
        ):
            text, status = document_parser.extract_text(Path("technical.odt"))

        self.assertEqual(status, "ok")
        self.assertEqual(text, "Поставка устройств пробоотборных")
        libreoffice.assert_called_once()



    def test_sanitize_filename_preserves_extension_on_long_cyrillic_names(self) -> None:
        long_name = "??????-???????? ????? ?????? ???????????? ?????????? ??? PUMP TYPE Z 12-205 JMW ? ?????? ????????? PUMP TYPE Z 12-125 JMW - ???????????.docx"
        sanitized = document_parser.sanitize_filename(long_name)
        self.assertTrue(sanitized.endswith(".docx"))
        self.assertLessEqual(len(sanitized.encode("utf-8")), 180)

    def test_extract_text_auto_detects_docx_without_extension(self) -> None:
        from docx import Document
        with tempfile.TemporaryDirectory() as tmp:
            no_ext_file = Path(tmp) / "docx_file_no_ext"
            doc = Document()
            doc.add_paragraph("??????-???????? ????? ?????? ???????????? ?????????? ??? PUMP TYPE Z 12-205 JMW")
            doc.save(str(no_ext_file))
            
            text, status = document_parser.extract_text(no_ext_file)
            self.assertTrue(status in ("ok", "docx_libreoffice_ok"))
            self.assertIn("PUMP TYPE Z 12-205 JMW", text)

    def test_docx_with_scanned_images_triggers_media_ocr(self) -> None:
        with (
            patch.object(document_parser, "_extract_docx", return_value=""),
            patch.object(document_parser, "_extract_via_libreoffice", return_value="\ufeff\n\n\n\n\n"),
            patch.object(document_parser, "_extract_docx_xml", return_value=""),
            patch.object(document_parser, "_extract_docx_media_ocr", return_value="Технические требования на АКБ"),
        ):
            text, status = document_parser.extract_text(Path("scanned_tz.docx"))

        self.assertEqual(status, "docx_ocr_ok")
        self.assertIn("Технические требования на АКБ", text)

    def test_doc_with_short_text_triggers_office_pdf_ocr(self) -> None:
        with (
            patch.object(document_parser, "_extract_doc", return_value=""),
            patch.object(document_parser, "_extract_office_via_pdf_ocr", return_value="Технические условия ГОСТ 12345"),
        ):
            text, status = document_parser.extract_text(Path("scanned_tz.doc"))

        self.assertEqual(status, "doc_ocr_ok")
        self.assertIn("Технические условия ГОСТ 12345", text)

    def test_docx_with_nested_tables_extracts_all_nested_content(self) -> None:
        from docx import Document

        with tempfile.TemporaryDirectory() as tmp:
            doc_file = Path(tmp) / "nested.docx"
            doc = Document()
            doc.add_paragraph("Техническое задание на поставку")
            outer_table = doc.add_table(rows=1, cols=1)
            outer_cell = outer_table.cell(0, 0)
            nested_table = outer_cell.add_table(rows=2, cols=2)
            nested_table.cell(0, 0).text = "Прибор"
            nested_table.cell(0, 1).text = "Количество"
            nested_table.cell(1, 0).text = "Гауссметр цифровой"
            nested_table.cell(1, 1).text = "2 шт"
            doc.save(str(doc_file))

            text, status = document_parser.extract_text(doc_file)
            self.assertEqual(status, "ok")
            self.assertIn("Гауссметр цифровой", text)
            self.assertIn("2 шт", text)

    def test_docx_with_empty_table_marker_triggers_fallback(self) -> None:
        with (
            patch.object(document_parser, "_extract_docx", return_value="Описание объекта закупки\n=== TABLE 1 ==="),
            patch.object(document_parser, "_extract_via_libreoffice", return_value="Описание объекта закупки\nПолная спецификация: Тесламетр 1 шт") as libreoffice,
        ):
            text, status = document_parser.extract_text(Path("empty_table.docx"))

        self.assertEqual(status, "docx_libreoffice_ok")
        self.assertIn("Тесламетр", text)
        libreoffice.assert_called_once()

    def test_docx_xml_extracts_nested_table_rows(self) -> None:
        from docx import Document

        with tempfile.TemporaryDirectory() as tmp:
            doc_file = Path(tmp) / "nested_xml.docx"
            doc = Document()
            doc.add_paragraph("Вводный текст")
            outer_table = doc.add_table(rows=1, cols=1)
            outer_cell = outer_table.cell(0, 0)
            nested_table = outer_cell.add_table(rows=2, cols=2)
            nested_table.cell(0, 0).text = "Оборудование"
            nested_table.cell(0, 1).text = "Параметр"
            nested_table.cell(1, 0).text = "Анализатор спектра"
            nested_table.cell(1, 1).text = "6.5 ГГц"
            doc.save(str(doc_file))

            xml_text = document_parser._extract_docx_xml(doc_file)
            self.assertIn("Анализатор спектра", xml_text)
            self.assertIn("6.5 ГГц", xml_text)

    def test_organize_procurement_documents_single_file(self) -> None:
        items = [("file1.docx", "Некоторый текст ТЗ")]
        reordered, subject = document_parser.organize_procurement_documents(items)
        self.assertEqual(reordered, items)
        self.assertEqual(subject, "")

    def test_organize_procurement_documents_ai_ordering(self) -> None:
        import json
        from unittest.mock import patch, AsyncMock
        items = [
            ("Проект_ГК.docx", "Проект государственного контракта на поставку оборудования..."),
            ("ТЗ_Спецификация.docx", "Техническое задание на поставку матрасов..."),
            ("НМЦК.xlsx", "Расчет начальной максимальной цены..."),
        ]
        mock_resp = json.dumps({
            "procurement_subject": "Поставка матрасов",
            "documents": [
                {"index": 0, "category": "contract_draft", "priority": 4},
                {"index": 1, "category": "technical_spec", "priority": 1},
                {"index": 2, "category": "pricing_nmck", "priority": 3},
            ]
        })
        with patch("app.exact_product.llm_bridge.call_llm", new=AsyncMock(return_value=mock_resp)):
            reordered, subject = document_parser.organize_procurement_documents(items)
            self.assertEqual(subject, "Поставка матрасов")
            self.assertEqual(reordered[0][0], "ТЗ_Спецификация.docx")
            self.assertEqual(reordered[1][0], "НМЦК.xlsx")
            self.assertEqual(reordered[2][0], "Проект_ГК.docx")


if __name__ == "__main__":
    unittest.main()



