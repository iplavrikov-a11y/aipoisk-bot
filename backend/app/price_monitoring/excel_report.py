from __future__ import annotations

import logging
import re
from collections import Counter
from typing import List, Optional
from urllib.parse import urlparse

from openpyxl import Workbook
from openpyxl.cell.rich_text import TextBlock, CellRichText
from openpyxl.cell.text import InlineFont
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .cleaner import extract_generic_nomenclature_codes, is_closed_factory_drawing
from .models import ItemPriceResult, SupplierQuote

logger = logging.getLogger(__name__)


def _source_position_note(item: ItemPriceResult) -> str:
    note = [f"Позиция исходного перечня: {item.source_position_no or item.row_number}"]
    for label, value in (("ID", item.position_id), ("Файл", item.source_file), ("Лист", item.source_sheet),
                         ("Строка", item.source_row), ("Количество", item.quantity)):
        if value:
            note.append(f"{label}: {value}")
    note.extend(f"{requirement.get('param_name')}: {requirement.get('tz_requirement')}" for requirement in item.requirements)
    return "\n".join(note)


def _missing_price_status(item: ItemPriceResult) -> str:
    if item.status == "error":
        return "Ошибка обработки позиции; требуется повторный поиск"
    if any(q.price_rub > 0 and q.is_outlier for q in item.quotes):
        return "Цены не образуют сопоставимую группу"
    reason = getattr(item, "not_found_reason", "")
    if reason in {"budget", "budget_exhausted"}:
        return "Цена не найдена: исчерпан лимит поиска"
    if reason in {"network", "network_error"}:
        return "Цена не подтверждена: источники недоступны"
    if reason == "all_rejected":
        return "Цена не подтверждена: предложения не прошли проверку"
    codes = extract_generic_nomenclature_codes(item.clean_name)
    if any(is_closed_factory_drawing(c, item.clean_name) for c in codes):
        return "По точному обозначению цена в этом поиске не найдена"
    return "Предложения с ценой в этом поиске не найдены"


def _price_sources_status(quotes: list[SupplierQuote]) -> str:
    count = len(quotes)
    if count == 1:
        return "Найдена цена: 1 источник; требуется проверка"
    if count == 2:
        return "Найдены цены: 2 источника; мало данных"
    return f"Найдены цены: {count} источников"


def _quote_evidence_note(q: SupplierQuote) -> str:
    level = {
        "page_content": "Цена найдена в тексте страницы",
        "search_snippet": "Источник цены: поисковый сниппет; требуется проверка страницы",
    }.get(q.evidence_level, "Способ подтверждения источника не указан")
    return level + (f"; получено: {q.retrieved_at}" if q.retrieved_at else "")


def _quote_geo_note(q: SupplierQuote) -> str:
    status = q.supplier_distance_radius_status
    source_note = "; расчёт между городами по контактному адресу" if q.supplier_location_source == "contact_address_city" else ""
    if status == "within_radius":
        return f"География: {q.supplier_city or 'город указан'}; {q.supplier_distance_km} км, в радиусе{source_note}"
    if status == "outside_radius":
        return f"География: {q.supplier_city or 'город указан'}; {q.supplier_distance_km} км, вне радиуса (добор по РФ){source_note}"
    if status == "unknown_distance":
        return "География: адрес поставщика не подтверждён; расстояние не определено"
    return ""


def format_short_url(url: str) -> str:
    """Форматирует компактную кликабельную ссылку без визуального захламления ячейки."""
    if not url:
        return ""
    try:
        parsed = urlparse(url)
        domain = parsed.netloc.lower()
        if domain.startswith("www."):
            domain = domain[4:]
        path = parsed.path.strip("/")
        if not path:
            return domain
        parts = [p for p in path.split("/") if p]
        if len(parts) == 1:
            slug = parts[0]
            if len(slug) > 24:
                slug = slug[:22] + ".."
            return f"{domain}/{slug}"
        else:
            last_part = parts[-1]
            if len(last_part) > 24:
                last_part = last_part[:22] + ".."
            return f"{domain}/.../{last_part}"
    except Exception:
        return url[:35]


def _extract_phones_and_emails(contacts: Optional[str]) -> tuple[str, str]:
    """
    Извлекает исключительно телефоны и email поставщика, исключая любые физические адреса и города.
    Возвращает (phones_str, emails_str).
    """
    if not contacts:
        return "", ""

    raw_parts = [p.strip() for p in re.split(r'[,;\n]+', contacts) if p.strip()]
    phones: list[str] = []
    emails: list[str] = []
    seen: set[str] = set()

    for p in raw_parts:
        clean_p = re.sub(r'\s+', ' ', p).strip()
        key = clean_p.lower()
        if key in seen:
            continue
        seen.add(key)

        if "@" in clean_p and "." in clean_p:
            clean_email = re.sub(r'^(?:email[\.:]?|e-mail[\.:]?|эл\.?\s*почта[\.:]?)\s*', '', clean_p, flags=re.IGNORECASE).strip()
            # Исключаем технические и сервисные адреса (sentry, noreply, mailer-daemon, abuse, postmaster)
            em_low = clean_email.lower()
            if any(tech in em_low for tech in ("sentry", "noreply", "no-reply", "mailer-daemon", "postmaster", "abuse")):
                continue
            emails.append(clean_email)
        elif any(c.isdigit() for c in clean_p) and (len(re.sub(r'\D', '', clean_p)) in (10, 11) or "800" in clean_p or "+7" in clean_p or "tel" in clean_p.lower()):
            clean_phone = re.sub(r'^(?:тел[\.:]?|tel[\.:]?|телефон[\.:]?)\s*', '', clean_p, flags=re.IGNORECASE).strip()
            phones.append(clean_phone)

    phones_str = ("Тел: " + ", ".join(phones[:2])) if phones else ""
    emails_str = ("Email: " + ", ".join(emails[:2])) if emails else ""
    return phones_str, emails_str


def _format_contacts_multiline(contacts: Optional[str]) -> str:
    """
    Форматирует контакты для шапки таблицы (без физического адреса).
    """
    phones_str, emails_str = _extract_phones_and_emails(contacts)
    parts = [p for p in (phones_str, emails_str) if p]
    return "\n".join(parts) if parts else "контакты на сайте"


def _format_contacts_compact(contacts: Optional[str]) -> str:
    """
    Форматирует контакты компактно через запятую (только телефон и email, без адреса).
    """
    phones_str, emails_str = _extract_phones_and_emails(contacts)
    parts = [p for p in (phones_str, emails_str) if p]
    return ", ".join(parts) if parts else "контакты на сайте"


def _build_cell_comment(q: SupplierQuote, item_unit: str) -> Comment:
    """Формирует всплывающую подсказку (комментарий Excel) с реквизитами поставщика и пересчетом."""
    lines = [
        f"Поставщик: {q.company_name}",
        f"ИНН: {q.inn or 'не указан'}" + (f" / КПП: {q.kpp}" if q.kpp else ""),
        f"Контакты: {q.contacts or 'см. на сайте'}",
    ]
    if q.raw_price_rub > 0 and q.site_unit:
        lines.append(f"Цена на сайте: {q.raw_price_rub:,.2f} ₽ ({q.site_unit})")
    lines.append(f"Цена в отчете: {q.price_rub:,.2f} ₽ за {item_unit}")
    if q.conversion_comment:
        lines.append(f"Пересчет: {q.conversion_comment}")
    if q.comment:
        lines.append(f"Примечание: {q.comment}")
    geo_note = _quote_geo_note(q)
    if geo_note:
        lines.append(geo_note)
    lines.append(_quote_evidence_note(q))
    if q.source_excerpt:
        lines.append(f"Фрагмент источника: {q.source_excerpt}")
    lines.append(f"Полная ссылка: {q.url}")

    text = "\n".join(lines)
    comm = Comment(text, "TenderLex")
    comm.width = 380
    comm.height = 240
    return comm


def generate_price_monitoring_excel(
    items: List[ItemPriceResult],
    output_path: str,
    max_suppliers: int = 5,
    max_variation_pct: float = 33.0,
    delivery_city: str = "",
    geo_radius_km: Optional[int] = None,
) -> str:
    """
    Генерирует Excel-файл со структурой и стилями, в точности соответствующими шаблону заказчика.
    Лист 1: 'Пример отчета' (сводная матрица с реквизитами в шапке, компактными ссылками и всплывающими подсказками).
    Лист 2: 'Детализированный реестр' (полный перечень для 1С/закупщика с ИНН, КПП, контактами и формулами пересчета).
    """
    wb = Workbook()

    # --- ЛИСТ 1: 'Пример отчета' по шаблону заказчика ---
    ws = wb.active
    ws.title = "Пример отчета"

    # Стили
    font_calibri_bold = Font(name="Calibri", size=11, bold=True)
    font_calibri_regular = Font(name="Calibri", size=11)
    font_header_bold = Font(name="Calibri", size=10, bold=True)
    font_header_regular = Font(name="Calibri", size=10)
    font_link = Font(name="Calibri", size=10, color="0563C1", underline="single")
    font_price_bold = Font(name="Calibri", size=11, bold=True, color="166534")
    font_price_outlier = Font(name="Calibri", size=11, bold=True, color="B45309")
    font_rich_regular = InlineFont(rFont="Calibri", sz=10, color="1E293B")
    font_rich_link = InlineFont(rFont="Calibri", sz=10, color="0563C1", u="single")

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    header_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    sup_header_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_right = Alignment(horizontal="right", vertical="center")

    # Колонки A и B
    ws.merge_cells("A1:A4")
    cell_a = ws["A1"]
    cell_a.value = "Номенклатура"
    cell_a.font = font_calibri_bold
    cell_a.alignment = align_center
    if delivery_city:
        clean_city = str(delivery_city).strip()
        rad_info = f" (в радиусе до {geo_radius_km} км)" if geo_radius_km else ""
        comm_loc = Comment(f"Локация / поставка: г. {clean_city}{rad_info}", "TenderLex")
        comm_loc.width = 280
        comm_loc.height = 60
        cell_a.comment = comm_loc

    ws.merge_cells("B1:B4")
    cell_b = ws["B1"]
    cell_b.value = "Ед. изм."
    cell_b.font = font_calibri_bold
    cell_b.alignment = align_center

    # Определяем основных поставщиков по массиву позиций
    supplier_freq: Counter[str] = Counter()
    supplier_best_meta: dict[str, SupplierQuote] = {}
    for item in items:
        for q in item.quotes:
            if q.company_name and q.price_rub > 0:
                supplier_freq[q.company_name] += 1
                if q.company_name not in supplier_best_meta or (q.inn and not supplier_best_meta[q.company_name].inn):
                    supplier_best_meta[q.company_name] = q

    top_suppliers = [sup_name for sup_name, _ in supplier_freq.most_common(max_suppliers)]

    # Поставщик считается доминирующим ТОЛЬКО если закрывает почти всю закупку (>= 80% позиций)
    dominant_threshold = max(1, int(len(items) * 0.8)) if len(items) > 1 else 1

    # Шапка поставщиков 1..max_suppliers (строки 1-4)
    for s_idx in range(1, max_suppliers + 1):
        col_start = 3 + (s_idx - 1) * 2  # C=3, E=5, G=7, I=9, K=11
        col_end = col_start + 1          # D=4, F=6, H=8, J=10, L=12
        col_start_let = get_column_letter(col_start)
        col_end_let = get_column_letter(col_end)

        # Выбираем метаданные поставщика для шапки столбца
        sup_quote: Optional[SupplierQuote] = None
        if s_idx <= len(top_suppliers):
            sup_quote = supplier_best_meta.get(top_suppliers[s_idx - 1])
        elif items and any(len(it.quotes) >= s_idx for it in items):
            for it in items:
                if len(it.quotes) >= s_idx and it.quotes[s_idx - 1].price_rub > 0:
                    sup_quote = it.quotes[s_idx - 1]
                    break

        is_dominant = False
        if sup_quote and supplier_freq:
            item_coverage = supplier_freq.get(sup_quote.company_name, 0)
            is_dominant = item_coverage >= dominant_threshold

        if sup_quote and is_dominant:
            c_name = sup_quote.company_name.strip()
            if not any(c_name.startswith(p) for p in ("ООО", "АО", "ЗАО", "ПАО", "ИП", "Общество", "Акционерное")):
                c_display = f'ООО "{c_name}"'
            else:
                c_display = c_name

            sup_name_text = f"Поставщик {s_idx} ({c_display})"
            inn_parts = []
            if sup_quote.inn:
                inn_parts.append(f"ИНН {sup_quote.inn}")
            if sup_quote.kpp:
                inn_parts.append(f"КПП {sup_quote.kpp}")
            inn_kpp_text = " / ".join(inn_parts) if inn_parts else "ИНН: уточняется"
            contacts_text = _format_contacts_multiline(sup_quote.contacts)
        else:
            sup_name_text = f"Поставщик {s_idx} ( наименование )"
            inn_kpp_text = "ИНН/КПП"
            contacts_text = "контакты"

        # Строка 1: Поставщик N (наименование)
        ws.merge_cells(f"{col_start_let}1:{col_end_let}1")
        cell_sup = ws[f"{col_start_let}1"]
        cell_sup.value = sup_name_text
        cell_sup.font = font_header_bold
        cell_sup.alignment = align_center

        # Строка 2: ИНН/КПП
        ws.merge_cells(f"{col_start_let}2:{col_end_let}2")
        cell_inn = ws[f"{col_start_let}2"]
        cell_inn.value = inn_kpp_text
        cell_inn.font = font_header_regular
        cell_inn.alignment = align_center

        # Строка 3: контакты
        ws.merge_cells(f"{col_start_let}3:{col_end_let}3")
        cell_cont = ws[f"{col_start_let}3"]
        cell_cont.value = contacts_text
        cell_cont.font = font_header_regular
        cell_cont.alignment = align_center

        # Строка 4: Цена руб./ед.
        cell_p = ws[f"{col_start_let}4"]
        cell_p.value = "Цена руб./ед."
        cell_p.font = font_header_bold
        cell_p.alignment = align_center

        # Строка 4: источник (ссылка)
        cell_src = ws[f"{col_end_let}4"]
        cell_src.value = "источник (ссылка)"
        cell_src.font = font_header_regular
        cell_src.alignment = align_center

    # Сводные колонки: Мин. цена, Средняя цена и Коэффициент вариации
    col_min_idx = 3 + max_suppliers * 2
    col_avg_idx = col_min_idx + 1
    col_var_idx = col_avg_idx + 1
    col_min_let = get_column_letter(col_min_idx)
    col_avg_let = get_column_letter(col_avg_idx)
    col_var_let = get_column_letter(col_var_idx)

    ws.merge_cells(f"{col_min_let}1:{col_min_let}4")
    cell_min_hdr = ws[f"{col_min_let}1"]
    cell_min_hdr.value = "Мин. цена, руб."
    cell_min_hdr.font = font_calibri_bold
    cell_min_hdr.alignment = align_center

    ws.merge_cells(f"{col_avg_let}1:{col_avg_let}4")
    cell_avg_hdr = ws[f"{col_avg_let}1"]
    cell_avg_hdr.value = "Средняя цена, руб."
    cell_avg_hdr.font = font_calibri_bold
    cell_avg_hdr.alignment = align_center

    ws.merge_cells(f"{col_var_let}1:{col_var_let}4")
    cell_var_hdr = ws[f"{col_var_let}1"]
    cell_var_hdr.value = "Коэффициент вариации"
    cell_var_hdr.font = font_calibri_bold
    cell_var_hdr.alignment = align_center

    # Применяем границы и фон к заголовкам
    max_col = col_var_idx
    for r in range(1, 5):
        for c in range(1, max_col + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = thin_border
            if r == 4 or c <= 2 or c >= col_min_idx:
                cell.fill = header_fill
            else:
                cell.fill = sup_header_fill

    # Список колонок с ценами для формул MIN и AVERAGE
    price_col_letters = [get_column_letter(3 + (s - 1) * 2) for s in range(1, max_suppliers + 1)]

    # Заполняем строки данными
    current_row = 5
    for item in items:
        # A: Номенклатура
        cell_nom = ws.cell(row=current_row, column=1, value=item.raw_name)
        cell_nom.comment = Comment(_source_position_note(item), "TenderLex")
        cell_nom.font = font_calibri_regular
        cell_nom.alignment = align_left
        cell_nom.border = thin_border

        # B: Ед. изм.
        cell_unit = ws.cell(row=current_row, column=2, value=item.unit)
        cell_unit.font = font_calibri_regular
        cell_unit.alignment = align_center
        cell_unit.border = thin_border

        display_quotes = [q for q in item.quotes if q.price_rub > 0 and not q.is_outlier]

        if not display_quotes:
            missing_label = _missing_price_status(item)

            cell_p1 = ws.cell(row=current_row, column=3, value="-")
            cell_p1.border = thin_border
            cell_p1.alignment = align_center
            cell_p1.font = font_calibri_regular

            cell_l1 = ws.cell(row=current_row, column=4, value=missing_label)
            cell_l1.border = thin_border
            cell_l1.font = Font(name="Calibri", size=10, italic=True, color="64748B")
            cell_l1.alignment = align_left

            for s_idx in range(2, max_suppliers + 1):
                col_start = 3 + (s_idx - 1) * 2
                col_end = col_start + 1
                c_p = ws.cell(row=current_row, column=col_start, value="")
                c_p.border = thin_border
                c_l = ws.cell(row=current_row, column=col_end, value="")
                c_l.border = thin_border

            cell_min = ws.cell(row=current_row, column=col_min_idx, value="")
            cell_min.border = thin_border
            cell_avg = ws.cell(row=current_row, column=col_avg_idx, value="")
            cell_avg.border = thin_border
            cell_var = ws.cell(row=current_row, column=col_var_idx, value="")
            cell_var.border = thin_border

            current_row += 1
            continue

        # Заполняем котировки поставщиков (СТРОГО согласованный кластер цен)
        for s_idx in range(1, max_suppliers + 1):
            col_start = 3 + (s_idx - 1) * 2
            col_end = col_start + 1

            cell_price = ws.cell(row=current_row, column=col_start)
            cell_price.border = thin_border
            cell_price.alignment = align_right

            cell_link = ws.cell(row=current_row, column=col_end)
            cell_link.border = thin_border
            cell_link.alignment = align_left

            if s_idx <= len(display_quotes):
                q = display_quotes[s_idx - 1]
                cell_price.value = q.price_rub
                cell_price.number_format = "#,##0.00"
                cell_price.font = font_price_bold
                if q.url.startswith("http"):
                    cell_price.hyperlink = q.url

                c_name = q.company_name.strip()
                if not any(c_name.startswith(p) for p in ("ООО", "АО", "ЗАО", "ПАО", "ИП", "Общество", "Акционерное")):
                    clean_c = c_name.strip(" '\"")
                    if clean_c != format_short_url(q.url) and "." not in clean_c:
                        c_display = f'ООО "{clean_c}"'
                    else:
                        c_display = clean_c
                else:
                    c_display = c_name

                # Извлекаем контакты строго без физических адресов
                phone_str, email_str = _extract_phones_and_emails(q.contacts)

                # Формируем RichText: компания, ИНН и телефон — обычным текстом, email и ссылка на источник — как ссылки
                runs: list[TextBlock] = []
                line1 = c_display + (f" | ИНН {q.inn}" if q.inn else "")
                header_text = line1 + "\n"
                if q.supplier_city:
                    dist_info = f" (~{q.supplier_distance_km} км)" if q.supplier_distance_km is not None else ""
                    header_text += f"📍 {q.supplier_city}{dist_info}\n"
                if phone_str:
                    header_text += phone_str + "\n"
                runs.append(TextBlock(font_rich_regular, header_text))

                if email_str:
                    runs.append(TextBlock(font_rich_link, email_str + "\n"))

                short_link = format_short_url(q.url)
                runs.append(TextBlock(font_rich_link, short_link))

                cell_link.value = CellRichText(*runs)
                if q.url.startswith("http"):
                    cell_link.hyperlink = q.url

                # Формируем информативную всплывающую подсказку в ячейке
                comment_obj = _build_cell_comment(q, item.unit)
                cell_link.comment = comment_obj
                cell_price.comment = comment_obj
            else:
                cell_price.value = ""
                cell_link.value = ""

        # Формулы Мин. цена и Средняя цена (СТРОГО по отображаемым столбцам поставщиков, защита от циклических ссылок)
        displayed_count = min(max_suppliers, len(display_quotes))
        clean_col_letters = [
            get_column_letter(3 + (s - 1) * 2)
            for s in range(1, displayed_count + 1)
        ]

        cell_min = ws.cell(row=current_row, column=col_min_idx)
        cell_min.border = thin_border
        cell_min.alignment = align_right
        cell_min.number_format = "#,##0.00"
        cell_min.font = font_price_bold

        cell_avg = ws.cell(row=current_row, column=col_avg_idx)
        cell_avg.border = thin_border
        cell_avg.alignment = align_right
        cell_avg.number_format = "#,##0.00"
        cell_avg.font = font_calibri_bold

        cell_var = ws.cell(row=current_row, column=col_var_idx)
        cell_var.border = thin_border
        cell_var.alignment = align_center
        cell_var.font = font_calibri_regular

        if clean_col_letters:
            args_str = ",".join(f"{col_let}{current_row}" for col_let in clean_col_letters)
            cell_min.value = f"=MIN({args_str})"
            cell_avg.value = f"=AVERAGE({args_str})"
            cell_var.value = f'=_xlfn.IFS(COUNTA({args_str})>=3,_xlfn.STDEV.S({args_str})/AVERAGE({args_str})*100,COUNTA({args_str})<3,"недостаточно данных")'
        else:
            cell_min.value = ""
            cell_avg.value = ""
            cell_var.value = ""

        current_row += 1

    # Настройка высоты строк шапки
    ws.row_dimensions[1].height = 28  # Поставщик N
    ws.row_dimensions[2].height = 22  # ИНН / КПП
    ws.row_dimensions[3].height = 36  # Контакты
    ws.row_dimensions[4].height = 24  # Заголовки колонок "Цена руб./ед.", "источник (ссылка)"

    # Динамическая подгонка высоты строк под фактический объем текста (исключает обрезание контактов и ссылок)
    for r in range(5, current_row):
        max_lines_in_row = 1

        # Колонка A (Номенклатура, ширина 48)
        nom_val = str(ws.cell(row=r, column=1).value or "").strip()
        if nom_val:
            nom_lines = sum(max(1, (len(p) + 41) // 42) for p in nom_val.split("\n"))
            max_lines_in_row = max(max_lines_in_row, nom_lines)

        # Колонки поставщиков (ширина 42)
        for s_idx in range(1, max_suppliers + 1):
            col_link_idx = 4 + (s_idx - 1) * 2
            cell_val = str(ws.cell(row=r, column=col_link_idx).value or "").strip()
            if cell_val:
                sup_lines = sum(max(1, (len(p) + 38) // 39) for p in cell_val.split("\n"))
                max_lines_in_row = max(max_lines_in_row, sup_lines)

        ws.row_dimensions[r].height = max(36, max_lines_in_row * 17 + 10)

    # Настройка ширины колонок Листа 1
    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 12
    for s_idx in range(1, max_suppliers + 1):
        col_start_let = get_column_letter(3 + (s_idx - 1) * 2)
        col_end_let = get_column_letter(4 + (s_idx - 1) * 2)
        ws.column_dimensions[col_start_let].width = 16
        ws.column_dimensions[col_end_let].width = 42

    ws.column_dimensions[col_min_let].width = 18
    ws.column_dimensions[col_avg_let].width = 18
    ws.column_dimensions[col_var_let].width = 24

    # --- ЛИСТ 2: Подробный детализированный реестр для закупщика/1С ---
    ws2 = wb.create_sheet(title="Детализированный реестр")
    headers_reg = ["№", "Номенклатура", "Ед. изм. заявки", "Статус", "Мин. цена (руб)"]
    for s_idx in range(1, max_suppliers + 1):
        headers_reg.extend([
            f"Поставщик {s_idx}",
            f"ИНН {s_idx}",
            f"КПП {s_idx}",
            f"Контакты {s_idx}",
            f"Город и расстояние {s_idx}",
            f"Цена {s_idx} (руб)",
            f"Исходная цена на сайте {s_idx}",
            f"Ед. на сайте {s_idx}",
            f"Пересчет / Примечание {s_idx}",
            f"Ссылка {s_idx}",
        ])

    ws2.append(headers_reg)
    for col_idx in range(1, len(headers_reg) + 1):
        cell = ws2.cell(row=1, column=col_idx)
        cell.font = font_calibri_bold
        cell.fill = header_fill
        cell.alignment = align_center
        cell.border = thin_border

    row_num = 2
    for idx, item in enumerate(items, 1):
        display_quotes = [q for q in item.quotes if q.price_rub > 0 and not q.is_outlier]
        if display_quotes:
            status_text = _price_sources_status(display_quotes)
        else:
            status_text = _missing_price_status(item)
        best_p = min(q.price_rub for q in display_quotes) if display_quotes else "-"

        row_vals = [
            item.source_position_no or item.row_number,
            item.raw_name,
            item.unit,
            status_text,
            best_p,
        ]
        for s_idx in range(1, max_suppliers + 1):
            if s_idx <= len(display_quotes):
                q = display_quotes[s_idx - 1]
                loc_val = "—"
                if q.supplier_city and q.supplier_distance_km is not None:
                    loc_val = f"{q.supplier_city} (~{q.supplier_distance_km} км)"
                    if q.supplier_distance_radius_status == "outside_radius":
                        loc_val += " (вне радиуса)"
                elif q.supplier_city:
                    loc_val = q.supplier_city

                row_vals.extend([
                    q.company_name or "-",
                    q.inn or "-",
                    q.kpp or "-",
                    q.contacts or "-",
                    loc_val,
                    q.price_rub if q.price_rub > 0 else "по запросу",
                    q.raw_price_rub if q.raw_price_rub > 0 else "-",
                    q.site_unit or "-",
                    "; ".join(filter(None, [q.conversion_comment, q.comment, _quote_evidence_note(q)])),
                    q.url,
                ])
            else:
                row_vals.extend(["-", "-", "-", "-", "-", "-", "-", "-", "-", "-"])
        ws2.append(row_vals)
        ws2.cell(row=row_num, column=2).comment = Comment(_source_position_note(item), "TenderLex")
        for col_idx in range(1, len(headers_reg) + 1):
            cell = ws2.cell(row=row_num, column=col_idx)
            cell.font = font_calibri_regular
            cell.border = thin_border
            if col_idx == 4:
                strong_sources = len(display_quotes) >= 3 and all(q.evidence_level == "page_content" for q in display_quotes)
                cell.font = Font(name="Calibri", bold=True, color="166534" if strong_sources else "92400E" if display_quotes else "64748B")
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            elif col_idx > 5 and (col_idx - 5) % 10 == 0:  # Ссылка
                val = str(cell.value or "")
                if val.startswith("http"):
                    cell.hyperlink = val
                    cell.font = font_link
        row_num += 1

    # Высота и ширина колонок листа 2
    ws2.row_dimensions[1].height = 26
    for r in range(2, row_num):
        ws2.row_dimensions[r].height = 42
    ws2.column_dimensions["A"].width = 6
    ws2.column_dimensions["B"].width = 46
    ws2.column_dimensions["C"].width = 14
    ws2.column_dimensions["D"].width = 40
    ws2.column_dimensions["E"].width = 18

    wb.save(output_path)
    logger.info("Saved price monitoring excel to %s", output_path)
    return output_path
