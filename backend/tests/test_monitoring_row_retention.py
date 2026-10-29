import io
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

import openpyxl
import pytest

from app.price_monitoring.pipeline import run_price_monitoring_pipeline
from app.price_monitoring.models import ItemPriceResult


@pytest.mark.asyncio
async def test_processing_exception_keeps_input_row_and_excel_number(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.append(["Наименование", "Ед. изм."])
    for name in ["Клапан", "Насос", "Фильтр"]:
        wb.active.append([name, "шт"])
    stream = io.BytesIO()
    wb.save(stream)

    async def process(item, *args, **kwargs):
        number, name, unit = item
        if number == 2:
            raise RuntimeError("network failed")
        return ItemPriceResult(number, name, name, unit, status="not_found")

    with patch("app.price_monitoring.pipeline.process_single_item", side_effect=process), patch("app.price_monitoring.pipeline.YandexSearchEngine") as engine:
        engine.return_value.total_requests = 0
        engine.return_value.close = AsyncMock()
        result = await run_price_monitoring_pipeline(stream.getvalue(), SimpleNamespace(yandex_search_folder_id="", yandex_search_api_key=""), str(tmp_path / "report.xlsx"))
    assert [row.row_number for row in result.items] == [1, 2, 3]
    assert result.items[1].status == "error"
    assert result.items[1].not_found_reason == "processing_error"
    output = openpyxl.load_workbook(tmp_path / "report.xlsx")
    assert output["Детализированный реестр"].cell(3, 1).value == 2


def test_detailed_register_keeps_nonconsecutive_source_numbers(tmp_path):
    from app.price_monitoring.excel_report import generate_price_monitoring_excel
    path = tmp_path / "gaps.xlsx"
    generate_price_monitoring_excel([ItemPriceResult(2, "Насос", "Насос", "шт"), ItemPriceResult(9, "Клапан", "Клапан", "шт")], str(path))
    workbook = openpyxl.load_workbook(path)
    assert [workbook["Детализированный реестр"].cell(i, 1).value for i in [2, 3]] == [2, 9]
    assert "9" in workbook["Пример отчета"].cell(6, 1).comment.text


@pytest.mark.asyncio
async def test_interrupted_run_retries_error_row_and_preserves_other_rows(tmp_path):
    from app.price_monitoring.pipeline import _save_checkpoint
    checkpoint = tmp_path / "checkpoint.json"
    _save_checkpoint(str(checkpoint), [ItemPriceResult(1, "Насос", "Насос", "шт", status="not_found"), ItemPriceResult(2, "Клапан", "Клапан", "шт", status="error", not_found_reason="processing_error")], 2)
    async def process(item, *args, **kwargs):
        return ItemPriceResult(item[0], item[1], item[1], item[2], status="not_found")
    with patch("app.price_monitoring.pipeline.process_single_item", side_effect=process) as task, patch("app.price_monitoring.pipeline.YandexSearchEngine") as engine:
        engine.return_value.total_requests = 0
        engine.return_value.close = AsyncMock()
        result = await run_price_monitoring_pipeline(b"", SimpleNamespace(yandex_search_folder_id="", yandex_search_api_key=""), str(tmp_path / "resume.xlsx"), items=[("Насос", "шт"), ("Клапан", "шт")], checkpoint_path=str(checkpoint))
    assert task.call_count == 1
    assert task.call_args.args[0][0] == 2
    assert [item.row_number for item in result.items] == [1, 2]
    assert all(item.status != "error" for item in result.items)
