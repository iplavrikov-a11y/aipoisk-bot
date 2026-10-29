from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from app.document_parser import _extract_xlsx


def _write_spec(path: Path, *, formatted_far_cell: bool) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = 'ТЗ'
    for row, values in enumerate([
        ('Наименование', 'Количество', 'Характеристики'),
        ('Товар А', 2, 'Размер 10×20 мм'),
        ('Товар Б', 3, 'Длина 30 мм'),
    ], 1):
        for column, value in zip((1, 4, 15), values):
            sheet.cell(row, column, value)
    if formatted_far_cell:
        sheet['XFD51'].font = Font(bold=True)
    book.save(path)
    book.close()


def test_formatted_empty_columns_do_not_change_procurement_text(tmp_path):
    normal = tmp_path / 'normal.xlsx'
    expanded = tmp_path / 'expanded.xlsx'
    _write_spec(normal, formatted_far_cell=False)
    _write_spec(expanded, formatted_far_cell=True)

    text = _extract_xlsx(expanded)

    assert text == _extract_xlsx(normal)
    assert len(text) < 1000
    assert 'Товар А' in text[:120000]
    assert 'Товар Б' in text[:120000]
    assert 'Размер 10×20 мм' in text
    # Interior gaps still describe the same columns; only empty tails disappear.
    assert 'Товар А |  |  | 2' in text


def test_empty_formatted_template_does_not_gain_product_rows(tmp_path):
    path = tmp_path / 'empty.xlsx'
    book = Workbook()
    book.active['XFD51'].font = Font(bold=True)
    book.save(path)
    book.close()

    text = _extract_xlsx(path)

    assert text.strip() == '=== SHEET: Sheet ==='
    assert len(text) < 100


def test_far_column_with_real_value_is_not_discarded(tmp_path):
    path = tmp_path / 'real-far-column.xlsx'
    book = Workbook()
    book.active['A1'] = 'Товар'
    book.active['XFD1'] = 'Значимое требование'
    book.save(path)
    book.close()

    text = _extract_xlsx(path)

    assert 'Товар' in text
    assert 'Значимое требование' in text
