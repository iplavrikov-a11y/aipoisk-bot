from __future__ import annotations

import asyncio
from copy import deepcopy
import io
import json
import os
from pathlib import Path
from tempfile import TemporaryDirectory
import threading
import unittest
from unittest.mock import patch
from types import SimpleNamespace

import app.supplier_search as supplier_search
from app.supplier_search import Candidate, CandidateMatch, CandidateRerank, ProcurementItem, ProcurementProfile


class SupplierDiscoveryFlowTests(unittest.IsolatedAsyncioTestCase):
    def test_base_domain_preserves_registrable_domain_for_multi_label_suffixes(self) -> None:
        self.assertEqual(supplier_search.base_domain("https://shop.vacuum.com.ru/catalog"), "vacuum.com.ru")
        self.assertEqual(supplier_search.base_domain("metiz.com.tw"), "metiz.com.tw")
        self.assertEqual(supplier_search.base_domain("https://www.example.co.uk"), "example.co.uk")

    def test_confidence_rejects_ambiguous_zero_to_ten_scale(self) -> None:
        with self.assertRaisesRegex(RuntimeError, "ambiguous 0-10"):
            supplier_search._confidence_percent(9)
        self.assertEqual(supplier_search._confidence_percent(0.91), 91)
        self.assertEqual(supplier_search._confidence_percent(91), 91)

    def test_verified_contacts_must_exist_in_extracted_page_contacts(self) -> None:
        self.assertEqual(
            supplier_search._verified_email("invented@wrong.example", ["real@supplier.example"]),
            "real@supplier.example",
        )
        self.assertEqual(
            supplier_search._verified_phone("+7 (617) 324-13-09", ["+86 173 2413 0960"]),
            "+8617324130960",
        )
        self.assertEqual(supplier_search._normalize_phone("+886-2-278-45675"), "+886227845675")

    def test_profile_merges_labeled_okpd2_from_source_context(self) -> None:
        profile = ProcurementProfile(
            summary="Установка",
            items=(ProcurementItem(id="item-1", name="Вакуумная установка"),),
        )
        merged = supplier_search._merge_deterministic_okpd2(
            profile,
            "Код ОКПД2: 28.21.13.117. ГОСТ 12.2.003-91.",
        )
        self.assertEqual(merged.items[0].okpd2_codes, ("28.21.13.117",))

    def test_detects_russian_origin_advantage_without_making_registry_mandatory(self) -> None:
        self.assertEqual(
            supplier_search._deterministic_measure_type(
                "Предоставляется преимущество в отношении товара российского происхождения"
            ),
            "advantage",
        )

    async def test_post_verification_reserve_pass_skips_primary_provider(self) -> None:
        calls: list[tuple[str, int]] = []

        async def yandex(*_args, **_kwargs):
            raise AssertionError("primary provider must not run in reserve-only pass")

        async def google(_settings, _queries, max_results, **_kwargs):
            calls.append(("google", max_results))
            return [Candidate(url="https://reserve.example", domain="reserve.example", source="google")]

        with (
            patch.object(supplier_search, "_search_with_yandex", yandex),
            patch.object(supplier_search, "_search_with_google", google),
        ):
            candidates, meta = await supplier_search.discover_candidates(
                SimpleNamespace(supplier_search_provider_order="google"),
                ["товар поставщик"],
                max_results=80,
                fallback_candidate_limit=8,
                reserve_only=True,
            )

        self.assertEqual(calls, [("google", 8)])
        self.assertEqual([candidate.domain for candidate in candidates], ["reserve.example"])
        self.assertTrue(meta["strategy"]["reserve_only"])
        self.assertEqual(meta["reports"][0]["status"], "skipped_post_verification_reserve")

    def test_technical_fit_dominates_contact_completeness_in_quality_score(self) -> None:
        common = {
            "evidence_status": "verified",
            "site": "https://supplier.example",
            "evidence_url": "https://supplier.example/product",
            "contact_url": "https://supplier.example/contact",
            "search_query": "оборудование поставщик",
            "company_name": "Supplier",
        }
        exact_without_phone = supplier_search._supplier_quality_score(
            {**common, "match_level": "exact", "product_fit": "exact", "phone": "", "email": ""}
        )
        category_with_contacts = supplier_search._supplier_quality_score(
            {
                **common,
                "match_level": "profile",
                "product_fit": "category",
                "phone": "+79991112233",
                "email": "sales@supplier.example",
            }
        )
        self.assertGreater(exact_without_phone, category_with_contacts)

    async def test_candidate_verification_limit_is_shared_between_concurrent_reviews(self) -> None:
        original_verify = supplier_search.verify_candidate
        active = 0
        max_active = 0

        async def fake_verify(*args, **kwargs):
            nonlocal active, max_active
            active += 1
            max_active = max(max_active, active)
            await asyncio.sleep(0.02)
            active -= 1
            return {"evidence_status": "verified"}

        supplier_search.verify_candidate = fake_verify
        supplier_search._supplier_verification_limiter.reset_for_tests()
        candidate = Candidate(url="https://supplier.example", domain="supplier.example")
        try:
            with (
                patch.object(supplier_search.config, "supplier_verification_concurrency", 2),
                patch.object(supplier_search.config, "supplier_verification_timeout_seconds", 1.0),
            ):
                await asyncio.gather(
                    *[
                        supplier_search._verify_candidate_with_limits(
                            SimpleNamespace(),
                            candidate,
                            "context",
                        )
                        for _ in range(6)
                    ]
                )
        finally:
            supplier_search.verify_candidate = original_verify
            supplier_search._supplier_verification_limiter.reset_for_tests()

        self.assertEqual(max_active, 2)

    async def test_candidate_verification_timeout_returns_no_result(self) -> None:
        original_verify = supplier_search.verify_candidate

        async def slow_verify(*args, **kwargs):
            await asyncio.sleep(0.1)
            return {"evidence_status": "verified"}

        supplier_search.verify_candidate = slow_verify
        supplier_search._supplier_verification_limiter.reset_for_tests()
        try:
            with patch.object(
                supplier_search,
                "_supplier_verification_timeout_seconds",
                return_value=0.01,
            ):
                result = await supplier_search._verify_candidate_with_limits(
                    SimpleNamespace(),
                    Candidate(url="https://supplier.example", domain="supplier.example"),
                    "context",
                )
        finally:
            supplier_search.verify_candidate = original_verify
            supplier_search._supplier_verification_limiter.reset_for_tests()

        self.assertIsNone(result)

    async def test_candidate_verification_retries_one_browser_infrastructure_failure(self) -> None:
        original_verify = supplier_search.verify_candidate
        calls = 0

        async def flaky_verify(*args, **kwargs):
            nonlocal calls
            calls += 1
            if calls == 1:
                raise supplier_search.SupplierBrowserInfrastructureError("temporary")
            return {"evidence_status": "verified"}

        supplier_search.verify_candidate = flaky_verify
        supplier_search._supplier_verification_limiter.reset_for_tests()
        try:
            result = await supplier_search._verify_candidate_with_limits(
                SimpleNamespace(),
                Candidate(url="https://supplier.example", domain="supplier.example"),
                "context",
            )
        finally:
            supplier_search.verify_candidate = original_verify
            supplier_search._supplier_verification_limiter.reset_for_tests()

        self.assertEqual(result, {"evidence_status": "verified"})
        self.assertEqual(calls, 2)

    async def test_candidate_verification_skips_persistent_browser_infrastructure_failure(self) -> None:
        original_verify = supplier_search.verify_candidate
        calls = 0

        async def failing_verify(*args, **kwargs):
            nonlocal calls
            calls += 1
            raise supplier_search.SupplierBrowserInfrastructureError("temporary")

        supplier_search.verify_candidate = failing_verify
        supplier_search._supplier_verification_limiter.reset_for_tests()
        try:
            result = await supplier_search._verify_candidate_with_limits(
                SimpleNamespace(),
                Candidate(url="https://supplier.example", domain="supplier.example"),
                "context",
            )
        finally:
            supplier_search.verify_candidate = original_verify
            supplier_search._supplier_verification_limiter.reset_for_tests()

        self.assertIsNone(result)
        self.assertEqual(calls, 2)

    async def test_candidate_verification_limit_is_shared_across_event_loops(self) -> None:
        original_verify = supplier_search.verify_candidate
        active = 0
        max_active = 0
        state_lock = threading.Lock()

        async def fake_verify(*args, **kwargs):
            nonlocal active, max_active
            with state_lock:
                active += 1
                max_active = max(max_active, active)
            await asyncio.sleep(0.02)
            with state_lock:
                active -= 1
            return {"evidence_status": "verified"}

        async def run_reviews() -> None:
            candidate = Candidate(url="https://supplier.example", domain="supplier.example")
            await asyncio.gather(
                *[
                    supplier_search._verify_candidate_with_limits(
                        SimpleNamespace(), candidate, "context"
                    )
                    for _ in range(6)
                ]
            )

        supplier_search.verify_candidate = fake_verify
        supplier_search._supplier_verification_limiter.reset_for_tests()
        try:
            with (
                patch.object(supplier_search.config, "supplier_verification_concurrency", 2),
                patch.object(supplier_search.config, "supplier_verification_timeout_seconds", 1.0),
            ):
                await asyncio.gather(
                    asyncio.to_thread(asyncio.run, run_reviews()),
                    asyncio.to_thread(asyncio.run, run_reviews()),
                )
        finally:
            supplier_search.verify_candidate = original_verify
            supplier_search._supplier_verification_limiter.reset_for_tests()

        self.assertEqual(max_active, 2)

    async def test_candidate_timeout_budget_starts_after_capacity_is_acquired(self) -> None:
        original_verify = supplier_search.verify_candidate

        async def fake_verify(*args, **kwargs):
            await asyncio.sleep(0.08)
            return {"evidence_status": "verified"}

        supplier_search.verify_candidate = fake_verify
        supplier_search._supplier_verification_limiter.reset_for_tests()
        candidate = Candidate(url="https://supplier.example", domain="supplier.example")
        try:
            with (
                patch.object(supplier_search.config, "supplier_verification_concurrency", 1),
                patch.object(supplier_search, "_supplier_verification_timeout_seconds", return_value=0.12),
            ):
                results = await asyncio.gather(
                    supplier_search._verify_candidate_with_limits(
                        SimpleNamespace(), candidate, "context"
                    ),
                    supplier_search._verify_candidate_with_limits(
                        SimpleNamespace(), candidate, "context"
                    ),
                )
        finally:
            supplier_search.verify_candidate = original_verify
            supplier_search._supplier_verification_limiter.reset_for_tests()

        self.assertEqual(
            results,
            [
                {"evidence_status": "verified"},
                {"evidence_status": "verified"},
            ],
        )

    async def test_collect_pages_adds_browser_rendered_page_when_http_has_no_contact(self) -> None:
        original_fetch_page = supplier_search.fetch_page
        original_browser = supplier_search.fetch_page_with_browser
        original_links = supplier_search.extract_internal_links

        async def fake_fetch_page(*args, **kwargs) -> dict:
            return {"url": "https://supplier.example", "html": "<html></html>", "text": "Каталог поставщика без опубликованного контакта"}

        async def fake_browser(url: str) -> dict:
            return {
                "url": url,
                "html": "<html></html>",
                "text": "Каталог поставщика. Контакты sales@supplier.example +7 999 111 22 33",
            }

        supplier_search.fetch_page = fake_fetch_page
        supplier_search.fetch_page_with_browser = fake_browser
        supplier_search.extract_internal_links = lambda *_args, **_kwargs: []
        try:
            pages = await supplier_search.collect_pages("https://supplier.example")
        finally:
            supplier_search.fetch_page = original_fetch_page
            supplier_search.fetch_page_with_browser = original_browser
            supplier_search.extract_internal_links = original_links

        self.assertEqual(pages[0]["text"], "Каталог поставщика. Контакты sales@supplier.example +7 999 111 22 33")
        self.assertEqual(len(pages), 1)

    async def test_build_procurement_profile_requires_ai_and_normalizes_items(self) -> None:
        original_call_llm = supplier_search.call_llm

        async def fake_call_llm(*args, **kwargs) -> str:
            return """
            {
              "summary": "Поставка двух товарных групп",
              "items": [
                {
                  "id": "main",
                  "name": "Сварочный полуавтомат",
                  "aliases": ["MIG/MAG"],
                  "okpd2_codes": ["28.29.70.110"],
                  "category_terms": ["сварочное оборудование"],
                  "exact_terms": ["500А"],
                  "required_terms": ["500А"]
                },
                {"name": "Блок жидкостного охлаждения", "excluded_terms": ["доставка"]}
              ],
              "excluded_terms": ["ТОРГ-12"]
            }
            """

        supplier_search.call_llm = fake_call_llm
        try:
            profile = await supplier_search.build_procurement_profile(
                SimpleNamespace(has_active_ai_provider=True),
                "Поставка сварочного полуавтомата и блока охлаждения.",
            )
        finally:
            supplier_search.call_llm = original_call_llm

        self.assertEqual(profile.summary, "Поставка двух товарных групп")
        self.assertEqual([item.id for item in profile.items], ["main", "item-2"])
        self.assertEqual(profile.items[0].aliases, ("MIG/MAG",))
        self.assertEqual(profile.items[0].okpd2_codes, ("28.29.70.110",))
        self.assertEqual(profile.items[0].category_terms, ("сварочное оборудование",))
        self.assertEqual(profile.items[0].exact_terms, ("500А",))
        self.assertEqual(profile.excluded_terms, ("ТОРГ-12",))

    async def test_procurement_profile_filters_regulatory_codes_from_okpd2(self) -> None:
        original_call_llm = supplier_search.call_llm

        async def fake_call_llm(*args, **kwargs) -> str:
            return """
            {
              "summary": "Поставка экологической лаборатории",
              "items": [
                {
                  "id": "main",
                  "name": "Передвижная экологическая лаборатория",
                  "okpd2_codes": ["52.04.840", "28.99.39.190"],
                  "exact_terms": ["РД 52.04.840-2015", "ГОСТ Р 8.589-2001"],
                  "category_terms": ["передвижные лаборатории"]
                }
              ],
              "excluded_terms": []
            }
            """

        supplier_search.call_llm = fake_call_llm
        try:
            profile = await supplier_search.build_procurement_profile(
                SimpleNamespace(has_active_ai_provider=True),
                "Передвижная лаборатория. Требования: РД 52.04.840-2015.",
            )
        finally:
            supplier_search.call_llm = original_call_llm

        self.assertEqual(profile.items[0].okpd2_codes, ("28.99.39.190",))

    async def test_minprom_registry_requirement_is_ai_gated(self) -> None:
        original_call_llm = supplier_search.call_llm
        calls: list[dict] = []

        async def fake_call_llm(*args, **kwargs) -> str:
            calls.append(kwargs)
            return """
            {
              "required": true,
              "measure_type": "prohibition",
              "evidence": "Запрет действует. Требуются реестровые записи Минпромторга.",
              "reason": "В ТЗ прямо указано требование реестровой записи"
            }
            """

        supplier_search.call_llm = fake_call_llm
        try:
            requirement = await supplier_search.assess_minprom_registry_requirement(
                SimpleNamespace(has_active_ai_provider=True),
                "Нацрежим: запрет действует. Требуются реестровые записи Минпромторга.",
            )
        finally:
            supplier_search.call_llm = original_call_llm

        self.assertTrue(requirement.required)
        self.assertEqual(requirement.measure_type, "prohibition")
        self.assertEqual(calls[0]["routing_key"], "minprom_registry_requirement")

    async def test_minprom_registry_requirement_requires_ai_provider(self) -> None:
        with self.assertRaises(RuntimeError):
            await supplier_search.assess_minprom_registry_requirement(
                SimpleNamespace(has_active_ai_provider=False),
                "Требуются реестровые записи Минпромторга.",
            )

    def test_minprom_registry_cache_uses_shared_sqlite_data_dir(self) -> None:
        with TemporaryDirectory() as tmp:
            data_dir = Path(tmp) / "shared-data"
            database_path = data_dir / "aipoisk.db"
            old_legacy_dir = os.environ.pop("SUPPLIER_MINPROM_REGISTRY_CACHE_DIR", None)
            try:
                with (
                    patch.object(supplier_search.config, "minprom_registry_cache_dir", ""),
                    patch.object(supplier_search.config, "database_url", f"sqlite:///{database_path}"),
                ):
                    self.assertEqual(
                        supplier_search._minprom_registry_cache_dir(),
                        data_dir / "minprom_registry",
                    )
            finally:
                if old_legacy_dir is not None:
                    os.environ["SUPPLIER_MINPROM_REGISTRY_CACHE_DIR"] = old_legacy_dir

    def test_minprom_registry_cache_dir_explicit_setting_overrides_database_location(self) -> None:
        with TemporaryDirectory() as tmp:
            explicit_dir = Path(tmp) / "registry-cache"
            old_legacy_dir = os.environ.pop("SUPPLIER_MINPROM_REGISTRY_CACHE_DIR", None)
            try:
                with (
                    patch.object(supplier_search.config, "minprom_registry_cache_dir", str(explicit_dir)),
                    patch.object(supplier_search.config, "database_url", "sqlite:////unrelated/data/aipoisk.db"),
                ):
                    self.assertEqual(supplier_search._minprom_registry_cache_dir(), explicit_dir)
            finally:
                if old_legacy_dir is not None:
                    os.environ["SUPPLIER_MINPROM_REGISTRY_CACHE_DIR"] = old_legacy_dir

    async def test_minprom_registry_search_uses_local_jsonl_index(self) -> None:
        with TemporaryDirectory() as tmp:
            index_path = Path(tmp) / "registry.jsonl"
            sqlite_path = Path(tmp) / "registry.sqlite"
            rows = [
                {
                    "manufacturer": 'АО "Катайский насосный завод"',
                    "product": "Насос центробежный типа Д",
                    "inn": "4509000018",
                    "registry_number": "РПП-НАСОС",
                    "source_url": "https://gisp.gov.ru/pp719v2/pub/prod/",
                    "evidence": "Производитель: АО Катайский насосный завод",
                    "row_text": "Насос центробежный типа Д РПП-НАСОС",
                },
                {
                    "manufacturer": 'ООО "Шум"',
                    "product": "Металлическая мебель",
                    "inn": "7700000000",
                    "registry_number": "РПП-ШУМ",
                    "row_text": "Металлическая мебель РПП-ШУМ",
                },
            ]
            index_path.write_text("\n".join(json.dumps(row, ensure_ascii=False) for row in rows), encoding="utf-8")
            old_env = {
                "SUPPLIER_MINPROM_REGISTRY_INDEX_PATH": os.environ.get("SUPPLIER_MINPROM_REGISTRY_INDEX_PATH"),
                "SUPPLIER_MINPROM_REGISTRY_SQLITE_PATH": os.environ.get("SUPPLIER_MINPROM_REGISTRY_SQLITE_PATH"),
                "SUPPLIER_MINPROM_REGISTRY_XLSX_CACHE_PATH": os.environ.get("SUPPLIER_MINPROM_REGISTRY_XLSX_CACHE_PATH"),
            }
            os.environ["SUPPLIER_MINPROM_REGISTRY_INDEX_PATH"] = str(index_path)
            os.environ["SUPPLIER_MINPROM_REGISTRY_SQLITE_PATH"] = str(sqlite_path)
            os.environ["SUPPLIER_MINPROM_REGISTRY_XLSX_CACHE_PATH"] = str(Path(tmp) / "missing.xlsx")
            try:
                entries = await supplier_search.search_minprom_registry_entries(["насос центробежный"], max_results=5)
                self.assertEqual(entries[0]["registry_number"], "РПП-НАСОС")
                self.assertEqual(entries[0]["source"], "minprom_registry_local_sqlite")
                self.assertTrue(sqlite_path.is_file())
            finally:
                for key, value in old_env.items():
                    if value is None:
                        os.environ.pop(key, None)
                    else:
                        os.environ[key] = value

    async def test_minprom_registry_search_reports_missing_local_index(self) -> None:
        with TemporaryDirectory() as tmp:
            old_env = {
                "SUPPLIER_MINPROM_REGISTRY_INDEX_PATH": os.environ.get("SUPPLIER_MINPROM_REGISTRY_INDEX_PATH"),
                "SUPPLIER_MINPROM_REGISTRY_SQLITE_PATH": os.environ.get("SUPPLIER_MINPROM_REGISTRY_SQLITE_PATH"),
                "SUPPLIER_MINPROM_REGISTRY_XLSX_CACHE_PATH": os.environ.get("SUPPLIER_MINPROM_REGISTRY_XLSX_CACHE_PATH"),
            }
            os.environ["SUPPLIER_MINPROM_REGISTRY_INDEX_PATH"] = str(Path(tmp) / "missing.jsonl")
            os.environ["SUPPLIER_MINPROM_REGISTRY_SQLITE_PATH"] = str(Path(tmp) / "missing.sqlite")
            os.environ["SUPPLIER_MINPROM_REGISTRY_XLSX_CACHE_PATH"] = str(Path(tmp) / "missing.xlsx")
            try:
                with self.assertRaises(RuntimeError) as raised:
                    await supplier_search.search_minprom_registry_entries(["насос центробежный"], max_results=5)
            finally:
                for key, value in old_env.items():
                    if value is None:
                        os.environ.pop(key, None)
                    else:
                        os.environ[key] = value

        self.assertIn("локальный индекс реестра Минпромторга отсутствует", str(raised.exception))

    async def test_minprom_registry_preflight_requires_ready_local_indexes(self) -> None:
        with TemporaryDirectory() as tmp:
            xlsx_path = Path(tmp) / "registry.xlsx"
            index_path = Path(tmp) / "registry.jsonl"
            sqlite_path = Path(tmp) / "registry.sqlite"
            old_env = {
                "SUPPLIER_MINPROM_REGISTRY_INDEX_PATH": os.environ.get("SUPPLIER_MINPROM_REGISTRY_INDEX_PATH"),
                "SUPPLIER_MINPROM_REGISTRY_SQLITE_PATH": os.environ.get("SUPPLIER_MINPROM_REGISTRY_SQLITE_PATH"),
                "SUPPLIER_MINPROM_REGISTRY_XLSX_CACHE_PATH": os.environ.get("SUPPLIER_MINPROM_REGISTRY_XLSX_CACHE_PATH"),
            }
            os.environ["SUPPLIER_MINPROM_REGISTRY_INDEX_PATH"] = str(index_path)
            os.environ["SUPPLIER_MINPROM_REGISTRY_SQLITE_PATH"] = str(sqlite_path)
            os.environ["SUPPLIER_MINPROM_REGISTRY_XLSX_CACHE_PATH"] = str(xlsx_path)
            try:
                self.assertIn(
                    "Локальный реестр Минпромторга не готов",
                    supplier_search.minprom_registry_preflight_error(supplier_search.SUPPLIER_POLICY_MINPROM_ONLY),
                )
                xlsx_path.write_bytes(b"PK\x03\x04")
                index_path.write_text(
                    json.dumps(
                        {
                            "manufacturer": 'АО "Катайский насосный завод"',
                            "product": "Насос центробежный типа Д",
                            "inn": "4509000018",
                            "registry_number": "РПП-НАСОС",
                            "row_text": "Насос центробежный типа Д РПП-НАСОС",
                        },
                        ensure_ascii=False,
                    )
                    + "\n",
                    encoding="utf-8",
                )
                await supplier_search.search_minprom_registry_entries(["насос центробежный"], max_results=5)
                self.assertEqual(
                    supplier_search.minprom_registry_preflight_error(supplier_search.SUPPLIER_POLICY_MINPROM_PRIORITY),
                    "",
                )
            finally:
                for key, value in old_env.items():
                    if value is None:
                        os.environ.pop(key, None)
                    else:
                        os.environ[key] = value

    async def test_minprom_registry_sqlite_empty_result_does_not_scan_jsonl_fallback(self) -> None:
        with TemporaryDirectory() as tmp:
            index_path = Path(tmp) / "registry.jsonl"
            sqlite_path = Path(tmp) / "registry.sqlite"
            index_path.write_text(
                json.dumps(
                    {
                        "manufacturer": 'АО "Катайский насосный завод"',
                        "product": "Насос центробежный типа Д",
                        "registry_number": "РПП-НАСОС",
                        "row_text": "Насос центробежный типа Д РПП-НАСОС",
                    },
                    ensure_ascii=False,
                )
                + "\n",
                encoding="utf-8",
            )
            old_env = {
                "SUPPLIER_MINPROM_REGISTRY_INDEX_PATH": os.environ.get("SUPPLIER_MINPROM_REGISTRY_INDEX_PATH"),
                "SUPPLIER_MINPROM_REGISTRY_SQLITE_PATH": os.environ.get("SUPPLIER_MINPROM_REGISTRY_SQLITE_PATH"),
                "SUPPLIER_MINPROM_REGISTRY_XLSX_CACHE_PATH": os.environ.get("SUPPLIER_MINPROM_REGISTRY_XLSX_CACHE_PATH"),
            }
            original_jsonl = supplier_search._search_minprom_registry_jsonl
            os.environ["SUPPLIER_MINPROM_REGISTRY_INDEX_PATH"] = str(index_path)
            os.environ["SUPPLIER_MINPROM_REGISTRY_SQLITE_PATH"] = str(sqlite_path)
            os.environ["SUPPLIER_MINPROM_REGISTRY_XLSX_CACHE_PATH"] = str(Path(tmp) / "missing.xlsx")
            try:
                await supplier_search.search_minprom_registry_entries(["насос центробежный"], max_results=5)
                supplier_search._search_minprom_registry_jsonl = lambda *_args, **_kwargs: (_ for _ in ()).throw(
                    AssertionError("JSONL fallback must not run when SQLite is ready")
                )
                entries = await supplier_search.search_minprom_registry_entries(["персональный компьютер"], max_results=5)
            finally:
                supplier_search._search_minprom_registry_jsonl = original_jsonl
                for key, value in old_env.items():
                    if value is None:
                        os.environ.pop(key, None)
                    else:
                        os.environ[key] = value

        self.assertEqual(entries, [])

    async def test_minprom_registry_xlsx_upload_builds_local_indexes(self) -> None:
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Предприятие-изготовитель", "Продукция", "ИНН", "Реестровый номер", "Срок действия"])
        sheet.append(['АО "Катайский насосный завод"', "Насос центробежный типа Д", "4509000018", "РПП-НАСОС", "31.12.2028"])
        buffer = io.BytesIO()
        workbook.save(buffer)
        workbook.close()

        with TemporaryDirectory() as tmp:
            xlsx_path = Path(tmp) / "registry.xlsx"
            index_path = Path(tmp) / "registry.jsonl"
            sqlite_path = Path(tmp) / "registry.sqlite"
            old_env = {
                "SUPPLIER_MINPROM_REGISTRY_INDEX_PATH": os.environ.get("SUPPLIER_MINPROM_REGISTRY_INDEX_PATH"),
                "SUPPLIER_MINPROM_REGISTRY_SQLITE_PATH": os.environ.get("SUPPLIER_MINPROM_REGISTRY_SQLITE_PATH"),
                "SUPPLIER_MINPROM_REGISTRY_XLSX_CACHE_PATH": os.environ.get("SUPPLIER_MINPROM_REGISTRY_XLSX_CACHE_PATH"),
            }
            os.environ["SUPPLIER_MINPROM_REGISTRY_INDEX_PATH"] = str(index_path)
            os.environ["SUPPLIER_MINPROM_REGISTRY_SQLITE_PATH"] = str(sqlite_path)
            os.environ["SUPPLIER_MINPROM_REGISTRY_XLSX_CACHE_PATH"] = str(xlsx_path)
            try:
                status = supplier_search.store_minprom_registry_xlsx_cache(buffer.getvalue(), filename="registry.xlsx")
                entries = await supplier_search.search_minprom_registry_entries(["насос центробежный"], max_results=5)
            finally:
                for key, value in old_env.items():
                    if value is None:
                        os.environ.pop(key, None)
                    else:
                        os.environ[key] = value

        self.assertTrue(status["xlsx_exists"])
        self.assertTrue(status["index_exists"])
        self.assertTrue(status["sqlite_ready"])
        self.assertEqual(status["index_count"], 1)
        self.assertEqual(entries[0]["registry_number"], "РПП-НАСОС")

    async def test_discover_minprom_registry_context_reports_source_errors(self) -> None:
        originals = {
            "build_minprom_registry_queries": supplier_search.build_minprom_registry_queries,
            "search_minprom_registry_entries": supplier_search.search_minprom_registry_entries,
        }

        async def fake_queries(*args, **kwargs):
            return ["канат стальной"]

        async def fake_search(*args, **kwargs):
            raise RuntimeError("GISP registry search failed: browser unavailable")

        supplier_search.build_minprom_registry_queries = fake_queries
        supplier_search.search_minprom_registry_entries = fake_search
        try:
            context = await supplier_search.discover_minprom_registry_context(
                SimpleNamespace(has_active_ai_provider=True),
                "ТЗ: запрет, требуется реестровая запись",
                ProcurementProfile(summary="Канат", items=(ProcurementItem(id="item-1", name="Канат стальной"),)),
                supplier_search.MinpromRegistryRequirement(required=True, measure_type="prohibition"),
            )
        finally:
            for name, original in originals.items():
                setattr(supplier_search, name, original)

        self.assertEqual(context.status, "error")
        self.assertIn("GISP registry search failed", context.error)

    def test_minprom_supplier_queries_are_only_for_required_prohibition(self) -> None:
        profile = ProcurementProfile(
            summary="Канат стальной",
            items=(ProcurementItem(id="item-1", name="Канат стальной", category_terms=("стальные канаты",)),),
        )
        required = supplier_search.MinpromRegistryRequirement(required=True, measure_type="prohibition")
        not_required = supplier_search.MinpromRegistryRequirement(required=False, measure_type="restriction")
        registry_context = supplier_search.MinpromRegistryContext(
            requirement=required,
            entries=(
                {
                    "registry_number": "123",
                    "manufacturer": "Канатный завод",
                    "product": "Канат стальной",
                },
            ),
            status="ok",
        )

        queries = supplier_search._build_minprom_supplier_queries(profile, registry_context)
        skipped = supplier_search._build_minprom_supplier_queries(
            profile,
            supplier_search.MinpromRegistryContext(requirement=not_required, status="not_required"),
        )

        self.assertEqual(skipped, [])
        self.assertTrue(any("Канатный завод" in query for query in queries))
        self.assertTrue(any("реестр Минпромторга" in query or "ГИСП" in query for query in queries))

    def test_minprom_supplier_queries_prioritize_product_terms_over_procurement_titles(self) -> None:
        profile = ProcurementProfile(
            summary="Поставка функциональной мебели для ГБОУ школа № 353",
            items=(
                ProcurementItem(
                    id="item-1",
                    name="Поставка функциональной мебели для ГБОУ школа № 353",
                    category_terms=("функциональная мебель", "школьная мебель"),
                ),
            ),
        )
        registry_context = supplier_search.MinpromRegistryContext(
            requirement=supplier_search.MinpromRegistryRequirement(required=True, measure_type="prohibition"),
            status="error",
            error="GISP timeout",
        )

        queries = supplier_search._build_minprom_supplier_queries(profile, registry_context, limit=6)

        self.assertTrue(queries[0].startswith('"функциональная мебель"'))
        self.assertIn("производитель", queries[0])
        self.assertTrue(any("школьная мебель" in query for query in queries))
        self.assertFalse(any("для ГБОУ" in query for query in queries[:4]))

    def test_minprom_queries_use_okpd2_hierarchy_and_points(self) -> None:
        profile = ProcurementProfile(
            summary="Поставка мебели",
            items=(
                ProcurementItem(
                    id="item-1",
                    name="Стол обеденный",
                    okpd2_codes=("31.09.12.131",),
                    category_terms=("стол обеденный",),
                ),
            ),
        )
        registry_context = supplier_search.MinpromRegistryContext(
            requirement=supplier_search.MinpromRegistryRequirement(required=True, measure_type="prohibition"),
            status="error",
            error="GISP timeout",
        )

        registry_queries = supplier_search._build_minprom_registry_code_queries(profile, limit=20)
        supplier_queries = supplier_search._build_minprom_supplier_queries(profile, registry_context, limit=14)

        self.assertIn("ОКПД2 31.09.12.131 реестр Минпромторга", registry_queries)
        self.assertIn('"31.09.12.131" ПП 719 баллы', registry_queries)
        self.assertIn("ОКПД2 31.09.12 реестр Минпромторга", registry_queries)
        self.assertTrue(any("31.09.1" in query for query in registry_queries))
        self.assertTrue(any("31.09.12.131" in query and "производитель" in query for query in supplier_queries))

    def test_minprom_comment_claims_are_removed_without_registry_entries(self) -> None:
        context = supplier_search.MinpromRegistryContext(
            requirement=supplier_search.MinpromRegistryRequirement(required=True, measure_type="prohibition"),
            status="error",
            error="GISP timeout",
        )

        comment = supplier_search._sanitize_minprom_comment_claims(
            (
                "Поставщик предлагает релевантный товар. "
                "Модель включена в реестр Минпромторга РФ, что удовлетворяет требованиям национального режима."
            ),
            context,
        )

        self.assertIn("Поставщик предлагает релевантный товар", comment)
        self.assertNotRegex(comment, r"(?i)минпромторг|гисп|реестров|национальн")

    async def test_minprom_registry_context_filters_irrelevant_raw_candidates(self) -> None:
        original_build_queries = supplier_search.build_minprom_registry_queries
        original_search_entries = supplier_search.search_minprom_registry_entries
        original_call_llm = supplier_search.call_llm

        async def fake_build_queries(*args, **kwargs):
            return ["видеоспектральный компаратор", "Regula 4308"]

        async def fake_search_entries(*args, **kwargs):
            return [
                {
                    "manufacturer": 'АО "ВРЕМЯ-Ч"',
                    "product": "Компаратор частотный VCH-314",
                    "inn": "5262007965",
                    "registry_number": "",
                },
                {
                    "manufacturer": 'ООО "ВАРТОН"',
                    "product": "Светодиодный светильник VARTON архитектурный Regula 2.0 1200",
                    "inn": "7731470910",
                    "registry_number": "",
                },
            ]

        async def fake_call_llm(*args, **kwargs):
            return '{"accepted_indexes":[]}'

        supplier_search.build_minprom_registry_queries = fake_build_queries
        supplier_search.search_minprom_registry_entries = fake_search_entries
        supplier_search.call_llm = fake_call_llm
        try:
            context = await supplier_search.discover_minprom_registry_context(
                SimpleNamespace(has_active_ai_provider=True),
                "ТЗ: компаратор видеоспектральный Regula 4308",
                ProcurementProfile(
                    summary="Поставка видеоспектрального компаратора",
                    items=(ProcurementItem(id="item-1", name="Компаратор видеоспектральный"),),
                ),
                supplier_search.MinpromRegistryRequirement(required=True, measure_type="restriction"),
            )
        finally:
            supplier_search.build_minprom_registry_queries = original_build_queries
            supplier_search.search_minprom_registry_entries = original_search_entries
            supplier_search.call_llm = original_call_llm

        self.assertEqual(context.status, "empty")
        self.assertEqual(context.candidate_count, 2)
        self.assertEqual(context.entries, ())

    def test_minprom_registry_match_annotates_supplier_origin(self) -> None:
        registry_context = supplier_search.MinpromRegistryContext(
            requirement=supplier_search.MinpromRegistryRequirement(required=True, measure_type="prohibition"),
            entries=(
                {
                    "registry_number": "РПП-НАСОС",
                    "manufacturer": 'АО "Катайский насосный завод"',
                    "product": "Насос центробежный типа Д",
                    "inn": "4509000018",
                    "evidence": "Производитель: АО Катайский насосный завод",
                },
            ),
            status="ok",
        )
        accepted = [
            {
                "company_name": 'АО "Катайский насосный завод"',
                "product": "Насосы центробежные",
                "site": "https://knz.example",
                "comments": "Официальный сайт производителя.",
                "evidence_status": "verified",
                "quality_score": 90,
            }
        ]

        supplier_search._annotate_minprom_registry_matches(
            accepted,
            registry_context,
            supplier_search.SUPPLIER_POLICY_MINPROM_ONLY,
        )

        self.assertEqual(accepted[0]["supplier_search_origin"], "minprom_registry")
        self.assertTrue(accepted[0]["minprom_registry_match"]["matched"])
        self.assertIn(accepted[0]["minprom_registry_match"]["method"], {"manufacturer", "manufacturer_product"})

    def test_minprom_registry_filter_rejects_generic_registry_claim_without_entry_match(self) -> None:
        registry_context = supplier_search.MinpromRegistryContext(
            requirement=supplier_search.MinpromRegistryRequirement(required=True, measure_type="prohibition"),
            entries=(
                {
                    "registry_number": "РПП-КАНАТ",
                    "manufacturer": "Канатный завод",
                    "product": "Канат стальной",
                    "inn": "7700000001",
                },
            ),
            status="ok",
        )
        accepted = [
            {
                "company_name": "Посторонний поставщик",
                "product": "Канат стальной",
                "comments": "Есть реестровая запись Минпромторга.",
                "site": "https://supplier.example",
                "evidence_status": "verified",
                "quality_score": 90,
            }
        ]

        filtered = supplier_search._filter_minprom_verified_suppliers(accepted, registry_context)

        self.assertEqual(filtered, [])

    async def test_strict_minprom_preserves_verified_rows_for_trustworthy_zero_fallback(self) -> None:
        profile = ProcurementProfile(
            summary="Волоконный усилитель",
            items=(ProcurementItem(id="item-1", name="Волоконный усилитель"),),
        )
        candidate = Candidate(
            url="https://supplier.example",
            domain="supplier.example",
            title="Поставщик",
        )

        async def fake_profile(*_args, **_kwargs):
            return profile

        async def fake_queries(*_args, **_kwargs):
            return ["волоконный усилитель поставщик"]

        async def fake_candidates(*_args, **_kwargs):
            return [candidate], {"reports": [{"provider": "test", "status": "ok"}]}

        async def fake_rerank(*_args, **_kwargs):
            return CandidateRerank([candidate], {"status": "ok"})

        for registry_status, expected_reason in (
            ("empty", "registry_no_relevant_entries"),
            ("ok", "registry_entries_no_supplier_match"),
        ):
            with self.subTest(registry_status=registry_status):
                registry_context = supplier_search.MinpromRegistryContext(
                    requirement=supplier_search.MinpromRegistryRequirement(
                        required=True,
                        measure_type="prohibition",
                    ),
                    entries=(
                        {
                            "registry_number": "РПП-ДРУГОЙ-ТОВАР",
                            "manufacturer": "Другой завод",
                            "product": "Другой товар",
                            "inn": "7700000000",
                        },
                    ) if registry_status == "ok" else (),
                    status=registry_status,
                )
                source_rows = [
                    {
                        "company_name": f"Поставщик {index}",
                        "site": f"https://supplier-{index}.example",
                        "email": f"sales{index}@supplier.example",
                        "product": "Волоконный усилитель",
                        "product_fit": "exact",
                        "evidence_status": "verified",
                        "quality_score": 90 - index,
                    }
                    for index in range(3)
                ]
                filtered_input: list[dict] = []
                filter_seen = False
                real_filter = supplier_search._filter_minprom_verified_suppliers

                async def fake_registry(*_args, **_kwargs):
                    if filter_seen:
                        raise AssertionError("registry discovery ran after strict filter")
                    return registry_context

                async def fake_review(*_args, **_kwargs):
                    if filter_seen:
                        raise AssertionError("candidate review ran after strict filter")
                    return source_rows, [], {"reviewed_count": len(source_rows)}

                async def unexpected_recovery(*_args, **_kwargs):
                    raise AssertionError("no recovery/search call is expected after three verified rows")

                def recording_filter(rows, context):
                    nonlocal filter_seen
                    filtered_input.extend(deepcopy(rows))
                    filter_seen = True
                    return real_filter(rows, context)

                with (
                    patch.object(supplier_search, "build_procurement_profile", fake_profile),
                    patch.object(supplier_search, "build_supplier_queries", fake_queries),
                    patch.object(supplier_search, "discover_minprom_registry_context", fake_registry),
                    patch.object(supplier_search, "discover_candidates", fake_candidates),
                    patch.object(supplier_search, "ai_rerank_candidates", fake_rerank),
                    patch.object(supplier_search, "_review_candidates_until_target", fake_review),
                    patch.object(supplier_search, "_build_supplier_recovery_queries_with_ai", unexpected_recovery),
                    patch.object(supplier_search, "_filter_minprom_verified_suppliers", recording_filter),
                ):
                    accepted, evidence = await supplier_search._discover_suppliers_impl(
                        SimpleNamespace(has_active_ai_provider=True),
                        "ТЗ на волоконный усилитель",
                        target=3,
                        supplier_search_policy=supplier_search.SUPPLIER_POLICY_MINPROM_ONLY,
                    )

                alternative = evidence["non_registry_alternative"]
                self.assertEqual(accepted, [])
                self.assertEqual(evidence["accepted"], [])
                self.assertEqual(evidence["accepted_count"], 0)
                self.assertEqual(evidence["registry_result"], {"status": registry_status, "verified_count": 0})
                self.assertTrue(alternative["available"])
                self.assertEqual(alternative["verified_count"], 3)
                self.assertEqual(alternative["reason_code"], expected_reason)
                self.assertEqual(alternative["verified_rows"], filtered_input)
                self.assertTrue(
                    all(row["supplier_search_origin"] == "ordinary_fallback" for row in alternative["verified_rows"])
                )

                source_rows[0]["company_name"] = "Изменено после формирования evidence"
                self.assertEqual(alternative["verified_rows"][0]["company_name"], "Поставщик 0")

    async def test_strict_minprom_registry_error_never_enables_fallback_offer(self) -> None:
        rows = [{"company_name": "Поставщик", "site": "https://supplier.example"}]
        context = supplier_search.MinpromRegistryContext(
            requirement=supplier_search.MinpromRegistryRequirement(required=True, measure_type="prohibition"),
            status="error",
            error="registry unavailable",
        )
        supplier_search._annotate_minprom_registry_matches(
            rows,
            context,
            supplier_search.SUPPLIER_POLICY_MINPROM_ONLY,
        )
        alternative = supplier_search._non_registry_alternative_evidence(
            rows,
            context,
            registry_verified_count=0,
        )

        self.assertFalse(alternative["available"])
        self.assertEqual(alternative["reason_code"], "")
        self.assertEqual(alternative["verified_count"], 1)

    def test_supplier_search_blocks_tender_and_registry_mirror_domains(self) -> None:
        for domain in ("poisktenderov.ru", "awindex.ru", "torgs.ru", "ruscable.ru", "zakupki44fz.ru", "dzen.ru"):
            self.assertTrue(supplier_search.is_blocked(domain), domain)

    async def test_discover_suppliers_passes_required_minprom_registry_context_to_review(self) -> None:
        originals = {
            "build_procurement_profile": supplier_search.build_procurement_profile,
            "build_supplier_queries": supplier_search.build_supplier_queries,
            "discover_candidates": supplier_search.discover_candidates,
            "ai_rerank_candidates": supplier_search.ai_rerank_candidates,
            "assess_minprom_registry_requirement": supplier_search.assess_minprom_registry_requirement,
            "discover_minprom_registry_context": supplier_search.discover_minprom_registry_context,
            "_review_candidates_until_target": supplier_search._review_candidates_until_target,
        }
        captured: dict = {}
        profile = ProcurementProfile(
            summary="Канат стальной",
            items=(ProcurementItem(id="item-1", name="Канат стальной", category_terms=("стальные канаты",)),),
        )
        candidate = Candidate(url="https://supplier.example", domain="supplier.example", title="поставщик")
        requirement = supplier_search.MinpromRegistryRequirement(
            required=True,
            measure_type="prohibition",
            evidence="Требуется реестровая запись",
            reason="Запрет действует",
        )
        registry_context = supplier_search.MinpromRegistryContext(
            requirement=requirement,
            queries=("стальные канаты ГИСП",),
            entries=({"registry_number": "123", "manufacturer": "Канатный завод", "product": "канат стальной"},),
            status="ok",
        )

        async def fake_profile(*args, **kwargs):
            return profile

        async def fake_queries(*args, **kwargs):
            return ["стальные канаты производитель"]

        async def fake_discover(*args, **kwargs):
            captured["queries"] = list(args[1])
            return [candidate], {"reports": [{"provider": "test", "status": "ok"}]}

        async def fake_rerank(*args, **kwargs):
            return CandidateRerank([candidate], {"status": "ok"})

        async def fake_assess(*args, **kwargs):
            return requirement

        async def fake_registry(*args, **kwargs):
            return registry_context

        async def fake_review(*args, registry_context=None, **kwargs):
            captured["registry_context"] = registry_context
            return (
                [
                    {
                        "company_name": "Канатный завод",
                        "site": "https://supplier.example",
                        "phone": "+7 999 111 22 33",
                        "email": "sales@supplier.example",
                        "comments": "Есть реестровая запись Минпромторга.",
                        "evidence_status": "verified",
                        "quality_score": 90,
                    }
                ],
                [],
                {"reviewed_count": 1},
            )

        supplier_search.build_procurement_profile = fake_profile
        supplier_search.build_supplier_queries = fake_queries
        supplier_search.discover_candidates = fake_discover
        supplier_search.ai_rerank_candidates = fake_rerank
        supplier_search.assess_minprom_registry_requirement = fake_assess
        supplier_search.discover_minprom_registry_context = fake_registry
        supplier_search._review_candidates_until_target = fake_review
        try:
            accepted, evidence = await supplier_search.discover_suppliers(
                SimpleNamespace(has_active_ai_provider=True, supplier_search_provider_order="test"),
                "ТЗ: запрет действует, требуется реестровая запись Минпромторга.",
                target=1,
            )
        finally:
            for name, original in originals.items():
                setattr(supplier_search, name, original)

        self.assertEqual(len(accepted), 1)
        self.assertEqual(accepted[0]["supplier_search_origin"], "minprom_registry")
        self.assertTrue(accepted[0]["minprom_registry_match"]["matched"])
        self.assertEqual(accepted[0]["minprom_registry_match"]["method"], "manufacturer")
        self.assertIs(captured["registry_context"], registry_context)
        self.assertIn('"Канатный завод" официальный сайт', captured["queries"][:6])
        self.assertIn("стальные канаты производитель", captured["queries"])
        self.assertTrue(evidence["minprom_registry"]["required"])
        self.assertEqual(evidence["minprom_registry"]["entries_count"], 1)
        self.assertTrue(any("Канатный завод" in query for query in evidence["minprom_supplier_queries"]))

    async def test_ai_rerank_candidates_keeps_ai_ranked_supplier_candidates(self) -> None:
        original_call_llm = supplier_search.call_llm

        async def fake_call_llm(*args, **kwargs) -> str:
            return """
            {
              "ranked": [
                {"id": "1", "keep": true, "confidence": 0.91, "procurement_item_id": "item-1", "reason": "официальный дилер"},
                {"id": "0", "keep": false, "confidence": 10, "procurement_item_id": "item-1", "reason": "агрегатор"}
              ]
            }
            """

        candidates = [
            Candidate(url="https://bad.example", domain="bad.example", title="каталог"),
            Candidate(url="https://good.example", domain="good.example", title="производитель"),
        ]
        supplier_search.call_llm = fake_call_llm
        try:
            rerank = await supplier_search.ai_rerank_candidates(
                SimpleNamespace(has_active_ai_provider=True),
                ProcurementProfile(summary="ТЗ", items=(ProcurementItem(id="item-1", name="Товар"),)),
                candidates,
                target=1,
            )
        finally:
            supplier_search.call_llm = original_call_llm

        self.assertEqual([candidate.domain for candidate in rerank.candidates], ["good.example"])
        self.assertEqual(rerank.candidates[0].procurement_item_id, "item-1")
        self.assertEqual(rerank.candidates[0].ai_rank_confidence, 91)

    async def test_ai_rerank_candidates_receives_minprom_context_when_required(self) -> None:
        original_call_llm = supplier_search.call_llm
        prompts: list[str] = []

        async def fake_call_llm(*args, **kwargs) -> str:
            prompts.append(str(args[1]))
            return """
            {
              "ranked": [
                {"id": "0", "keep": true, "confidence": 0.9, "procurement_item_id": "item-1", "reason": "производитель с реестровым сигналом"}
              ]
            }
            """

        registry_context = supplier_search.MinpromRegistryContext(
            requirement=supplier_search.MinpromRegistryRequirement(required=True, measure_type="prohibition"),
            entries=({"registry_number": "123", "manufacturer": "Канатный завод", "product": "Канат стальной"},),
            status="ok",
        )
        supplier_search.call_llm = fake_call_llm
        try:
            rerank = await supplier_search.ai_rerank_candidates(
                SimpleNamespace(has_active_ai_provider=True),
                ProcurementProfile(summary="Канат стальной", items=(ProcurementItem(id="item-1", name="Канат стальной"),)),
                [Candidate(url="https://kanat.example", domain="kanat.example", title="Канатный завод")],
                target=1,
                registry_context=registry_context,
            )
        finally:
            supplier_search.call_llm = original_call_llm

        self.assertEqual([candidate.domain for candidate in rerank.candidates], ["kanat.example"])
        self.assertIn("Контекст Минпромторга", prompts[0])
        self.assertIn("Канатный завод", prompts[0])
        self.assertIn("реестр Минпромторга", prompts[0])

    async def test_ai_rerank_candidates_uses_ai_expansion_when_initial_pool_is_too_small(self) -> None:
        original_call_llm = supplier_search.call_llm
        calls = 0

        async def fake_call_llm(*args, **kwargs) -> str:
            nonlocal calls
            calls += 1
            if calls == 1:
                return """
                {
                  "ranked": [
                    {"id": "0", "keep": true, "confidence": 0.95, "procurement_item_id": "item-1", "reason": "точная страница"},
                    {"id": "1", "keep": true, "confidence": 0.88, "procurement_item_id": "item-1", "reason": "профильный поставщик"}
                  ]
                }
                """
            return """
            {
              "ranked": [
                {"id": "2", "keep": true, "confidence": 0.8, "procurement_item_id": "item-1", "reason": "производитель категории"},
                {"id": "3", "keep": true, "confidence": 0.78, "procurement_item_id": "item-1", "reason": "завод категории"},
                {"id": "4", "keep": true, "confidence": 0.76, "procurement_item_id": "item-1", "reason": "дилер категории"},
                {"id": "5", "keep": true, "confidence": 0.74, "procurement_item_id": "item-1", "reason": "поставщик категории"}
              ]
            }
            """

        candidates = [
            Candidate(
                url=f"https://supplier-{index}.example/catalog",
                domain=f"supplier-{index}.example",
                title="стальные канаты производитель",
                snippet="поставщик канатной продукции",
                source="test",
                query="производитель стальных канатов",
            )
            for index in range(8)
        ]
        supplier_search.call_llm = fake_call_llm
        try:
            rerank = await supplier_search.ai_rerank_candidates(
                SimpleNamespace(has_active_ai_provider=True),
                ProcurementProfile(
                    summary="Канат стальной",
                    items=(
                        ProcurementItem(
                            id="item-1",
                            name="Канат стальной",
                            category_terms=("стальные канаты", "канатная продукция"),
                        ),
                    ),
                ),
                candidates,
                target=3,
            )
        finally:
            supplier_search.call_llm = original_call_llm

        self.assertEqual(calls, 2)
        self.assertEqual([candidate.domain for candidate in rerank.candidates[:6]], [f"supplier-{index}.example" for index in range(6)])
        self.assertEqual(rerank.meta["initial_kept_count"], 2)
        self.assertEqual(rerank.meta["expanded_kept_count"], 4)

    async def test_build_supplier_queries_requires_ai_provider(self) -> None:
        with self.assertRaises(RuntimeError):
            await supplier_search.build_supplier_queries(
                SimpleNamespace(has_active_ai_provider=False),
                "Поставка промышленного оборудования",
                target=5,
            )

    async def test_discover_suppliers_requires_ai_provider(self) -> None:
        with self.assertRaises(RuntimeError):
            await supplier_search.discover_suppliers(
                SimpleNamespace(has_active_ai_provider=False),
                "Поставка промышленного оборудования",
                target=5,
            )

    async def test_discover_candidates_filters_previously_found_domains(self) -> None:
        original_yandex = supplier_search._search_with_yandex
        captured: dict = {}

        async def fake_yandex(settings, queries: list[str], max_results: int, *, existing_domains=None):
            captured["existing_domains"] = set(existing_domains or set())
            return [
                Candidate(url="https://old.example/catalog", domain="old.example", title="old", source="yandex", query=queries[0]),
                Candidate(url="https://new.example/catalog", domain="new.example", title="new", source="yandex", query=queries[0]),
            ], 1

        supplier_search._search_with_yandex = fake_yandex
        try:
            candidates, _meta = await supplier_search.discover_candidates(
                SimpleNamespace(supplier_search_provider_order="yandex"),
                ["поставщик"],
                max_results=10,
                excluded_domains={"old.example"},
            )
        finally:
            supplier_search._search_with_yandex = original_yandex

        self.assertIn("old.example", captured["existing_domains"])
        self.assertEqual([candidate.domain for candidate in candidates], ["new.example"])

    async def test_build_supplier_queries_prefers_ai_queries_over_deterministic_fallback(self) -> None:
        original_call_llm = supplier_search.call_llm
        profile = ProcurementProfile(
            summary="Лазерный станок резки металла",
            items=(ProcurementItem(id="item-1", name="лазерный станок резки металла"),),
        )

        async def fake_call_llm(*args, **kwargs) -> str:
            return '{"queries": ["лазерный станок резки металла производитель", "официальный дилер лазерных станков"]}'

        supplier_search.call_llm = fake_call_llm
        try:
            queries = await supplier_search.build_supplier_queries(
                SimpleNamespace(has_active_ai_provider=True),
                "Техническое задание. Поставка лазерного станка резки металла.",
                target=5,
                profile=profile,
            )
        finally:
            supplier_search.call_llm = original_call_llm

        self.assertEqual(queries[:2], ["лазерный станок резки металла производитель", "официальный дилер лазерных станков"])

    async def test_build_supplier_queries_revises_overly_exact_query_set_with_ai(self) -> None:
        original_call_llm = supplier_search.call_llm
        calls: list[str] = []
        profile = ProcurementProfile(
            summary="Поставка стального каната 31 мм ЛК-РО",
            items=(
                ProcurementItem(
                    id="item-1",
                    name="Канат стальной",
                    aliases=("стальной канат ЛК-РО",),
                    required_terms=("диаметр 31 мм", "тип ЛК-РО", "ГОСТ 7668-80"),
                    category_terms=("стальные канаты", "канатная продукция"),
                    exact_terms=("31 мм", "ЛК-РО", "ГОСТ 7668-80"),
                ),
            ),
        )

        async def fake_call_llm(_settings, prompt: str, *args, **kwargs) -> str:
            calls.append(prompt)
            if len(calls) == 1:
                return """
                {"queries": [
                  "производитель стальных канатов ЛК-РО 31 мм",
                  "завод стальных канатов 31 мм ГОСТ 7668-80",
                  "купить канат стальной 31 мм ЛК-РО оптом"
                ]}
                """
            self.assertIn("слишком узко", prompt)
            return """
            {"queries": [
              "производитель стальных канатов",
              "завод канатной продукции",
              "поставщик стальных канатов оптом",
              "дилер стальных канатов",
              "метизный завод стальные канаты"
            ]}
            """

        supplier_search.call_llm = fake_call_llm
        try:
            queries = await supplier_search.build_supplier_queries(
                SimpleNamespace(has_active_ai_provider=True),
                "Поставка стального каната 31 мм ЛК-РО",
                target=15,
                profile=profile,
            )
        finally:
            supplier_search.call_llm = original_call_llm

        self.assertEqual(len(calls), 2)
        self.assertIn("производитель стальных канатов", queries)
        self.assertIn("завод канатной продукции", queries)
        self.assertTrue(any("31 мм" not in query and "ЛК-РО" not in query for query in queries))

    async def test_build_supplier_queries_retries_empty_message_ai_failure(self) -> None:
        original_call_llm = supplier_search.call_llm
        calls = 0
        profile = ProcurementProfile(
            summary="Сотовый поликарбонат",
            items=(ProcurementItem(id="item-1", name="сотовый поликарбонат"),),
        )

        async def fake_call_llm(*args, **kwargs) -> str:
            nonlocal calls
            calls += 1
            if calls == 1:
                raise TimeoutError()
            return '{"queries": ["сотовый поликарбонат производитель", "поликарбонатные панели поставщик"]}'

        supplier_search.call_llm = fake_call_llm
        try:
            queries = await supplier_search.build_supplier_queries(
                SimpleNamespace(has_active_ai_provider=True),
                "Поставка сотового поликарбоната",
                target=5,
                profile=profile,
            )
        finally:
            supplier_search.call_llm = original_call_llm

        self.assertEqual(calls, 2)
        self.assertEqual(queries[:2], ["сотовый поликарбонат производитель", "поликарбонатные панели поставщик"])

    async def test_build_supplier_queries_failure_includes_exception_type_after_retry(self) -> None:
        original_call_llm = supplier_search.call_llm
        calls = 0
        profile = ProcurementProfile(
            summary="Сотовый поликарбонат",
            items=(ProcurementItem(id="item-1", name="сотовый поликарбонат"),),
        )

        async def fake_call_llm(*args, **kwargs) -> str:
            nonlocal calls
            calls += 1
            raise TimeoutError()

        supplier_search.call_llm = fake_call_llm
        try:
            with self.assertRaises(RuntimeError) as raised:
                await supplier_search.build_supplier_queries(
                    SimpleNamespace(has_active_ai_provider=True),
                    "Поставка сотового поликарбоната",
                    target=5,
                    profile=profile,
                )
        finally:
            supplier_search.call_llm = original_call_llm

        self.assertEqual(calls, 2)
        self.assertIn("after retry", str(raised.exception))
        self.assertIn("TimeoutError", str(raised.exception))

    async def test_verify_candidate_does_not_fallback_to_keyword_acceptance_after_ai_reject(self) -> None:
        original_collect_pages = supplier_search.collect_pages
        original_ai_verify = supplier_search.ai_verify

        async def fake_collect_pages(url: str) -> list[dict]:
            return [
                {
                    "url": url,
                    "text": "Каталог. Сварочный полуавтомат MIG/MAG купить. Производитель. +7 999 111 22 33 sales@supplier.example",
                }
            ]

        async def fake_ai_verify(*args, **kwargs) -> dict:
            return {
                "action": "reject",
                "confidence": 20,
                "site_type": "directory",
                "product_fit": "category",
                "comments": "ИИ видит справочник, а не сайт поставщика.",
            }

        supplier_search.collect_pages = fake_collect_pages
        supplier_search.ai_verify = fake_ai_verify
        try:
            result = await supplier_search.verify_candidate(
                SimpleNamespace(has_active_ai_provider=True),
                Candidate(
                    url="https://supplier.example/catalog",
                    domain="supplier.example",
                    title="Сварочный полуавтомат купить",
                    snippet="производитель, каталог, контакты",
                    source="test",
                    query="сварочный полуавтомат поставщик",
                ),
                "Поставка сварочного полуавтомата MIG/MAG",
            )
        finally:
            supplier_search.collect_pages = original_collect_pages
            supplier_search.ai_verify = original_ai_verify

        self.assertIsNotNone(result)
        self.assertEqual(result["evidence_status"], "weak")
        self.assertIn("ИИ", result["comments"])

    async def test_verify_candidate_accepts_only_after_ai_accepts_when_ai_is_configured(self) -> None:
        original_collect_pages = supplier_search.collect_pages
        original_ai_verify = supplier_search.ai_verify

        async def fake_collect_pages(url: str) -> list[dict]:
            return [
                {
                    "url": url,
                    "text": "Официальный сайт производителя. Сварочный полуавтомат MIG/MAG. Контакты +7 999 111 22 33 sales@supplier.example",
                }
            ]

        async def fake_ai_verify(*args, **kwargs) -> dict:
            return {
                "action": "accept",
                "confidence": 0.92,
                "site_type": "manufacturer",
                "product_fit": "exact",
                "company_name": "Supplier",
                "status": "завод",
                "product": "Сварочный полуавтомат MIG/MAG",
                "phone": "+7 999 111 22 33",
                "email": "sales@supplier.example",
                "evidence_url": "https://supplier.example/catalog",
                "contact_url": "https://supplier.example/catalog",
                "evidence_snippet": "Сварочный полуавтомат MIG/MAG",
                "contact_evidence_snippet": "+7 999 111 22 33 sales@supplier.example",
                "comments": "ИИ подтвердил сайт производителя и контакт.",
            }

        supplier_search.collect_pages = fake_collect_pages
        supplier_search.ai_verify = fake_ai_verify
        try:
            result = await supplier_search.verify_candidate(
                SimpleNamespace(has_active_ai_provider=True),
                Candidate(
                    url="https://supplier.example/catalog",
                    domain="supplier.example",
                    title="Сварочный полуавтомат MIG/MAG",
                    snippet="производитель, каталог, контакты",
                    source="test",
                    query="сварочный полуавтомат поставщик",
                ),
                "Поставка сварочного полуавтомата MIG/MAG",
            )
        finally:
            supplier_search.collect_pages = original_collect_pages
            supplier_search.ai_verify = original_ai_verify

        self.assertIsNotNone(result)
        self.assertEqual(result["evidence_status"], "verified")
        self.assertEqual(result["company_name"], "Supplier")
        self.assertEqual(result["match_level"], "exact")
        self.assertEqual(result["ai_confidence"], 92)

    async def test_verify_candidate_replaces_ai_contact_placeholders_with_extracted_contacts(self) -> None:
        original_collect_pages = supplier_search.collect_pages
        original_ai_verify = supplier_search.ai_verify

        async def fake_collect_pages(url: str) -> list[dict]:
            return [
                {
                    "url": url,
                    "text": "Официальный сайт производителя. Сотовый поликарбонат. Контакты +7 999 111 22 33 sales@supplier.example",
                }
            ]

        async def fake_ai_verify(*args, **kwargs) -> dict:
            return {
                "action": "accept",
                "confidence": 0.92,
                "site_type": "manufacturer",
                "product_fit": "exact",
                "company_name": "Supplier",
                "status": "завод",
                "product": "Сотовый поликарбонат",
                "phone": "не найдено на главной",
                "email": "не найдено на главной",
                "evidence_url": "https://supplier.example/catalog",
                "contact_url": "https://supplier.example/catalog",
                "evidence_snippet": "Сотовый поликарбонат",
                "contact_evidence_snippet": "+7 999 111 22 33 sales@supplier.example",
                "comments": "ИИ подтвердил сайт производителя и контакт.",
            }

        supplier_search.collect_pages = fake_collect_pages
        supplier_search.ai_verify = fake_ai_verify
        try:
            result = await supplier_search.verify_candidate(
                SimpleNamespace(has_active_ai_provider=True),
                Candidate(
                    url="https://supplier.example/catalog",
                    domain="supplier.example",
                    title="Сотовый поликарбонат",
                    snippet="производитель, каталог, контакты",
                    source="test",
                    query="сотовый поликарбонат поставщик",
                ),
                "Поставка сотового поликарбоната",
            )
        finally:
            supplier_search.collect_pages = original_collect_pages
            supplier_search.ai_verify = original_ai_verify

        self.assertIsNotNone(result)
        self.assertEqual(result["phone"], "+7 (999) 111-22-33")
        self.assertEqual(result["email"], "sales@supplier.example")

    def test_verified_phone_normalizes_and_rejects_broken_numbers(self) -> None:
        self.assertEqual(supplier_search._verified_phone("8 800 119 00 00", []), "")
        self.assertEqual(supplier_search._verified_phone("8000\n1190000", []), "")
        self.assertEqual(
            supplier_search._verified_phone("не найдено", ["+7 999 111 22 33"]),
            "+7 (999) 111-22-33",
        )

    async def test_verify_candidate_rejects_ai_contact_placeholders_without_extracted_contacts(self) -> None:
        original_collect_pages = supplier_search.collect_pages
        original_ai_verify = supplier_search.ai_verify

        async def fake_collect_pages(url: str) -> list[dict]:
            return [{"url": url, "text": "Официальный сайт производителя. Сотовый поликарбонат."}]

        async def fake_ai_verify(*args, **kwargs) -> dict:
            return {
                "action": "accept",
                "confidence": 0.92,
                "site_type": "manufacturer",
                "product_fit": "exact",
                "company_name": "Supplier",
                "status": "завод",
                "product": "Сотовый поликарбонат",
                "phone": "не найдено на главной",
                "email": "не найдено на главной",
                "evidence_url": "https://supplier.example/catalog",
                "contact_url": "https://supplier.example/catalog",
                "evidence_snippet": "Сотовый поликарбонат",
                "contact_evidence_snippet": "не найдено на главной",
                "comments": "ИИ подтвердил сайт производителя и контакт.",
            }

        supplier_search.collect_pages = fake_collect_pages
        supplier_search.ai_verify = fake_ai_verify
        try:
            result = await supplier_search.verify_candidate(
                SimpleNamespace(has_active_ai_provider=True),
                Candidate(
                    url="https://supplier.example/catalog",
                    domain="supplier.example",
                    title="Сотовый поликарбонат",
                    snippet="производитель, каталог",
                    source="test",
                    query="сотовый поликарбонат поставщик",
                ),
                "Поставка сотового поликарбоната",
            )
        finally:
            supplier_search.collect_pages = original_collect_pages
            supplier_search.ai_verify = original_ai_verify

        self.assertIsNotNone(result)
        self.assertEqual(result["evidence_status"], "weak")
        self.assertIn("телефон", result["comments"].lower())

    async def test_ai_verify_prompt_forbids_full_match_claims_for_non_exact_fit(self) -> None:
        original_call_llm = supplier_search.call_llm
        captured: dict[str, str] = {}

        async def fake_call_llm(_settings, prompt: str, *args, **kwargs) -> str:
            captured["prompt"] = prompt
            return '{"action": "reject", "comments": "test"}'

        supplier_search.call_llm = fake_call_llm
        try:
            await supplier_search.ai_verify(
                SimpleNamespace(has_active_ai_provider=True),
                Candidate(
                    url="https://supplier.example/catalog",
                    domain="supplier.example",
                    title="Каталог поставщика",
                    snippet="производитель, каталог, контакты",
                    source="test",
                    query="сварочный полуавтомат поставщик",
                ),
                "Поставка сварочного полуавтомата MIG/MAG",
                [
                    {
                        "url": "https://supplier.example/catalog",
                        "text": "Каталог поставщика. Контакты +7 999 111 22 33 sales@supplier.example",
                    }
                ],
                ["sales@supplier.example"],
                ["+7 999 111 22 33"],
                CandidateMatch(True, "profile", "Промышленное оборудование", "Профильная категория"),
            )
        finally:
            supplier_search.call_llm = original_call_llm

        self.assertIn("Классификация product_fit", captured["prompt"])
        self.assertIn('запрещено писать "полностью соответствует"', captured["prompt"])

    def test_supplier_quality_score_caps_non_exact_product_fit(self) -> None:
        base = {
            "evidence_status": "verified",
            "site": "https://supplier.example/catalog",
            "evidence_url": "https://supplier.example/catalog",
            "contact_url": "https://supplier.example/contacts",
            "phone": "+7 999 111 22 33",
            "email": "sales@supplier.example",
            "search_query": "сварочный полуавтомат поставщик",
            "company_name": "Supplier",
        }

        exact_score = supplier_search._supplier_quality_score({**base, "match_level": "exact", "product_fit": "exact"})
        category_score = supplier_search._supplier_quality_score({**base, "match_level": "profile", "product_fit": "category"})
        profile_score = supplier_search._supplier_quality_score({**base, "match_level": "profile", "product_fit": "profile"})
        analog_score = supplier_search._supplier_quality_score({**base, "match_level": "adjacent", "product_fit": "analog"})

        self.assertGreaterEqual(exact_score, 80)
        self.assertLess(category_score, 80)
        self.assertLess(profile_score, 80)
        self.assertLess(analog_score, 80)

    def test_ai_rejection_reason_requires_stronger_category_fit(self) -> None:
        low_confidence = supplier_search._ai_rejection_reason(
            {
                "action": "accept",
                "confidence": 50,
                "site_type": "supplier",
                "product_fit": "category",
                "evidence_snippet": "Поставщик сварочного оборудования",
            }
        )
        missing_snippet = supplier_search._ai_rejection_reason(
            {
                "action": "accept",
                "confidence": 80,
                "site_type": "supplier",
                "product_fit": "profile",
                "evidence_snippet": "",
            }
        )
        accepted = supplier_search._ai_rejection_reason(
            {
                "action": "accept",
                "confidence": 70,
                "site_type": "supplier",
                "product_fit": "category",
                "evidence_snippet": "Каталог: сварочные полуавтоматы, источники MIG/MAG и расходные материалы",
            }
        )

        self.assertIn("низкая уверенность", low_confidence)
        self.assertIn("нет фрагмента сайта", missing_snippet)
        self.assertEqual(accepted, "")

    async def test_discover_suppliers_returns_already_reviewed_verified_results_above_minimum(self) -> None:
        original_build = supplier_search.build_supplier_queries
        original_profile = supplier_search.build_procurement_profile
        original_candidates = supplier_search.discover_candidates
        original_rerank = supplier_search.ai_rerank_candidates
        original_verify = supplier_search.verify_candidate
        original_assess = supplier_search.assess_minprom_registry_requirement
        calls: list[str] = []
        progress_events: list[tuple[int, str]] = []

        async def fake_profile(settings, context: str) -> ProcurementProfile:
            return ProcurementProfile(
                summary="Промышленное оборудование",
                items=(ProcurementItem(id="item-1", name="промышленное оборудование"),),
            )

        async def fake_build(settings, context: str, target: int, profile=None) -> list[str]:
            return ["query"]

        async def fake_candidates(settings, queries: list[str], max_results: int, **kwargs):
            return (
                [
                    Candidate(
                        url=f"https://supplier-{index}.ru",
                        domain=f"supplier-{index}.ru",
                        title="производитель",
                        snippet="оборудование",
                        source="test",
                        query="query",
                    )
                    for index in range(30)
                ],
                {"provider_order": ["test"], "reports": []},
            )

        async def fake_rerank(settings, profile: ProcurementProfile, candidates: list[Candidate], target: int, **kwargs) -> CandidateRerank:
            return CandidateRerank(candidates=candidates, meta={"status": "test", "kept_count": len(candidates)})

        async def fake_assess(settings, context: str):
            return supplier_search.MinpromRegistryRequirement(required=False, measure_type="unknown", reason="Требование не найдено")

        async def fake_verify(settings, candidate: Candidate, context: str, *, profile=None, registry_context=None):
            calls.append(candidate.domain)
            return {
                "company_name": candidate.domain,
                "site": candidate.url,
                "evidence_url": candidate.url,
                "contact_url": candidate.url,
                "phone": "+7 999 111 22 33",
                "email": "sales@example.ru",
                "evidence_status": "verified",
                "match_level": "exact",
                "source": candidate.source,
                "search_query": candidate.query,
            }

        async def fake_progress(progress: int, message: str) -> None:
            progress_events.append((progress, message))

        supplier_search.build_supplier_queries = fake_build
        supplier_search.build_procurement_profile = fake_profile
        supplier_search.discover_candidates = fake_candidates
        supplier_search.ai_rerank_candidates = fake_rerank
        supplier_search.assess_minprom_registry_requirement = fake_assess
        supplier_search.verify_candidate = fake_verify
        try:
            accepted, evidence = await supplier_search.discover_suppliers(
                SimpleNamespace(has_active_ai_provider=True),
                "поставка промышленного оборудования",
                target=3,
                progress_callback=fake_progress,
            )
        finally:
            supplier_search.build_supplier_queries = original_build
            supplier_search.build_procurement_profile = original_profile
            supplier_search.discover_candidates = original_candidates
            supplier_search.ai_rerank_candidates = original_rerank
            supplier_search.assess_minprom_registry_requirement = original_assess
            supplier_search.verify_candidate = original_verify

        self.assertEqual(len(accepted), 12)
        self.assertGreaterEqual(len(accepted), 3)
        self.assertLess(len(calls), 30)
        self.assertTrue(evidence["ai_required"])
        self.assertTrue(evidence["ai_used"])
        self.assertIn("supplier_candidate_verifier", evidence["ai_required_stages"])
        self.assertTrue(evidence["review"]["early_stop"])
        self.assertEqual(evidence["accepted_count"], 12)
        self.assertGreaterEqual(len(progress_events), 5)
        self.assertIn("Анализирую ТЗ", progress_events[0][1])
        self.assertTrue(any("Проверяю сайты" in message or "Проверено сайтов" in message for _, message in progress_events))

    async def test_discover_suppliers_runs_ai_recovery_search_when_first_round_underfills(self) -> None:
        original_build = supplier_search.build_supplier_queries
        original_profile = supplier_search.build_procurement_profile
        original_candidates = supplier_search.discover_candidates
        original_rerank = supplier_search.ai_rerank_candidates
        original_verify = supplier_search.verify_candidate
        original_assess = supplier_search.assess_minprom_registry_requirement
        original_recovery = supplier_search._build_supplier_recovery_queries_with_ai
        search_calls: list[list[str]] = []
        recovery_inputs: dict = {}
        progress_events: list[tuple[int, str]] = []

        async def fake_profile(settings, context: str) -> ProcurementProfile:
            return ProcurementProfile(
                summary="Промышленная номенклатура",
                items=(
                    ProcurementItem(
                        id="item-1",
                        name="промышленная номенклатура",
                        category_terms=("промышленная номенклатура", "заводы производители"),
                    ),
                ),
            )

        async def fake_build(settings, context: str, target: int, profile=None) -> list[str]:
            return ["точная позиция поставщик"]

        async def fake_candidates(settings, queries: list[str], max_results: int, **kwargs):
            search_calls.append(list(queries))
            if len(search_calls) == 1:
                candidates = [
                    Candidate(url=f"https://first-{index}.ru", domain=f"first-{index}.ru", title="поставщик", source="test", query=queries[0])
                    for index in range(4)
                ]
            else:
                candidates = [
                    Candidate(
                        url=f"https://recovery-{index}.ru",
                        domain=f"recovery-{index}.ru",
                        title="производитель категории",
                        source="test",
                        query=queries[0],
                    )
                    for index in range(4)
                ]
            return candidates, {"provider_order": ["test"], "reports": [{"provider": "test", "status": "ok"}]}

        async def fake_rerank(settings, profile: ProcurementProfile, candidates: list[Candidate], target: int, **kwargs) -> CandidateRerank:
            return CandidateRerank(candidates=candidates, meta={"status": "test", "kept_count": len(candidates)})

        async def fake_assess(settings, context: str):
            return supplier_search.MinpromRegistryRequirement(required=False, measure_type="unknown", reason="Требование не найдено")

        async def fake_recovery(settings, context: str, profile: ProcurementProfile, initial_queries, reviewed, accepted, target: int) -> list[str]:
            recovery_inputs["initial_queries"] = list(initial_queries)
            recovery_inputs["reviewed_count"] = len(reviewed)
            recovery_inputs["accepted_count"] = len(accepted)
            return ["производитель промышленной номенклатуры"]

        async def fake_verify(settings, candidate: Candidate, context: str, *, profile=None, registry_context=None):
            if candidate.domain == "first-0.ru" or candidate.domain.startswith("recovery-"):
                return {
                    "company_name": candidate.domain,
                    "site": candidate.url,
                    "evidence_url": candidate.url,
                    "contact_url": candidate.url,
                    "phone": "+7 999 111 22 33",
                    "email": "sales@example.ru",
                    "evidence_status": "verified",
                    "match_level": "exact",
                    "product_fit": "exact",
                    "ai_confidence": 95,
                    "source": candidate.source,
                    "search_query": candidate.query,
                }
            return {
                "company_name": candidate.domain,
                "site": candidate.url,
                "evidence_status": "weak",
                "source": candidate.source,
                "search_query": candidate.query,
                "comments": "AI-аудит отклонил кандидата.",
            }

        async def fake_progress(progress: int, message: str) -> None:
            progress_events.append((progress, message))

        supplier_search.build_supplier_queries = fake_build
        supplier_search.build_procurement_profile = fake_profile
        supplier_search.discover_candidates = fake_candidates
        supplier_search.ai_rerank_candidates = fake_rerank
        supplier_search.assess_minprom_registry_requirement = fake_assess
        supplier_search._build_supplier_recovery_queries_with_ai = fake_recovery
        supplier_search.verify_candidate = fake_verify
        try:
            accepted, evidence = await supplier_search.discover_suppliers(
                SimpleNamespace(has_active_ai_provider=True),
                "поставка промышленной номенклатуры",
                target=3,
                progress_callback=fake_progress,
            )
        finally:
            supplier_search.build_supplier_queries = original_build
            supplier_search.build_procurement_profile = original_profile
            supplier_search.discover_candidates = original_candidates
            supplier_search.ai_rerank_candidates = original_rerank
            supplier_search.assess_minprom_registry_requirement = original_assess
            supplier_search._build_supplier_recovery_queries_with_ai = original_recovery
            supplier_search.verify_candidate = original_verify

        self.assertGreaterEqual(len(accepted), 3)
        self.assertEqual(len(search_calls), 2)
        self.assertEqual(search_calls[1], ["производитель промышленной номенклатуры"])
        self.assertEqual(recovery_inputs["initial_queries"], ["точная позиция поставщик"])
        self.assertEqual(recovery_inputs["accepted_count"], 1)
        self.assertTrue(evidence["recovery_rounds"])
        self.assertEqual(evidence["recovery_rounds"][0]["status"], "ok")
        self.assertTrue(any("Расширяю поиск" in message for _, message in progress_events))

    async def test_ai_rerank_failure_includes_exception_type_when_message_is_empty(self) -> None:
        original_call_llm = supplier_search.call_llm
        calls = 0

        async def fake_call_llm(*args, **kwargs) -> str:
            nonlocal calls
            calls += 1
            raise TimeoutError()

        supplier_search.call_llm = fake_call_llm
        try:
            with self.assertRaises(RuntimeError) as raised:
                await supplier_search.ai_rerank_candidates(
                    SimpleNamespace(has_active_ai_provider=True),
                    ProcurementProfile(summary="ТЗ", items=(ProcurementItem(id="item-1", name="Товар"),)),
                    [Candidate(url="https://supplier.example", domain="supplier.example", title="поставщик")],
                    target=1,
                )
        finally:
            supplier_search.call_llm = original_call_llm

        self.assertEqual(calls, 2)
        self.assertIn("after retry", str(raised.exception))
        self.assertIn("TimeoutError", str(raised.exception))


if __name__ == "__main__":
    unittest.main()
