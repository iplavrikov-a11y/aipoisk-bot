import pytest

from app.price_monitoring.models import SupplierQuote
from app.supplier_locations import annotate_contacts_radius_status


def test_validate_geo_request_requires_a_known_city_and_bounded_pair():
    from app.geo_policy import validate_geo_request

    assert validate_geo_request("", None, mode="supplier_search") == ("", None)
    assert validate_geo_request("г. Казань", 300, mode="supplier_search") == ("Казань", 300)
    assert validate_geo_request("Казань", 300, mode="supplier_search_extra") == ("Казань", 300)

    with pytest.raises(ValueError, match="одновременно"):
        validate_geo_request("Казань", None, mode="supplier_search")
    with pytest.raises(ValueError, match="от 10 до 5000"):
        validate_geo_request("Казань", 0, mode="supplier_search")
    with pytest.raises(ValueError, match="не распознана"):
        validate_geo_request("Несуществующий город", 300, mode="supplier_search")
    with pytest.raises(ValueError, match="не поддерживает"):
        validate_geo_request("Казань", 300, mode="exact_product")


def test_region_without_supplier_city_stays_unknown_not_zero_distance():
    contacts = [{"company_name": "Склад без города", "supplier_city": "Московская область", "region": "Московская область"}]

    stats = annotate_contacts_radius_status(contacts, "Москва", 10)

    assert stats == {"within_radius": 0, "outside_radius": 0, "unknown_distance": 1}
    assert contacts[0]["supplier_distance_km"] is None
    assert contacts[0]["supplier_distance_within_radius"] is None
    assert contacts[0]["supplier_distance_radius_status"] == "unknown_distance"


def test_price_quotes_are_annotated_and_ranked_without_inventing_location():
    from app.supplier_locations import annotate_price_quotes_radius_status, sort_quotes_by_radius_priority

    quotes = [
        SupplierQuote(company_name="Далёкий", price_rub=100, contacts="Адрес: г. Москва"),
        SupplierQuote(company_name="Ближний", price_rub=110, contacts="Адрес: г. Шахты"),
        SupplierQuote(company_name="Без адреса", price_rub=90, contacts="+7 800 000-00-00"),
    ]

    stats = annotate_price_quotes_radius_status(quotes, "Каменск-Шахтинский", 150)
    sort_quotes_by_radius_priority(quotes)

    assert stats == {"within_radius": 1, "outside_radius": 1, "unknown_distance": 1}
    assert [quote.company_name for quote in quotes] == ["Ближний", "Без адреса", "Далёкий"]
    assert quotes[0].supplier_city == "Шахты"
    assert quotes[0].supplier_distance_within_radius is True
    assert quotes[0].supplier_location_source == "contact_address_city"
    assert quotes[0].supplier_location_confidence == "city_centroid"
    assert quotes[1].supplier_city == ""
    assert quotes[1].supplier_distance_within_radius is None
    assert quotes[2].supplier_distance_within_radius is False


def test_delivery_destination_is_not_treated_as_supplier_address():
    """A delivery promise is not evidence of a warehouse or legal address."""
    from app.supplier_locations import annotate_price_quotes_radius_status

    quote = SupplierQuote(
        company_name="Без подтверждённого адреса",
        price_rub=100,
        contacts="Доставка Москва и область; склад уточняйте",
    )

    stats = annotate_price_quotes_radius_status([quote], "Москва", 100)

    assert stats == {"within_radius": 0, "outside_radius": 0, "unknown_distance": 1}
    assert quote.supplier_city == ""
    assert quote.supplier_distance_within_radius is None


def test_city_next_to_an_unconfirmed_warehouse_is_not_address_evidence():
    from app.supplier_locations import annotate_price_quotes_radius_status

    quote = SupplierQuote(
        company_name="Адрес требуется уточнить",
        price_rub=100,
        contacts="Москва, склад уточняйте у менеджера",
    )

    stats = annotate_price_quotes_radius_status([quote], "Москва", 100)

    assert stats == {"within_radius": 0, "outside_radius": 0, "unknown_distance": 1}
    assert quote.supplier_distance_within_radius is None


def test_real_address_context_is_accepted_for_radius_calculation():
    """An actual address remains usable even without an explicit city prefix."""
    from app.supplier_locations import annotate_price_quotes_radius_status

    quote = SupplierQuote(
        company_name="Поставщик с адресом",
        price_rub=100,
        contacts="Юридический адрес: Москва, ул. Ленина, д. 5",
    )

    stats = annotate_price_quotes_radius_status([quote], "Москва", 100)

    assert stats == {"within_radius": 1, "outside_radius": 0, "unknown_distance": 0}
    assert quote.supplier_city == "Москва"
    assert quote.supplier_distance_within_radius is True


def test_contact_fields_do_not_combine_unrelated_address_and_city_fragments():
    """Address evidence and a city must come from the same source field."""
    contacts = [{
        "company_name": "Поставщик без города",
        "address": "Адрес уточняйте у менеджера",
        "comments": "Москва — условия поставки по запросу",
    }]

    stats = annotate_contacts_radius_status(contacts, "Москва", 100)

    assert stats == {"within_radius": 0, "outside_radius": 0, "unknown_distance": 1}
    assert contacts[0]["supplier_distance_within_radius"] is None


def test_price_report_geo_note_discloses_remote_and_unknown_offers():
    from app.price_monitoring.excel_report import _quote_geo_note

    remote = SupplierQuote(
        company_name="Далёкий",
        supplier_city="Москва",
        supplier_distance_km=900,
        supplier_distance_radius_status="outside_radius",
    )
    unknown = SupplierQuote(
        company_name="Без адреса",
        supplier_distance_radius_status="unknown_distance",
    )
    assert "вне радиуса" in _quote_geo_note(remote)
    assert "добор по РФ" in _quote_geo_note(remote)
    assert "расстояние не определено" in _quote_geo_note(unknown)


def test_supplier_xlsx_dedicated_location_column_and_clean_comment():
    import tempfile
    from pathlib import Path
    from openpyxl import load_workbook
    from app.report_builder import write_supplier_xlsx

    contacts = [
        {
            "company_name": "ООО Близкий",
            "site": "https://blizkiy.ru",
            "phone": "+7 999 111 22 33",
            "email": "info@blizkiy.ru",
            "product_fit": "exact",
            "comments": "Труба профильная 40х20х2",
            "supplier_city": "Шахты",
            "supplier_distance_km": 78,
            "supplier_distance_within_radius": True,
        },
        {
            "company_name": "ООО Далёкий",
            "site": "https://dalekiy.ru",
            "phone": "+7 999 444 55 66",
            "email": "sales@dalekiy.ru",
            "product_fit": "analog",
            "comments": "Труба профильная 40х20х2.5",
            "supplier_city": "Москва",
            "supplier_distance_km": 974,
            "supplier_distance_within_radius": False,
        },
    ]

    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "suppliers_geo.xlsx"
        write_supplier_xlsx(path, contacts, title="Тест труб", target=2, delivery_city="Каменск-Шахтинский", geo_radius_km=150)

        wb = load_workbook(path)
        ws = wb.active
        headers = [cell.value for cell in ws[5]]
        assert headers == [
            "Компания",
            "Сайт",
            "Телефоны",
            "Email",
            "Комментарий",
            "Город и расстояние",
            "Реестр Минпромторга",
        ]
        # Check row 6 (Близкий)
        c6_comment = ws["E6"].value
        c6_loc = ws["F6"].value
        assert "Локация:" not in c6_comment
        assert "Точное соответствие: Труба профильная 40х20х2." in c6_comment
        assert c6_loc == "Шахты (~78 км)"

        # Check row 7 (Далёкий)
        c7_comment = ws["E7"].value
        c7_loc = ws["F7"].value
        assert "Локация:" not in c7_comment
        assert "Возможный аналог: Труба профильная 40х20х2.5." in c7_comment
        assert c7_loc == "Москва (~974 км) (вне радиуса)"
        wb.close()


def test_price_monitoring_excel_dedicated_location_column_sheet2():
    import tempfile
    from pathlib import Path
    from openpyxl import load_workbook
    from app.price_monitoring.models import ItemPriceResult, SupplierQuote
    from app.price_monitoring.excel_report import generate_price_monitoring_excel

    q1 = SupplierQuote(
        company_name="ООО Локальный",
        price_rub=2500.0,
        raw_price_rub=2500.0,
        site_unit="шт",
        url="https://local.ru/item",
        contacts="+7 800 123-45-67, info@local.ru",
        supplier_city="Шахты",
        supplier_distance_km=78,
        supplier_distance_radius_status="within_radius",
        conversion_comment="Пересчет не требуется",
    )
    item = ItemPriceResult(
        row_number=1,
        raw_name="Задвижка стальная Ду50",
        clean_name="Задвижка стальная Ду50",
        unit="шт",
        status="ok",
        quotes=[q1],
    )

    with tempfile.TemporaryDirectory() as tmp:
        path = str(Path(tmp) / "price_geo.xlsx")
        generate_price_monitoring_excel([item], path, delivery_city="Каменск-Шахтинский", geo_radius_km=150)

        wb = load_workbook(path)
        ws2 = wb["Детализированный реестр"]
        headers_row = [cell.value for cell in ws2[1]]

        assert "Город и расстояние 1" in headers_row
        col_loc_idx = headers_row.index("Город и расстояние 1") + 1
        col_note_idx = headers_row.index("Пересчет / Примечание 1") + 1

        loc_val = ws2.cell(row=2, column=col_loc_idx).value
        note_val = ws2.cell(row=2, column=col_note_idx).value

        assert loc_val == "Шахты (~78 км)"
        assert "География:" not in note_val
        assert "Пересчет не требуется" in note_val
        wb.close()
