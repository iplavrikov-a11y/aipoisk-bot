from __future__ import annotations

import logging
from pathlib import Path
from typing import TYPE_CHECKING
from urllib.parse import urlparse

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Pt, RGBColor

if TYPE_CHECKING:
    from .models import ExactProductReport

logger = logging.getLogger("tenderlex.exact_product.docx")

# ---------------------------------------------------------------------------
# TenderLex Emerald Palette
# ---------------------------------------------------------------------------
BRAND_EMERALD = RGBColor(4, 120, 87)       # #047857 TenderLex Primary
DARK_EMERALD = "064E3B"                    # Deep Forest Emerald
MEDIUM_EMERALD = "0F766E"                  # Teal-Emerald
SUBTLE_EMERALD = "134E4A"                  # Slate-Emerald
BANNER_EMERALD = "064E3B"                  # Banner background
ZEBRA_MINT = "F4FBF7"                      # Subtle fresh mint
META_BG = "F8FAFC"
BORDER_COLOR = "CBD5E1"
TEXT_DARK = RGBColor(15, 23, 42)
TEXT_MUTED = RGBColor(71, 85, 105)
TEXT_GREEN = RGBColor(4, 120, 87)


def _set_docx_margins(doc: Document) -> None:
    for section in doc.sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.65)
        section.right_margin = Inches(0.65)
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)


def _set_cell_bg(cell, color_hex: str) -> None:
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)


def _set_table_fixed_width(table, width_in_inches: float = 6.97) -> None:
    tblPr = table._tbl.tblPr
    dxa_val = int(width_in_inches * 1440)
    tblW = parse_xml(f'<w:tblW {nsdecls("w")} w:w="{dxa_val}" w:type="dxa"/>')
    tblPr.append(tblW)


def _set_table_full_grid_borders(table, color: str = "CBD5E1") -> None:
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)


def _prevent_row_split(table) -> None:
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))


def _set_table_header_repeat(table) -> None:
    if table.rows:
        trPr = table.rows[0]._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))


def _add_docx_hyperlink(paragraph, url: str, text: str, color_hex="0284C7", underline=True):
    """Adds an interactive OpenXML hyperlink into a Word paragraph."""
    if not url or not str(url).strip():
        return paragraph.add_run(text)
    try:
        from docx.opc.constants import RELATIONSHIP_TYPE
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        part = paragraph.part
        r_id = part.relate_to(str(url).strip(), RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), r_id)

        new_run = OxmlElement("w:r")
        rPr = OxmlElement("w:rPr")
        if color_hex:
            c = OxmlElement("w:color")
            c.set(qn("w:val"), color_hex)
            rPr.append(c)
        if underline:
            u = OxmlElement("w:u")
            u.set(qn("w:val"), "single")
            rPr.append(u)
        new_run.append(rPr)
        new_run.text = text
        hyperlink.append(new_run)
        paragraph._p.append(hyperlink)
        return hyperlink
    except Exception as exc:
        logger.debug("add_docx_hyperlink_error: %s", exc)
        return paragraph.add_run(text)


def write_exact_product_docx(
    path: str | Path,
    report: ExactProductReport,
    *,
    title: str = "Отчёт о подборе товара и аналогов по ТЗ",
) -> Path:
    """
    Generates official Word DOCX report with positions overview, Form 2 parameters,
    domestic analogs comparison table, and verified datasheets/sources registry.
    """
    target_path = Path(path)
    target_path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    _set_docx_margins(doc)

    # 1. Document Header
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(0)
    h1.paragraph_format.space_after = Pt(2)
    r_brand = h1.add_run("TenderLex")
    r_brand.font.bold = True
    r_brand.font.size = Pt(14)
    r_brand.font.color.rgb = BRAND_EMERALD

    r_pipe = h1.add_run(" | ")
    r_pipe.font.size = Pt(14)
    r_pipe.font.color.rgb = RGBColor(148, 163, 184)

    r_title = h1.add_run(title)
    r_title.font.bold = True
    r_title.font.size = Pt(13)
    r_title.font.color.rgb = TEXT_DARK

    # 2. Metadata card
    meta_table = doc.add_table(rows=2, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False
    _set_table_fixed_width(meta_table, 6.97)

    col_widths_meta = [Inches(1.4), Inches(5.57)]
    for row in meta_table.rows:
        for i, w in enumerate(col_widths_meta):
            row.cells[i].width = w

    c00, c01 = meta_table.rows[0].cells[0], meta_table.rows[0].cells[1]
    p00 = c00.paragraphs[0]
    p00.add_run("Закупка / ТЗ:").font.bold = True
    p01 = c01.paragraphs[0]
    p01.add_run(report.procurement_title or "Спецификация технического задания")

    sources_str = "Открытый интернет (Яндекс Поиск), Реестр Минпромторга (ГИСП), PDF-паспорта заводов"
    if report.web_sources:
        sources_str += f" | Проверено источников: {len(report.web_sources)} ({', '.join(report.web_sources[:3])})"

    c10, c11 = meta_table.rows[1].cells[0], meta_table.rows[1].cells[1]
    p10 = c10.paragraphs[0]
    p10.add_run("Источники:").font.bold = True
    p11 = c11.paragraphs[0]
    p11.add_run(sources_str)

    for row in meta_table.rows:
        for cell in row.cells:
            _set_cell_bg(cell, META_BG)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            for r in p.runs:
                r.font.size = Pt(8.5)
                r.font.color.rgb = RGBColor(51, 65, 85)
    _set_table_full_grid_borders(meta_table, "E2E8F0")

    # 3. Summary
    if report.summary:
        sum_p = doc.add_paragraph()
        sum_p.paragraph_format.space_before = Pt(6)
        sum_p.paragraph_format.space_after = Pt(8)
        sum_p.paragraph_format.left_indent = Inches(0.05)
        r_sum_lbl = sum_p.add_run("Экспертное резюме: ")
        r_sum_lbl.font.bold = True
        r_sum_lbl.font.size = Pt(9.5)
        r_sum_lbl.font.color.rgb = BRAND_EMERALD

        r_sum_txt = sum_p.add_run(report.summary)
        r_sum_txt.font.size = Pt(9.5)
        r_sum_txt.font.color.rgb = TEXT_DARK

    # 4. Section 1: Overview
    sec1_p = doc.add_paragraph()
    sec1_p.paragraph_format.space_before = Pt(4)
    sec1_p.paragraph_format.space_after = Pt(3)
    sec1_run = sec1_p.add_run("1. Сводная ведомость подбора по позициям спецификации")
    sec1_run.font.bold = True
    sec1_run.font.size = Pt(10.5)
    sec1_run.font.color.rgb = BRAND_EMERALD

    sum_table = doc.add_table(rows=1, cols=7)
    sum_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sum_table.autofit = False
    _set_table_fixed_width(sum_table, 6.97)

    sum_headers = [
        ("№", Inches(0.35)),
        ("Позиция из ТЗ", Inches(1.55)),
        ("Выявленный товар / модель", Inches(1.30)),
        ("Завод-изготовитель", Inches(1.30)),
        ("Реестр ГИСП", Inches(0.82)),
        ("Соотв.", Inches(0.45)),
        ("Основной аналог (РФ)", Inches(1.20)),
    ]

    for i, (h_text, w) in enumerate(sum_headers):
        cell = sum_table.rows[0].cells[i]
        cell.width = w
        cell.text = h_text
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        for r in p.runs:
            r.font.bold = True
            r.font.size = Pt(8)
            r.font.color.rgb = RGBColor(255, 255, 255)
        _set_cell_bg(cell, DARK_EMERALD)

    for pos_idx, pos in enumerate(report.positions, start=1):
        row = sum_table.add_row()
        cells = row.cells
        cells[0].text = str(pos.position_no or pos_idx)
        cells[1].text = pos.name_in_tz
        cells[2].text = f"{pos.identified_brand} {pos.identified_model}".strip()
        cells[3].text = pos.manufacturer
        gisp_cell_text = "Не в реестре"
        if pos.gisp_match and pos.gisp_match.registry_number:
            gisp_cell_text = f"№ {pos.gisp_match.registry_number}"
            if pos.gisp_match.conclusion_number:
                gisp_cell_text += f"\n(Закл. № {pos.gisp_match.conclusion_number})"
        cells[4].text = gisp_cell_text
        conf_pct = int(round(pos.confidence * 100))
        if conf_pct >= 85:
            cells[5].text = f"{conf_pct}%"
        elif conf_pct >= 60:
            cells[5].text = f"{conf_pct}%"
        else:
            cells[5].text = f"{conf_pct}%\n(Откл.)"

        main_alt = pos.alternative_brands[0] if pos.alternative_brands else None
        cells[6].text = f"{main_alt.brand} {main_alt.model}" if main_alt else "—"

        fill_color = ZEBRA_MINT if pos_idx % 2 == 1 else "FFFFFF"

        for idx, c in enumerate(cells):
            c.width = sum_headers[idx][1]
            c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            _set_cell_bg(c, fill_color)
            p = c.paragraphs[0]
            p.paragraph_format.line_spacing = 1.05
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            for r in p.runs:
                r.font.size = Pt(8)
                r.font.color.rgb = TEXT_DARK
            if idx in (0, 4, 5):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if idx == 2:
                for r in p.runs:
                    r.font.bold = True
            if idx == 4 and (pos.gisp_match and pos.gisp_match.registry_number):
                for r in p.runs:
                    r.font.bold = True
                    r.font.color.rgb = TEXT_GREEN
            if idx == 5:
                for r in p.runs:
                    r.font.bold = True
                    if conf_pct >= 85:
                        r.font.color.rgb = TEXT_GREEN
                    elif conf_pct >= 60:
                        r.font.color.rgb = RGBColor(180, 83, 9)
                    else:
                        r.font.color.rgb = RGBColor(220, 38, 38)

    _set_table_full_grid_borders(sum_table, BORDER_COLOR)
    _prevent_row_split(sum_table)
    _set_table_header_repeat(sum_table)

    # 5. Section 2: Details for Form 2 & Alternates
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    sec2_p = doc.add_paragraph()
    sec2_p.paragraph_format.space_before = Pt(8)
    sec2_p.paragraph_format.space_after = Pt(4)
    sec2_run = sec2_p.add_run("2. Подробные характеристики для заявки и сравнение аналогов")
    sec2_run.font.bold = True
    sec2_run.font.size = Pt(10.5)
    sec2_run.font.color.rgb = BRAND_EMERALD

    for pos in report.positions:
        pos_banner = doc.add_table(rows=1, cols=1)
        pos_banner.alignment = WD_TABLE_ALIGNMENT.CENTER
        pos_banner.autofit = False
        _set_table_fixed_width(pos_banner, 6.97)
        pos_banner.rows[0].cells[0].width = Inches(6.97)
        b_cell = pos_banner.rows[0].cells[0]

        conf_pct = int(round(pos.confidence * 100))
        banner_bg = BANNER_EMERALD if conf_pct >= 60 else "991B1B"
        _set_cell_bg(b_cell, banner_bg)
        bp = b_cell.paragraphs[0]
        bp.paragraph_format.space_before = Pt(3)
        bp.paragraph_format.space_after = Pt(3)
        bp.paragraph_format.left_indent = Inches(0.08)

        gisp_str = "ГИСП: Не в реестре"
        if pos.gisp_match and pos.gisp_match.registry_number:
            gisp_str = f"ГИСП № {pos.gisp_match.registry_number}"
            if pos.gisp_match.conclusion_number:
                gisp_str += f" (Заключение № {pos.gisp_match.conclusion_number})"
        status_prefix = "" if conf_pct >= 60 else f"ВНИМАНИЕ — ОТКЛОНЕНИЕ ОТ ТЗ ({conf_pct}%): "
        brun = bp.add_run(
            f"ПОЗИЦИЯ №{pos.position_no}: {pos.name_in_tz}\n"
            f"{status_prefix}Товар: {pos.identified_brand} {pos.identified_model} ({pos.manufacturer})   |   {gisp_str}"
        )
        brun.font.bold = True
        brun.font.size = Pt(9)
        brun.font.color.rgb = RGBColor(255, 255, 255)

        if pos.source_url:
            r_src = bp.add_run("   |   Источник: ")
            r_src.font.bold = True
            r_src.font.size = Pt(9)
            r_src.font.color.rgb = RGBColor(255, 255, 255)

            parsed_u = urlparse(pos.source_url)
            host = parsed_u.netloc or pos.source_url
            path_part = parsed_u.path.strip("/").split("/")[0] if parsed_u.path else ""
            anchor_label = f"{host}/{path_part} ↗" if path_part else f"{host} ↗"
            if len(anchor_label) > 35:
                anchor_label = f"{host} ↗"
            _add_docx_hyperlink(bp, pos.source_url, anchor_label, color_hex="99F6E4", underline=True)

        _set_table_full_grid_borders(pos_banner, banner_bg)

        if pos.reasoning:
            rsn_p = doc.add_paragraph()
            rsn_p.paragraph_format.space_before = Pt(3)
            rsn_p.paragraph_format.space_after = Pt(3)
            rsn_p.paragraph_format.left_indent = Inches(0.05)
            lbl_title = "Обоснование соответствия ТЗ: " if conf_pct >= 60 else "Результат сверки с ТЗ (Отклонение): "
            lbl_color = BRAND_EMERALD if conf_pct >= 60 else RGBColor(220, 38, 38)
            r_lbl = rsn_p.add_run(lbl_title)
            r_lbl.font.bold = True
            r_lbl.font.size = Pt(8.5)
            r_lbl.font.color.rgb = lbl_color

            r_val = rsn_p.add_run(pos.reasoning)
            r_val.font.size = Pt(8.5)
            r_val.font.color.rgb = TEXT_MUTED

        # Specs breakdown table
        if pos.specs_breakdown:
            sp_lbl = doc.add_paragraph()
            sp_lbl.paragraph_format.space_before = Pt(2)
            sp_lbl.paragraph_format.space_after = Pt(2)
            sp_lbl_run = sp_lbl.add_run("Показатели для первой части заявки (проверка по первоисточникам):")
            sp_lbl_run.font.bold = True
            sp_lbl_run.font.size = Pt(8.5)
            sp_lbl_run.font.color.rgb = TEXT_DARK

            spec_table = doc.add_table(rows=1, cols=6)
            spec_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            spec_table.autofit = False
            _set_table_fixed_width(spec_table, 6.97)

            spec_headers = [
                ("№", Inches(0.35)),
                ("Наименование параметра", Inches(1.70)),
                ("Требование ТЗ", Inches(1.45)),
                ("Конкретный показатель товара", Inches(1.50)),
                ("Соответствие", Inches(0.77)),
                ("Примечание / Источник", Inches(1.20)),
            ]

            for i, (h_text, w) in enumerate(spec_headers):
                cell = spec_table.rows[0].cells[i]
                cell.width = w
                cell.text = h_text
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                for r in p.runs:
                    r.font.bold = True
                    r.font.size = Pt(8)
                    r.font.color.rgb = RGBColor(255, 255, 255)
                _set_cell_bg(cell, MEDIUM_EMERALD)

            for s_idx, spec in enumerate(pos.specs_breakdown, start=1):
                row = spec_table.add_row()
                cells = row.cells
                cells[0].text = str(s_idx)
                cells[1].text = spec.param_name
                cells[2].text = spec.tz_requirement
                cells[3].text = spec.product_fact
                status_text = (
                    "Подходит"
                    if spec.status == "match"
                    else "Уточнить (по паспорту)"
                    if spec.status == "clarify"
                    else "Отклонение"
                )
                cells[4].text = status_text
                cells[5].text = spec.comment

                fill_color = ZEBRA_MINT if s_idx % 2 == 1 else "FFFFFF"

                for idx, c in enumerate(cells):
                    c.width = spec_headers[idx][1]
                    c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                    _set_cell_bg(c, fill_color)
                    p = c.paragraphs[0]
                    p.paragraph_format.line_spacing = 1.05
                    p.paragraph_format.space_before = Pt(1.5)
                    p.paragraph_format.space_after = Pt(1.5)
                    for r in p.runs:
                        r.font.size = Pt(7.5)
                        r.font.color.rgb = TEXT_DARK
                    if idx in (0, 4):
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    if idx == 3:
                        for r in p.runs:
                            r.font.bold = True
                    if idx == 4:
                        for r in p.runs:
                            r.font.bold = True
                            if spec.status == "match":
                                r.font.color.rgb = TEXT_GREEN
                            elif spec.status == "clarify":
                                r.font.color.rgb = RGBColor(180, 83, 9)
                            else:
                                r.font.color.rgb = RGBColor(185, 28, 28)

            _set_table_full_grid_borders(spec_table, BORDER_COLOR)
            _prevent_row_split(spec_table)
            _set_table_header_repeat(spec_table)

        # Alternates table
        if pos.alternative_brands:
            alt_lbl = doc.add_paragraph()
            alt_lbl.paragraph_format.space_before = Pt(3)
            alt_lbl.paragraph_format.space_after = Pt(2)
            alt_lbl_run = alt_lbl.add_run("Сравнение взаимозаменяемых российских аналогов (по 44/223-ФЗ):")
            alt_lbl_run.font.bold = True
            alt_lbl_run.font.size = Pt(8.5)
            alt_lbl_run.font.color.rgb = TEXT_DARK

            alt_table = doc.add_table(rows=1, cols=6)
            alt_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            alt_table.autofit = False
            _set_table_fixed_width(alt_table, 6.97)

            alt_headers = [
                ("№", Inches(0.35)),
                ("Аналог (Бренд / Модель)", Inches(1.35)),
                ("Завод-изготовитель", Inches(1.25)),
                ("Страна / Реестр", Inches(0.85)),
                ("Совм.", Inches(0.52)),
                ("Сравнение ключевых характеристик и замена", Inches(2.65)),
            ]

            for i, (h_text, w) in enumerate(alt_headers):
                cell = alt_table.rows[0].cells[i]
                cell.width = w
                cell.text = h_text
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                for r in p.runs:
                    r.font.bold = True
                    r.font.size = Pt(8)
                    r.font.color.rgb = RGBColor(255, 255, 255)
                _set_cell_bg(cell, SUBTLE_EMERALD)

            for a_idx, alt in enumerate(pos.alternative_brands, start=1):
                row = alt_table.add_row()
                cells = row.cells
                cells[0].text = str(a_idx)
                cells[1].text = f"{alt.brand} {alt.model}"
                cells[2].text = alt.manufacturer
                cells[3].text = "РФ (Реестр)"
                cells[4].text = f"{int(alt.confidence * 100)}%"
                cells[5].text = alt.notes or "Взаимозаменяемый аналог по ГОСТ/ТУ"

                fill_color = ZEBRA_MINT if a_idx % 2 == 1 else "FFFFFF"

                for idx, c in enumerate(cells):
                    c.width = alt_headers[idx][1]
                    c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                    _set_cell_bg(c, fill_color)
                    p = c.paragraphs[0]
                    p.paragraph_format.line_spacing = 1.05
                    p.paragraph_format.space_before = Pt(1.5)
                    p.paragraph_format.space_after = Pt(1.5)
                    for r in p.runs:
                        r.font.size = Pt(7.5)
                        r.font.color.rgb = TEXT_DARK
                    if idx in (0, 3, 4):
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    if idx == 1:
                        for r in p.runs:
                            r.font.bold = True
                    if idx in (3, 4):
                        for r in p.runs:
                            r.font.bold = True
                            if alt.confidence >= 0.90:
                                r.font.color.rgb = TEXT_GREEN
                            else:
                                r.font.color.rgb = TEXT_DARK

            _set_table_full_grid_borders(alt_table, BORDER_COLOR)
            _prevent_row_split(alt_table)
            _set_table_header_repeat(alt_table)

            # Detailed specs for each analog
            for a_idx, alt in enumerate(pos.alternative_brands, start=1):
                if not alt.specs_breakdown:
                    continue

                alt_form2_p = doc.add_paragraph()
                alt_form2_p.paragraph_format.space_before = Pt(6)
                alt_form2_p.paragraph_format.space_after = Pt(2)
                gisp_alt_tag = (
                    f"ГИСП № {alt.gisp_match.registry_number}"
                    if (alt.gisp_match and alt.gisp_match.registry_number)
                    else "ГИСП: Реестр РФ"
                )
                r_a_hdr = alt_form2_p.add_run(
                    f"Показатели аналога №{a_idx} для первой части заявки — "
                    f"{alt.brand} {alt.model} ({alt.manufacturer})   |   {gisp_alt_tag}   |   Совместимость: {int(alt.confidence * 100)}%:"
                )
                r_a_hdr.font.bold = True
                r_a_hdr.font.size = Pt(8.5)
                r_a_hdr.font.color.rgb = BRAND_EMERALD

                alt_spec_table = doc.add_table(rows=1, cols=6)
                alt_spec_table.alignment = WD_TABLE_ALIGNMENT.CENTER
                alt_spec_table.autofit = False
                _set_table_fixed_width(alt_spec_table, 6.97)

                alt_spec_headers = [
                    ("№", Inches(0.35)),
                    ("Наименование параметра", Inches(1.70)),
                    ("Требование ТЗ", Inches(1.45)),
                    (f"Конкретный показатель ({alt.brand})", Inches(1.50)),
                    ("Соответствие", Inches(0.77)),
                    ("Примечание", Inches(1.20)),
                ]

                for i, (h_text, w) in enumerate(alt_spec_headers):
                    cell = alt_spec_table.rows[0].cells[i]
                    cell.width = w
                    cell.text = h_text
                    p = cell.paragraphs[0]
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    for r in p.runs:
                        r.font.bold = True
                        r.font.size = Pt(8)
                        r.font.color.rgb = RGBColor(255, 255, 255)
                    _set_cell_bg(cell, SUBTLE_EMERALD)

                for s_idx, spec in enumerate(alt.specs_breakdown, start=1):
                    row = alt_spec_table.add_row()
                    cells = row.cells
                    cells[0].text = str(s_idx)
                    cells[1].text = spec.param_name
                    cells[2].text = spec.tz_requirement
                    cells[3].text = spec.product_fact
                    status_text = (
                        "Подходит"
                        if spec.status == "match"
                        else "Уточнить (по паспорту)"
                        if spec.status == "clarify"
                        else "Отклонение"
                    )
                    cells[4].text = status_text
                    cells[5].text = spec.comment

                    fill_color = ZEBRA_MINT if s_idx % 2 == 1 else "FFFFFF"

                    for idx, c in enumerate(cells):
                        c.width = alt_spec_headers[idx][1]
                        c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                        _set_cell_bg(c, fill_color)
                        p = c.paragraphs[0]
                        p.paragraph_format.line_spacing = 1.05
                        p.paragraph_format.space_before = Pt(1.5)
                        p.paragraph_format.space_after = Pt(1.5)
                        for r in p.runs:
                            r.font.size = Pt(7.5)
                            r.font.color.rgb = TEXT_DARK
                        if idx in (0, 4):
                            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        if idx == 3:
                            for r in p.runs:
                                r.font.bold = True
                        if idx == 4:
                            for r in p.runs:
                                r.font.bold = True
                                if spec.status == "match":
                                    r.font.color.rgb = TEXT_GREEN
                                elif spec.status == "clarify":
                                    r.font.color.rgb = RGBColor(180, 83, 9)
                                else:
                                    r.font.color.rgb = RGBColor(185, 28, 28)

                _set_table_full_grid_borders(alt_spec_table, BORDER_COLOR)
                _prevent_row_split(alt_spec_table)
                _set_table_header_repeat(alt_spec_table)

    # 6. Section 3: Verified Datasheets & Web Sources Registry
    if report.verified_documents or report.web_sources:
        doc.add_paragraph().paragraph_format.space_after = Pt(4)
        sec3_p = doc.add_paragraph()
        sec3_p.paragraph_format.space_before = Pt(8)
        sec3_p.paragraph_format.space_after = Pt(4)
        sec3_run = sec3_p.add_run("3. Реестр проверенных открытых веб-источников и паспортов изделий")
        sec3_run.font.bold = True
        sec3_run.font.size = Pt(10.5)
        sec3_run.font.color.rgb = BRAND_EMERALD

        src_table = doc.add_table(rows=1, cols=4)
        src_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        src_table.autofit = False
        _set_table_fixed_width(src_table, 6.97)

        src_headers = [
            ("№", Inches(0.35)),
            ("Тип источника", Inches(1.30)),
            ("Наименование / Документ", Inches(2.32)),
            ("Ссылка / Домен", Inches(3.00)),
        ]

        for i, (h_text, w) in enumerate(src_headers):
            cell = src_table.rows[0].cells[i]
            cell.width = w
            cell.text = h_text
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            for r in p.runs:
                r.font.bold = True
                r.font.size = Pt(8)
                r.font.color.rgb = RGBColor(255, 255, 255)
            _set_cell_bg(cell, DARK_EMERALD)

        row_count = 0
        for doc_item in report.verified_documents[:8]:
            row_count += 1
            row = src_table.add_row()
            cells = row.cells
            cells[0].text = str(row_count)
            cells[1].text = "PDF Паспорт" if doc_item.get("type") == "pdf" else "Сайт завода"
            cells[2].text = str(doc_item.get("title") or "Техническая документация").strip()

            p3 = cells[3].paragraphs[0]
            p3.text = ""
            doc_url = str(doc_item.get("url") or doc_item.get("domain") or "").strip()
            doc_domain = str(doc_item.get("domain") or "").strip()

            if doc_url and doc_url.startswith("http"):
                parsed_u = urlparse(doc_url)
                host = parsed_u.netloc or doc_domain or doc_url
                if doc_item.get("type") == "pdf":
                    anchor_text = f"📄 Скачать PDF-паспорт ({host}) ↗"
                else:
                    anchor_text = f"🌐 {host} ↗"
                _add_docx_hyperlink(p3, doc_url, anchor_text, color_hex="0284C7", underline=True)
            elif doc_domain:
                anchor_text = f"🌐 {doc_domain} ↗"
                full_link = f"https://{doc_domain}" if not doc_domain.startswith("http") else doc_domain
                _add_docx_hyperlink(p3, full_link, anchor_text, color_hex="0284C7", underline=True)
            else:
                p3.add_run("—")

            fill_color = ZEBRA_MINT if row_count % 2 == 1 else "FFFFFF"
            for idx, c in enumerate(cells):
                c.width = src_headers[idx][1]
                c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                _set_cell_bg(c, fill_color)
                p = c.paragraphs[0]
                p.paragraph_format.space_before = Pt(1.5)
                p.paragraph_format.space_after = Pt(1.5)
                for r in p.runs:
                    r.font.size = Pt(7.5)
                    r.font.color.rgb = TEXT_DARK
                if idx in (0, 1):
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        if row_count == 0 and report.web_sources:
            for s_idx, src in enumerate(report.web_sources[:6], start=1):
                row = src_table.add_row()
                cells = row.cells
                cells[0].text = str(s_idx)
                cells[1].text = "Веб-поиск"
                cells[2].text = "Сайт производителя / поставщика"

                p3 = cells[3].paragraphs[0]
                p3.text = ""
                src_url = src if src.startswith("http") else f"https://{src}"
                _add_docx_hyperlink(p3, src_url, f"🌐 {src} ↗", color_hex="0284C7", underline=True)

                fill_color = ZEBRA_MINT if s_idx % 2 == 1 else "FFFFFF"
                for idx, c in enumerate(cells):
                    c.width = src_headers[idx][1]
                    c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                    _set_cell_bg(c, fill_color)
                    p = c.paragraphs[0]
                    p.paragraph_format.space_before = Pt(1.5)
                    p.paragraph_format.space_after = Pt(1.5)
                    for r in p.runs:
                        r.font.size = Pt(7.5)
                        r.font.color.rgb = TEXT_DARK
                    if idx in (0, 1):
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        _set_table_full_grid_borders(src_table, BORDER_COLOR)
        _prevent_row_split(src_table)
        _set_table_header_repeat(src_table)

    # 7. Disclaimer
    doc.add_paragraph().paragraph_format.space_before = Pt(8)
    disc_p = doc.add_paragraph()
    disc_p.paragraph_format.left_indent = Inches(0.05)
    r_disc = disc_p.add_run(f"Примечание: {report.disclaimer}")
    r_disc.font.size = Pt(8)
    r_disc.font.italic = True
    r_disc.font.color.rgb = TEXT_MUTED

    doc.save(str(target_path))
    return target_path
