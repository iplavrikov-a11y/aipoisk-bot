from unittest.mock import AsyncMock, patch
from types import SimpleNamespace

import pytest

from app.exact_product.pipeline import analyze_exact_product
from app.models import SystemSettings


@pytest.mark.asyncio
async def test_selected_rows_keep_identity_and_no_reparse_or_factory_fallback():
    selected = [{"id": "row-2", "position_no": 2, "name": "Насос", "requirements": []},
                {"id": "row-9", "position_no": 9, "name": "Клапан", "requirements": [{"param_name": "Диаметр", "tz_requirement": "25 мм"}]}]
    with patch("app.exact_product.pipeline.detect_exact_products_deep", new_callable=AsyncMock) as deep, patch("app.exact_product.matcher.detect_exact_products_characteristic_first", new_callable=AsyncMock, return_value=None):
        report = await analyze_exact_product(SystemSettings(), "Спецификация оборудования с техническими характеристиками", selected_positions=selected)
    deep.assert_not_called()
    assert [p.position_no for p in report.positions] == [2, 9]
    assert [p.position_id for p in report.positions] == ["row-2", "row-9"]
    assert [p.status for p in report.positions] == ["needs_clarification", "not_found"]
    assert all(not p.identified_model and not p.manufacturer and p.confidence == 0 for p in report.positions)
    assert report.yandex_requests_count == 0
    assert report.search_metrics_available


@pytest.mark.asyncio
async def test_selection_limit_is_explicit():
    with pytest.raises(ValueError, match="100"):
        await analyze_exact_product(SystemSettings(), "Спецификация оборудования с характеристиками", selected_positions=[{"name": str(i)} for i in range(101)])


@pytest.mark.asyncio
async def test_empty_search_never_invents_manufacturer():
    with patch("app.exact_product.pipeline.detect_exact_products_deep", new_callable=AsyncMock, return_value=[]):
        report = await analyze_exact_product(SystemSettings(), "Спецификация оборудования с техническими характеристиками")
    assert all(not p.manufacturer and p.confidence == 0 for p in report.positions)
    assert all(p.status != "found" for p in report.positions)


@pytest.mark.asyncio
async def test_prepare_spreadsheet_keeps_every_row_and_duplicate_names(tmp_path):
    import openpyxl
    from app.exact_product.pipeline import prepare_exact_product_positions
    workbook = openpyxl.Workbook()
    workbook.active.append(["Наименование", "Диаметр", "Кол-во", "Ед.изм."])
    for number in range(12):
        workbook.active.append(["Клапан", f"{number + 1} мм", 2, "шт"])
    path = tmp_path / "spec.xlsx"
    workbook.save(path)
    rows = await prepare_exact_product_positions(SystemSettings(), "", [path])
    assert len(rows) == 12
    assert rows[8]["position_no"] == 9
    assert rows[8]["requirements"] == [{"param_name": "Диаметр", "tz_requirement": "9 мм"}]
    assert rows[8]["source_row"] == 10


@pytest.mark.asyncio
async def test_actual_search_counter_isolated_between_parallel_runs():
    import asyncio
    from app.exact_product.search_metrics import record_search_request
    async def deep(**kwargs):
        count = 2 if kwargs["procurement_title"] == "first" else 3
        for _ in range(count):
            record_search_request()
            await asyncio.sleep(0)
        return []
    with patch("app.exact_product.pipeline.detect_exact_products_deep", side_effect=deep):
        first, second = await asyncio.gather(*[
            analyze_exact_product(SystemSettings(), "Спецификация оборудования с техническими характеристиками", procurement_title=name)
            for name in ["first", "second"]
        ])
    assert [first.yandex_requests_count, second.yandex_requests_count] == [2, 3]


@pytest.mark.asyncio
async def test_word_includes_unfound_original_position_numbers(tmp_path):
    from docx import Document
    from app.exact_product.docx_report import write_exact_product_docx
    with patch("app.exact_product.matcher.detect_exact_products_characteristic_first", new_callable=AsyncMock, return_value=None):
        report = await analyze_exact_product(SystemSettings(), "Спецификация оборудования с техническими характеристиками", selected_positions=[{"id": "9", "position_no": 9, "name": "Клапан", "requirements": []}])
    path = write_exact_product_docx(tmp_path / "selection.docx", report)
    text = "\n".join(p.text for p in Document(path).paragraphs)
    assert "ПОЗИЦИЯ №9: Клапан" in text
    assert "уточнить характеристики" in text
    assert "Завод промышленного оборудования РФ" not in text


@pytest.mark.asyncio
async def test_search_adapter_counts_actual_retry_requests():
    from unittest.mock import MagicMock
    from app.exact_product.search_metrics import SearchMetrics, active_search_metrics
    from app.exact_product.yandex_search import YandexSearchEngine
    metrics = SearchMetrics()
    token = active_search_metrics.set(metrics)
    client = MagicMock()
    limited = MagicMock(status_code=429)
    successful = MagicMock(status_code=200)
    successful.json.return_value = {"rawData": "encoded-xml"}
    client.post = AsyncMock(side_effect=[limited, successful])
    engine = YandexSearchEngine("folder", "key")
    try:
        with patch.object(engine, "_get_client", new_callable=AsyncMock, return_value=client), patch.object(engine, "_parse_v2_xml", return_value=[SimpleNamespace(domain="manufacturer.ru")]), patch("app.exact_product.yandex_search.asyncio.sleep", new_callable=AsyncMock):
            await engine._search_v2("насос", 1, False)
        assert metrics.requests == 2
    finally:
        active_search_metrics.reset(token)


@pytest.mark.asyncio
async def test_name_only_selected_row_is_searched_using_its_own_source():
    selected = [{"id": "row-9", "position_no": 9, "name": "Смазка WD-40 200мл", "unit": "шт", "requirements": []}]
    with patch("app.exact_product.matcher.detect_exact_products_characteristic_first", new_callable=AsyncMock, return_value=None) as matcher, patch("app.exact_product.pipeline.detect_exact_products_deep", new_callable=AsyncMock) as deep:
        report = await analyze_exact_product(SystemSettings(), "Спецификация оборудования с техническими характеристиками; UNRELATED-ROW", selected_positions=selected)
    matcher.assert_awaited_once()
    prepared = matcher.call_args.kwargs["prepared_positions"][0]
    assert prepared["name"] == "Смазка WD-40 200мл"
    assert prepared["requirements"] == [{"param_name": "Наименование из спецификации", "tz_requirement": "Смазка WD-40 200мл"}]
    assert "UNRELATED-ROW" not in matcher.call_args.args[0]
    assert matcher.call_args.kwargs["evidence_data"] == {}
    assert report.positions[0].position_id == "row-9"
    assert report.positions[0].position_no == 9
    assert report.positions[0].status == "needs_clarification"
    deep.assert_not_called()


@pytest.mark.asyncio
async def test_name_only_selected_row_reaches_real_search_without_document_reparse():
    selected = [{"id": "row-9", "position_no": 9, "name": "Смазка WD-40 200мл", "requirements": []}]
    with patch("app.exact_product.matcher.parse_tz_structure", new_callable=AsyncMock) as parse, patch("app.exact_product.matcher.build_characteristic_queries_ai", new_callable=AsyncMock, return_value=["Смазка WD-40 200мл"]) as queries, patch("app.exact_product.matcher._fetch_search_results_for_specs", new_callable=AsyncMock, return_value=[]) as search, patch("app.exact_product.evidence_miner.mine_procurement_evidence_ai", new_callable=AsyncMock) as mine:
        report = await analyze_exact_product(SystemSettings(), "Спецификация оборудования с техническими характеристиками", selected_positions=selected)
    queries.assert_awaited_once()
    assert queries.call_args.args[0] == "Смазка WD-40 200мл"
    assert search.await_count >= 1
    assert "WD-40" in search.call_args_list[0].args[0]
    parse.assert_not_called()
    mine.assert_not_called()
    assert report.positions[0].position_id == "row-9"
    assert report.positions[0].status == "needs_clarification"


@pytest.mark.asyncio
async def test_name_match_alone_never_claims_full_spec_verification():
    candidate = {"name_in_tz": "Смазка WD-40 200мл", "identified_brand": "WD-40", "identified_model": "200мл", "manufacturer": "WD-40 Company", "confidence": .8,
                 "specs_breakdown": [{"param_name": "Наименование из спецификации", "tz_requirement": "Смазка WD-40 200мл", "product_fact": "Смазка WD-40 200мл", "status": "match"}]}
    with patch("app.exact_product.matcher.detect_exact_products_characteristic_first", new_callable=AsyncMock, return_value=([candidate], [])), patch("app.exact_product.validation.validate_exact_products", return_value=[candidate]):
        report = await analyze_exact_product(SystemSettings(), "Спецификация оборудования с техническими характеристиками", selected_positions=[{"id": "r9", "position_no": 9, "name": "Смазка WD-40 200мл", "requirements": []}])
    assert report.positions[0].identified_brand == "WD-40"
    assert report.positions[0].status == "needs_clarification"
    assert "уточнить технические характеристики" in report.positions[0].reasoning


@pytest.mark.asyncio
async def test_requirement_selection_excludes_unrelated_document_rows():
    selected = [{"id": "row-9", "position_no": 9, "name": "Клапан", "quantity": "2", "unit": "шт", "requirements": [{"param_name": "Диаметр", "tz_requirement": "25 мм"}]}]
    with patch("app.exact_product.matcher.detect_exact_products_characteristic_first", new_callable=AsyncMock, return_value=None) as matcher:
        await analyze_exact_product(SystemSettings(), "Спецификация оборудования с техническими характеристиками; UNRELATED-ROW насос 500 кВт", selected_positions=selected)
    source = matcher.call_args.args[0]
    assert "UNRELATED-ROW" not in source
    assert "500" not in source
    assert "25 мм" in source and "Клапан" in source
    assert "Количество: 2" in source
    assert matcher.call_args.kwargs["evidence_data"] == {}
    assert matcher.call_args.kwargs["full_context"] == source


@pytest.mark.asyncio
async def test_mixed_spreadsheet_and_text_preserve_separate_source_inventories(tmp_path):
    import openpyxl
    from app.exact_product.pipeline import prepare_exact_product_positions
    workbook = openpyxl.Workbook()
    workbook.active.append(["Наименование", "Диаметр"])
    workbook.active.append(["Клапан", "25 мм"])
    workbook.active.append(["Клапан", "25 мм"])
    spreadsheet = tmp_path / "spec.xlsx"
    workbook.save(spreadsheet)
    text_file = tmp_path / "other.txt"
    text_file.write_text("Клапан — давление 4 бара", encoding="utf-8")
    with patch("app.exact_product.matcher.parse_tz_structure", new_callable=AsyncMock, return_value={"positions": [{"name": "Клапан", "requirements": [{"param_name": "Давление", "tz_requirement": "4 бара"}]}]}) as parse:
        rows = await prepare_exact_product_positions(SystemSettings(), "Клапан 25 мм\nКлапан 25 мм\nКлапан — давление 4 бара", [spreadsheet, text_file])
    assert len(rows) == 3
    assert [row["source_file"] for row in rows] == ["spec.xlsx", "spec.xlsx", "other.txt"]
    assert rows[0]["source_row"] == 2 and rows[1]["source_row"] == 3
    assert parse.await_count == 1
    assert "25 мм" not in parse.call_args.args[0]
    assert rows[2]["requirements"][0]["tz_requirement"] == "4 бара"
