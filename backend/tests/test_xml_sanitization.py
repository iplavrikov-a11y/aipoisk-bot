from __future__ import annotations

import json
from pathlib import Path

import docx
import openpyxl
import pytest

from app.document_parser import clean_xml_compatible, combined_document_context
from app.report_builder import (
    write_evidence,
    write_procurement_docx,
    write_quote_request_docx,
    write_supplier_xlsx,
)
from app.exact_product.models import (
    ExactProductPosition,
    ExactProductReport,
    GispRegistryMatch,
    SpecParameterMatch,
)
from app.exact_product.docx_report import write_exact_product_docx


def test_clean_xml_compatible_strips_illegal_and_keeps_valid():
    raw = "Valid: Привет мир 123!\t\n\rBad: \x00\x01\x08\x0e\x1f\udc80\ud800\ufffe\uffff"
    cleaned = clean_xml_compatible(raw)
    assert "\x00" not in cleaned
    assert "\x01" not in cleaned
    assert "\x08" not in cleaned
    assert "\x0e" not in cleaned
    assert "\x1f" not in cleaned
    assert "\udc80" not in cleaned
    assert "\ud800" not in cleaned
    assert "\ufffe" not in cleaned
    assert "\uffff" not in cleaned
    assert "Привет мир 123!" in cleaned
    assert "\t" in cleaned
    assert "\n" in cleaned


def test_clean_xml_compatible_normalizes_formfeed_and_vtab():
    raw = "Page 1\x0cPage 2\x0bLine 2"
    cleaned = clean_xml_compatible(raw)
    assert "\x0c" not in cleaned
    assert "\x0b" not in cleaned
    assert "Page 1\nPage 2\nLine 2" in cleaned


def test_write_procurement_docx_with_control_characters(tmp_path: Path):
    out_docx = tmp_path / "procurement_report.docx"
    bad_md = (
        "# Заголовок с нулём\x00 и переводом\x0c\n"
        "Параграф с управляющим кодом \x1f и суррогатом \udc80.\n\n"
        "| № | Наименование | Характеристики | Ед.изм. | Кол-во |\n"
        "|---|---|---|---|---|\n"
        "| 1 | Товар\x00 тест | Описание\x0b с табом | шт | 10 |\n"
    )
    res = write_procurement_docx(
        out_docx,
        bad_md,
        title="Закупка\x00 123",
    )
    assert res.exists()
    assert res.stat().st_size > 0

    # Ensure document can be opened and parsed by docx
    doc = docx.Document(res)
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "\x00" not in full_text
    assert "\x1f" not in full_text
    assert "Заголовок с нулём" in full_text


def test_write_supplier_xlsx_with_control_characters(tmp_path: Path):
    out_xlsx = tmp_path / "suppliers.xlsx"
    rows = [
        {
            "company_name": "ООО 'Рога и копыта'\x00\x08",
            "site": "https://example.com/test\x1f",
            "phone": "+7 999 123-45-67\x0b",
            "email": "info@example.com\x0c",
            "comment": "Поставщик\x00 сертифицирован\udc80",
            "minprom_registry_status": "matched",
            "minprom_registry_match": {"matched": True, "registry_number": "123\x00456"},
        }
    ]
    res = write_supplier_xlsx(
        out_xlsx,
        rows,
        title="Поставщики\x00",
        target=5,
        subject="Поставка труб\x1f",
    )
    assert res.exists()
    assert res.stat().st_size > 0

    # Ensure workbook can be loaded by openpyxl
    wb = openpyxl.load_workbook(res)
    ws = wb.active
    values = [cell.value for row in ws.rows for cell in row if cell.value is not None]
    for val in values:
        s = str(val)
        assert "\x00" not in s
        assert "\x1f" not in s
        assert "\x0b" not in s
        assert "\x0c" not in s


def test_write_exact_product_docx_with_control_characters(tmp_path: Path):
    out_docx = tmp_path / "exact_product.docx"
    report = ExactProductReport(
        procurement_title="Закупка кабелей\x00 и проводов\x1f",
        total_positions=1,
        summary="Резюме\x0c эксперта\udc80",
        disclaimer="Отказ\x0b от ответственности",
        web_sources=["https://source1.ru\x00"],
        positions=[
            ExactProductPosition(
                position_no=1,
                name_in_tz="Кабель ВВГнг\x00 3х2.5",
                identified_brand="Кабэкс\x1f",
                identified_model="ВВГнг-LS\x0b",
                manufacturer="Завод Кабэкс\x0c",
                confidence=0.95,
                source_url="https://kabex.ru/catalog\x00",
                reasoning="Полное\x00 соответствие\udc80 ТЗ",
                specs_breakdown=[
                    SpecParameterMatch(
                        param_name="Сечение\x00",
                        tz_requirement="не менее 2.5\x1f",
                        product_fact="2.5 мм2\x0b",
                        status="match",
                        comment="ГОСТ\x0c соблюдён",
                    )
                ],
                gisp_match=GispRegistryMatch(
                    registry_number="999\x00888",
                    manufacturer="Кабэкс",
                    product="Кабель",
                    conclusion_number="777\x1f",
                    matched=True,
                ),
            )
        ],
    )
    res = write_exact_product_docx(
        out_docx,
        report,
        title="Отчёт подбора\x00 товара",
    )
    assert res.exists()
    assert res.stat().st_size > 0

    doc = docx.Document(res)
    body_text = "\n".join(p.text for p in doc.paragraphs)
    table_text = "\n".join(p.text for t in doc.tables for row in t.rows for cell in row.cells for p in cell.paragraphs)
    full_text = f"{body_text}\n{table_text}"
    assert "\x00" not in full_text
    assert "\x1f" not in full_text
    assert "Кабель ВВГнг" in full_text


def test_write_evidence_with_surrogates(tmp_path: Path):
    out_json = tmp_path / "evidence.json"
    payload = {
        "error": "Ошибка в имени файла: архив_\udc80\udcff.sig",
        "file": "док_\udce0.docx",
    }
    res = write_evidence(out_json, payload)
    assert res.exists()
    content = res.read_text(encoding="utf-8")
    loaded = json.loads(content)
    assert "\udc80" not in loaded["error"]
    assert "\udce0" not in loaded["file"]


def test_combined_document_context_sanitization():
    items = [
        ("файл_\x00bad_\udc80.pdf", "Текст с \x00нулём и \x1fконтролом."),
        ("empty.txt", "   "),
    ]
    ctx = combined_document_context(items)
    assert "\x00" not in ctx
    assert "\x1f" not in ctx
    assert "\udc80" not in ctx
    assert "=== FILE: файл_bad_" in ctx
    assert "Текст с нулём и контролом." in ctx
