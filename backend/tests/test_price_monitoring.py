import pytest
from starlette.testclient import TestClient
from app.main import app
from app.db import SessionLocal
from app.models import Client, Job, SystemSettings, WebUser, BillingTransaction
from app.security import _admin_token
from app.billing import (
    KIND_PRICE_MONITORING,
    charge_price_monitoring_job,
    effective_price_kopeks,
)
from app.price_monitoring.models import (
    BatchMonitoringResult,
    ItemPriceResult,
    SupplierQuote,
)


@pytest.fixture
def db_session():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()


@pytest.fixture
def admin_client():
    token = _admin_token()
    client = TestClient(app, headers={"x-admin-token": token})
    return client


def test_price_monitoring_admin_settings_api(admin_client):
    # Get settings
    resp = admin_client.get("/api/price-monitoring/settings")
    assert resp.status_code == 200
    data = resp.json()
    assert "default_price_rub" in data
    assert "default_price_kopeks" in data

    # Update settings
    resp = admin_client.post("/api/price-monitoring/settings", json={"default_price_rub": 45.0})
    assert resp.status_code == 200
    data = resp.json()
    assert data["default_price_rub"] == 45.0
    assert data["default_price_kopeks"] == 4500

    # Reset back to 50.0
    admin_client.post("/api/price-monitoring/settings", json={"default_price_rub": 50.0})


def test_price_monitoring_client_access_and_override(db_session, admin_client):
    # List clients
    resp = admin_client.get("/api/price-monitoring/clients")
    assert resp.status_code == 200
    clients_list = resp.json().get("clients", [])
    assert isinstance(clients_list, list)

    # Pick or create a test client
    test_c = db_session.query(Client).filter(Client.telegram_id == "pm_test_999").first()
    if not test_c:
        test_c = Client(
            telegram_id="pm_test_999",
            name="PM Test Client",
            money_balance_kopeks=100000,
            allowed_price_monitoring=False,
        )
        db_session.add(test_c)
        db_session.commit()
        db_session.refresh(test_c)

    # Toggle access and set override price
    resp = admin_client.post(
        f"/api/price-monitoring/clients/{test_c.id}",
        json={"allowed": True, "custom_price_rub": 35.0},
    )
    assert resp.status_code == 200
    assert resp.json()["success"] is True

    # Refresh client and check
    db_session.refresh(test_c)
    assert test_c.allowed_price_monitoring is True

    # Verify effective price in billing
    eff_price = effective_price_kopeks(db_session, test_c, KIND_PRICE_MONITORING)
    assert eff_price == 3500

    # Clean up
    db_session.delete(test_c)
    db_session.commit()


def test_charge_price_monitoring_job(db_session):
    # Create test client with 10,000 kopeks balance
    test_c = Client(
        telegram_id="pm_billing_client",
        name="Billing Test Client",
        money_balance_kopeks=10000,
        allowed_price_monitoring=True,
    )
    db_session.add(test_c)
    db_session.commit()
    db_session.refresh(test_c)

    job = Job(
        client_id=test_c.id,
        mode="price_monitoring",
        status="completed",
        title="Test Spec",
    )
    db_session.add(job)
    db_session.commit()
    db_session.refresh(job)

    # Scenario 1: 0 items found -> charged 0
    charged_zero = charge_price_monitoring_job(db_session, job, found_count=0, price_per_item_kopeks=5000)
    assert charged_zero == 0
    db_session.refresh(test_c)
    assert test_c.money_balance_kopeks == 10000  # unchanged!

    # Scenario 2: 2 items found with price 5000 kopeks -> charged 10000
    charged_two = charge_price_monitoring_job(db_session, job, found_count=2, price_per_item_kopeks=5000)
    assert charged_two == 10000
    db_session.refresh(test_c)
    assert test_c.money_balance_kopeks == 0  # 10000 - 10000 = 0

    # Cleanup
    db_session.query(BillingTransaction).filter(BillingTransaction.client_id == test_c.id).delete()
    db_session.delete(job)
    db_session.delete(test_c)
    db_session.commit()


def test_cleaner_unit_and_hints():
    from app.price_monitoring.cleaner import clean_nomenclature_name, extract_package_hints, normalize_unit

    assert normalize_unit("КГ") == "кг"
    assert normalize_unit("М3") == "м3"
    assert normalize_unit("Т") == "т"
    assert normalize_unit("УП") == "уп"

    hints = extract_package_hints("Пропан сжиженный бал.50л ТУ 51-882", "КГ")
    assert hints["gas_type"] == "propane"
    assert hints["cylinder_volume_l"] == 50

    hints_ox = extract_package_hints("Кислород газ.ГОСТ 5583(м3)", "М3")
    assert hints_ox["gas_type"] == "oxygen"

    hints_pack = extract_package_hints("Пружина пластик.19мм 100шт/НЕ ИСП", "УП")
    assert hints_pack["pack_pieces"] == 100

    clean = clean_nomenclature_name("Пружина пластик.19мм 100шт/НЕ ИСП")
    assert "НЕ ИСП" not in clean


def test_unit_reconciliation_math():
    from app.price_monitoring.extractor import reconcile_quote_unit

    # 1. Propane 50L cylinder -> converted to kg (~21 kg gas)
    q_propane = SupplierQuote(
        company_name="Алан Газ",
        price_rub=2400.0,
        raw_price_rub=2400.0,
        site_unit="баллон 50л",
        unit="кг",
        url="https://alangas-msk.ru/catalog/propan/50l/",
    )
    rec_propane = reconcile_quote_unit(q_propane, "Пропан сжиженный бал.50л ТУ 51-882", "КГ")
    assert rec_propane.price_rub == 114.29
    assert "21 кг" in rec_propane.conversion_comment

    # 2. Oxygen 40L cylinder -> converted to m3 (~6.3 m3 gas)
    q_ox = SupplierQuote(
        company_name="Завод Уралтехгаз",
        price_rub=900.0,
        raw_price_rub=900.0,
        site_unit="баллон 40л",
        unit="м3",
        url="https://svar-ka.ru/shop/goods/kislorod",
    )
    rec_ox = reconcile_quote_unit(q_ox, "Кислород газ.ГОСТ 5583(м3)", "М3")
    assert rec_ox.price_rub == 142.86
    assert "6.3 м³" in rec_ox.conversion_comment

    # 3. Steel ton -> converted to kg
    q_steel = SupplierQuote(
        company_name="ПетроСталь",
        price_rub=60000.0,
        raw_price_rub=60000.0,
        site_unit="т",
        unit="кг",
        url="https://petrostal.ru/list",
    )
    rec_steel = reconcile_quote_unit(q_steel, "Лист 10 ВСт3сп ГОСТ 19903", "КГ")
    assert rec_steel.price_rub == 60.0


def test_cross_supplier_validation_and_outliers():
    from app.price_monitoring.pipeline import validate_and_reconcile_item_quotes

    # Given propane item with 1 quote in kg (89.62) and 1 quote in full cylinder (2400.0)
    item = ItemPriceResult(
        row_number=1,
        raw_name="Пропан сжиженный бал.50л ТУ 51-882",
        clean_name="Пропан сжиженный бал.50л ТУ 51-882",
        unit="КГ",
        quotes=[
            SupplierQuote(company_name="Lidergaz", price_rub=89.62, raw_price_rub=89.62, unit="кг"),
            SupplierQuote(company_name="Алан", price_rub=2400.0, raw_price_rub=2400.0, unit="кг"),
        ],
    )
    validate_and_reconcile_item_quotes(item)
    assert item.quotes[1].price_rub == 114.29
    assert item.quotes[1].is_outlier is False

    # Given an irreconcilable extreme outlier (e.g. 50,000 руб for a 100 руб item)
    item_outlier = ItemPriceResult(
        row_number=2,
        raw_name="Ветошь х/б",
        clean_name="Ветошь х/б",
        unit="КГ",
        quotes=[
            SupplierQuote(company_name="Атон", price_rub=61.2, raw_price_rub=61.2, unit="кг"),
            SupplierQuote(company_name="СпецСбыт", price_rub=65.0, raw_price_rub=65.0, unit="кг"),
            SupplierQuote(company_name="ОшибкаЦены", price_rub=5000.0, raw_price_rub=5000.0, unit="кг"),
        ],
    )
    validate_and_reconcile_item_quotes(item_outlier)
    assert any(q.is_outlier for q in item_outlier.quotes)


def test_extract_inn_kpp_contacts():
    from app.price_monitoring.extractor import _extract_inn_kpp_contacts

    html_sample = """
    <footer>
        <p>ООО "Машсервис", ИНН: 7452150212, КПП: 745201001, ОГРН 1187456012345</p>
        <p>Телефон: +7 (351) 220-80-74, 8 (800) 511-80-74</p>
        <p>Email: info@ms-74.ru, г. Челябинск, ул. Инженерная, 41</p>
    </footer>
    """
    inn, kpp, contacts = _extract_inn_kpp_contacts(html_sample)
    assert inn == "7452150212"
    assert kpp == "745201001"
    assert contacts is not None
    assert "info@ms-74.ru" in contacts


def test_excel_report_structure_and_links(tmp_path):
    import openpyxl
    from app.price_monitoring.excel_report import format_short_url, generate_price_monitoring_excel

    assert format_short_url("https://www.foroffice.ru/products/description/1129.html") == "foroffice.ru/.../1129.html"

    items = [
        ItemPriceResult(
            row_number=1,
            raw_name="Пружина пластик.19мм 100шт",
            clean_name="Пружина пластик.19мм 100шт",
            unit="УП",
            quotes=[
                SupplierQuote(
                    company_name="ForOffice",
                    inn="7724521481",
                    kpp="772401001",
                    contacts="+7 (495) 228-20-11, shop@foroffice.ru, Москва",
                    price_rub=723.0,
                    raw_price_rub=723.0,
                    site_unit="уп (100 шт)",
                    unit="УП",
                    url="https://www.foroffice.ru/products/description/1129.html",
                    conversion_comment="Цена за 1 уп (100 шт)",
                ),
            ],
        ),
    ]

    out_file = tmp_path / "test_report.xlsx"
    generate_price_monitoring_excel(items, str(out_file))

    wb = openpyxl.load_workbook(str(out_file))
    assert "Пример отчета" in wb.sheetnames
    assert "Детализированный реестр" in wb.sheetnames

    ws1 = wb["Пример отчета"]
    # Check rows 1-3 in supplier 1 column (C1:D1, C2:D2, C3:D3)
    assert "ForOffice" in str(ws1["C1"].value)
    assert "7724521481" in str(ws1["C2"].value)
    assert "+7 (495) 228-20-11" in str(ws1["C3"].value)

    # Check cell D5 (link cell): contains company, contacts, shortened link and has hyperlink!
    cell_d5 = ws1["D5"]
    assert "foroffice.ru/.../1129.html" in str(cell_d5.value)
    assert "ForOffice" in str(cell_d5.value)
    assert "7724521481" in str(cell_d5.value)
    assert cell_d5.hyperlink.target == "https://www.foroffice.ru/products/description/1129.html"
    assert cell_d5.comment is not None
    assert "ForOffice" in cell_d5.comment.text
    assert "7724521481" in cell_d5.comment.text

    # Check cell C5 (price cell): MUST be numeric, hyperlinked and have comment
    cell_c5 = ws1["C5"]
    assert cell_c5.value == 723.0
    assert cell_c5.hyperlink.target == "https://www.foroffice.ru/products/description/1129.html"
    assert cell_c5.comment is not None

    # Check summary formula columns (M: Мин. цена, N: Средняя цена)
    assert ws1["M1"].value == "Мин. цена, руб."
    assert ws1["N1"].value == "Средняя цена, руб."
    assert "=MIN(" in str(ws1["M5"].value)
    assert "=AVERAGE(" in str(ws1["N5"].value)


@pytest.mark.asyncio
async def test_dadata_quote_enrichment():
    from app.dadata_client import enrich_company_by_inn_or_query

    # Test by known INN
    res_inn = await enrich_company_by_inn_or_query(inn="7728820940")
    if res_inn:
        assert res_inn.get("inn") == "7728820940"
        assert "ДЕАЛМЕД" in res_inn.get("company_name", "").upper()
        assert res_inn.get("kpp") == "772801001"
        assert res_inn.get("status") == "ACTIVE"

    # Test by company query name
    res_q = await enrich_company_by_inn_or_query(query="Автопитер")
    if res_q:
        assert res_q.get("inn") in {"9714085165", "7820046291"}
        assert "АВТОПИТЕР" in res_q.get("company_name", "").upper()


def test_dynamic_pack_weight_from_site_unit():
    from app.price_monitoring.extractor import reconcile_quote_unit

    # Item doesn't have weight in name, but site_unit specifies 5.5 kg pack
    q = SupplierQuote(
        company_name="Сварка",
        price_rub=825.0,
        raw_price_rub=825.0,
        site_unit="пачка 5.5 кг",
        unit="пачка",
        url="https://svar.ru/el",
        is_price_found=True,
    )
    rec = reconcile_quote_unit(q, "Электрод сварочный ОЗС-12", "КГ")
    assert rec.price_rub == 150.0  # 825 / 5.5 = 150.0
    assert "5.5 кг" in rec.conversion_comment


def test_outlier_excluded_from_excel_formulas(tmp_path):
    import openpyxl
    from app.price_monitoring.excel_report import generate_price_monitoring_excel

    item = ItemPriceResult(
        row_number=1,
        raw_name="Подшипник 13630",
        clean_name="Подшипник 13630",
        unit="ШТ",
        quotes=[
            SupplierQuote(company_name="Русподшипник", price_rub=22000.0, raw_price_rub=22000.0, unit="ШТ", is_outlier=False),
            SupplierQuote(company_name="Уралпруток", price_rub=38000.0, raw_price_rub=38000.0, unit="ШТ", is_outlier=False),
            SupplierQuote(company_name="Аскет", price_rub=45000.0, raw_price_rub=45000.0, unit="ШТ", is_outlier=False),
            SupplierQuote(company_name="ЭТМ", price_rub=492.43, raw_price_rub=492.43, unit="ШТ", is_outlier=True),
        ],
    )

    assert item.best_price == 22000.0  # 492.43 excluded from best_price!

    out_file = tmp_path / "test_outliers.xlsx"
    generate_price_monitoring_excel([item], str(out_file))

    wb = openpyxl.load_workbook(str(out_file))
    ws = wb["Пример отчета"]

    # In row 5: Supplier 1 is C, Supplier 2 is E, Supplier 3 is G, Supplier 4 (outlier) is I
    min_formula = ws["M5"].value
    avg_formula = ws["N5"].value

    assert "C5" in min_formula
    assert "E5" in min_formula
    assert "G5" in min_formula
    assert "I5" not in min_formula  # Outlier column I MUST be excluded!
    assert "I5" not in avg_formula


def test_divergent_prices_suppressed_when_no_consensus_cluster(tmp_path):
    import openpyxl
    from app.price_monitoring.pipeline import cluster_and_reconcile_item_quotes
    from app.price_monitoring.excel_report import generate_price_monitoring_excel

    # Two wildly divergent quotes (14.40 vs 1950.00 = 135x difference)
    item = ItemPriceResult(
        row_number=1,
        raw_name="Броня радиальная К4-52",
        clean_name="Броня радиальная К4-52",
        unit="ШТ",
        quotes=[
            SupplierQuote(company_name="Лейрус", price_rub=14.4, raw_price_rub=14.4, unit="ШТ"),
            SupplierQuote(company_name="ЗаводДробмаш", price_rub=1950.0, raw_price_rub=1950.0, unit="ШТ"),
        ],
    )

    cluster_and_reconcile_item_quotes(item, cluster_ratio=2.2)

    # The wild outlier (1950.0) is marked as outlier, while best confirmed base quote (14.4) is kept
    assert any(q.is_outlier for q in item.quotes)
    assert item.has_prices is True

    out_file = tmp_path / "test_divergence.xlsx"
    generate_price_monitoring_excel([item], str(out_file))

    wb = openpyxl.load_workbook(str(out_file))
    ws = wb["Пример отчета"]

    # Price cell displays the best confirmed quote, never "разброс цен"
    assert ws["C5"].value == 14.4
    assert ws["C5"].value != "разброс цен"
    assert "разброс цен" not in str(ws["D5"].value)


def test_circular_reference_prevention_and_row_height_and_no_address(tmp_path):
    import openpyxl
    from app.price_monitoring.excel_report import generate_price_monitoring_excel

    # 5 quotes, but max_suppliers=3: column I (col 9) is Min price, column J (col 10) is Avg price!
    item = ItemPriceResult(
        row_number=1,
        raw_name="Лента диаграммная 0-100 120мм Р1763 с длинным наименованием для проверки переноса строк",
        clean_name="Лента диаграммная 0-100 120мм Р1763",
        unit="М2",
        quotes=[
            SupplierQuote(
                company_name="ООО Бумага",
                inn="7701234567",
                contacts="Тел: +7 (495) 123-45-67, Email: info@bumaga.ru, г. Москва, ул. Ленина, д. 5",
                url="https://bumaga.ru/lenta",
                price_rub=230.0,
                is_price_found=True,
            ),
            SupplierQuote(
                company_name="ООО Фортланд",
                inn="7814109659",
                contacts="+7 495 799-02-12, sales@fortland.ru, г. Санкт-Петербург",
                url="https://fortland.ru/item",
                price_rub=297.0,
                is_price_found=True,
            ),
            SupplierQuote(
                company_name="ООО ОФСИ",
                inn="5047123720",
                contacts="8 (800) 200-00-00, ofsi@mail.ru",
                url="https://ofsi.ru/p19",
                price_rub=310.0,
                is_price_found=True,
            ),
            SupplierQuote(
                company_name="ООО Спецбумага",
                inn="6670000000",
                contacts="8 800 111-22-33",
                url="https://specbumaga.ru/p",
                price_rub=320.0,
                is_price_found=True,
            ),
            SupplierQuote(
                company_name="ООО Полиграф",
                inn="5401000000",
                contacts="poly@nsk.ru",
                url="https://poligraph.ru/p",
                price_rub=330.0,
                is_price_found=True,
            ),
        ],
    )

    out_file = tmp_path / "test_no_circular.xlsx"
    generate_price_monitoring_excel([item], str(out_file), max_suppliers=3)

    wb = openpyxl.load_workbook(str(out_file), data_only=False)
    ws = wb["Пример отчета"]

    # Header check: columns C, E, G are suppliers 1, 2, 3; Column I is Min, Column J is Avg
    assert ws["I1"].value == "Мин. цена, руб."
    assert ws["J1"].value == "Средняя цена, руб."

    # Cell I5 formula: MUST strictly reference C5,E5,G5 and MUST NOT reference I5 or K5!
    min_formula = str(ws["I5"].value)
    avg_formula = str(ws["J5"].value)
    assert min_formula == "=MIN(C5,E5,G5)"
    assert avg_formula == "=AVERAGE(C5,E5,G5)"
    assert "I5" not in min_formula
    assert "K5" not in min_formula

    # Check cell D5 value: company and phone present, email present, NO address!
    d5_val = str(ws["D5"].value)
    assert "Бумага" in d5_val
    assert "7701234567" in d5_val
    assert "+7 (495) 123-45-67" in d5_val
    assert "info@bumaga.ru" in d5_val
    assert "Москва" not in d5_val
    assert "Ленина" not in d5_val

    # Check hyperlink is preserved
    assert ws["D5"].hyperlink is not None
    assert ws["D5"].hyperlink.target == "https://bumaga.ru/lenta"

    # Dynamic row height check: row 5 has multiple lines of text -> height must be >= 60
    assert ws.row_dimensions[5].height >= 60


def test_fallback_queries_short_names():
    from app.price_monitoring.cleaner import build_fallback_broad_queries, build_commercial_search_queries

    q_wheel = build_fallback_broad_queries("Колесо Л-04.524", "шт")
    assert len(q_wheel) >= 2
    assert any("Л-04.524" in q or "л-04.524" in q.lower() for q in q_wheel)

    q_rag = build_fallback_broad_queries("Ветошь х/б", "кг")
    assert len(q_rag) >= 2
    assert any("ветошь" in q.lower() for q in q_rag)

    q_pin = build_fallback_broad_queries("Палец К4-79", "шт")
    assert len(q_pin) >= 2
    assert any("к4-79" in q.lower() for q in q_pin)

    comm_wheel = build_commercial_search_queries("Колесо Л-04.524", "шт")
    assert len(comm_wheel) >= 3


def test_strict_cluster_ratio_and_marketplaces():
    from app.price_monitoring.pipeline import cluster_and_reconcile_item_quotes, KNOWN_MAJOR_RETAILERS
    from app.price_monitoring.extractor import is_valid_domain

    # 1. Verification: cluster_ratio=1.4 rejects 2x outlier
    # In live test on cable: [69.0, 79.0, 83.5, 95.0, 151.34]
    item = ItemPriceResult(
        row_number=1,
        raw_name="Кабель ВВГнг(А)-LS 3х2.5",
        clean_name="Кабель ВВГнг(А)-LS 3х2.5",
        unit="М",
        quotes=[
            SupplierQuote(company_name="Поставщик 1", price_rub=69.0, raw_price_rub=69.0, unit="М"),
            SupplierQuote(company_name="Поставщик 2", price_rub=79.0, raw_price_rub=79.0, unit="М"),
            SupplierQuote(company_name="Поставщик 3", price_rub=83.5, raw_price_rub=83.5, unit="М"),
            SupplierQuote(company_name="Поставщик 4", price_rub=95.0, raw_price_rub=95.0, unit="М"),
            SupplierQuote(company_name="Поставщик 5 (Завышенный)", price_rub=151.34, raw_price_rub=151.34, unit="М"),
        ],
    )

    cluster_and_reconcile_item_quotes(item, cluster_ratio=1.4)

    # 151.34 / 69.0 = 2.19x > 1.4 -> MUST BE MARKED AS OUTLIER
    valid_qs = [q for q in item.quotes if not q.is_outlier]
    outlier_qs = [q for q in item.quotes if q.is_outlier]

    assert len(valid_qs) == 4
    assert len(outlier_qs) == 1
    assert outlier_qs[0].price_rub == 151.34
    assert [q.price_rub for q in valid_qs] == [69.0, 79.0, 83.5, 95.0]

    # 2. Marketplaces validation: ozon, wb, avito are valid domains when filter_marketplaces=False
    assert is_valid_domain("https://www.ozon.ru/product/123", filter_marketplaces=False) is True
    assert is_valid_domain("https://www.wildberries.ru/catalog/123/detail.aspx", filter_marketplaces=False) is True
    assert is_valid_domain("https://www.avito.ru/item/123", filter_marketplaces=False) is True

    # 3. Known major retailers dictionary includes marketplaces with valid INN
    assert "ozon.ru" in KNOWN_MAJOR_RETAILERS
    assert KNOWN_MAJOR_RETAILERS["ozon.ru"]["inn"] == "7704217370"
    assert "wildberries.ru" in KNOWN_MAJOR_RETAILERS
    assert "avito.ru" in KNOWN_MAJOR_RETAILERS


def test_cookie_management_api(admin_client):
    # 1. List cookies catalog
    resp = admin_client.get("/api/cookies")
    assert resp.status_code == 200
    data = resp.json()
    assert "catalog" in data
    assert "total_configured" in data
    assert any(item["domain"] == "ozon.ru" for item in data["catalog"])

    # 2. Set cookie for test domain
    set_resp = admin_client.post("/api/cookies/testdomain.ru", json={"cookie": "test_sid=12345; auth=abc"})
    assert set_resp.status_code == 200
    assert set_resp.json()["success"] is True

    # 3. Verify it appears
    resp2 = admin_client.get("/api/cookies")
    test_item = next(it for it in resp2.json()["catalog"] if it["domain"] == "testdomain.ru")
    assert test_item["has_cookie"] is True
    assert test_item["cookie_length"] > 0

    # 4. Clean up / Delete cookie
    del_resp = admin_client.delete("/api/cookies/testdomain.ru")
    assert del_resp.status_code == 200
    assert del_resp.json()["success"] is True


@pytest.mark.asyncio
async def test_rule_14_ai_and_generic_query_generation(monkeypatch):
    from app.price_monitoring.cleaner import (
        extract_generic_nomenclature_codes,
        build_commercial_search_queries,
        build_fallback_broad_queries,
    )
    from app.price_monitoring.pipeline import generate_monitoring_queries_ai
    from app.models import SystemSettings

    # 1. Generic code extraction without hardcoded word dictionaries (Rule 14)
    codes_sensor = extract_generic_nomenclature_codes("Датчик давления 0065C21J1 GT13E2")
    assert "0065c21j1" in [c.lower() for c in codes_sensor]
    assert "gt13e2" in [c.lower() for c in codes_sensor]

    codes_crusher = extract_generic_nomenclature_codes("Билодержатель К4-15-01")
    assert "к4-15-01" in [c.lower() for c in codes_crusher]

    codes_tape = extract_generic_nomenclature_codes("Лента диаграммная Р-1763")
    assert "р-1763" in [c.lower() for c in codes_tape]

    # 2. Generic fallback queries work for any item
    broad_q = build_fallback_broad_queries("Датчик давления GT13E2", "шт")
    assert len(broad_q) >= 2
    assert any("gt13e2" in q.lower() for q in broad_q)

    # 3. AI query generation (Rule 14) parses LLM response
    settings = SystemSettings(
        custom_ai_providers_json='[{"id":"test","apiKey":"test-key","baseUrl":"https://api.example.com"}]'
    )

    async def fake_call_llm(s, prompt, **kwargs):
        return '["Датчик давления GT13E2 купить цена", "GT13E2 датчик в наличии склад", "Датчик GT13E2 цена руб"]'

    monkeypatch.setattr("app.price_monitoring.pipeline.call_llm", fake_call_llm)

    ai_queries = await generate_monitoring_queries_ai(
        clean_name="Датчик давления GT13E2",
        unit="шт",
        settings=settings,
        wave=1,
    )
    assert len(ai_queries) == 3
    assert "Датчик давления GT13E2 купить цена" in ai_queries


def test_price_monitoring_multiformat_input_extraction():
    from app.price_monitoring.cleaner import parse_items_from_text, extract_items_from_files_or_text
    from app.price_monitoring.pipeline import parse_input_excel
    import tempfile
    from pathlib import Path
    from openpyxl import Workbook

    # 1. Single item text (like Job 903)
    t1 = "Печь для бани Атмосфера М"
    res1 = parse_items_from_text(t1)
    assert res1 == [("Печь для бани Атмосфера М", "шт")]

    # 2. Numbered list with 10+ distinct items and diverse units
    t2 = """1. Кислород газообразный технический, м3
2) Подшипник 13630, шт
3 - Сталь листовая горячекатаная 10мм, т
4. Электрод ОЗС-12 4мм (кг)
5. Пропан сжиженный бал.50л ТУ 51-882, кг
6. Набивка ВАТИ-202 22х22мм, кг
7. Круг шлиф.ПП 80х20х20, шт
8. Ветошь х/б, кг
9. Смазка 1-13 ТУ 38.401-58-142, кг
10. Кабель ВВГнг-LS 3х2.5, м
11. Труба профильная 40х20х2, м
12. Болт М16х60 ГОСТ 7798-70, шт
"""
    res2 = parse_items_from_text(t2)
    assert len(res2) == 12
    assert res2[0] == ("Кислород газообразный технический", "м3")
    assert res2[1] == ("Подшипник 13630", "шт")
    assert res2[2] == ("Сталь листовая горячекатаная 10мм", "т")
    assert res2[3] == ("Электрод ОЗС-12 4мм", "кг")
    assert res2[9] == ("Кабель ВВГнг-LS 3х2.5", "м")

    # 3. Non-zip / plain text bytes passed to parse_input_excel (should NOT raise BadZipFile!)
    text_bytes = b"\xd0\x9f\xd0\xb5\xd1\x87\xd1\x8c \xd0\xb4\xd0\xbb\xd1\x8f \xd0\xb1\xd0\xb0\xd0\xbd\xd0\xb8 \xd0\x90\xd1\x82\xd0\xbc\xd0\xbe\xd1\x81\xd1\x84\xd0\xb5\xd1\x80\xd0\xb0 \xd0\x9c"
    res_bytes = parse_input_excel(text_bytes)
    assert len(res_bytes) == 1
    assert res_bytes[0][0] == "Печь для бани Атмосфера М"

    # 4. Multiple Excel files + text files combination
    with tempfile.TemporaryDirectory() as tmp:
        p1 = Path(tmp) / "file1.xlsx"
        wb1 = Workbook()
        ws1 = wb1.active
        ws1.append(["Номенклатура", "Ед. изм."])
        ws1.append(["Кислород газообразный", "м3"])
        ws1.append(["Подшипник 13630", "шт"])
        wb1.save(p1)

        p2 = Path(tmp) / "file2.xlsx"
        wb2 = Workbook()
        ws2 = wb2.active
        ws2.append(["Наименование", "Ед."])
        ws2.append(["Сталь листовая 10мм", "т"])
        ws2.append(["Подшипник 13630", "шт"]) # duplicate
        wb2.save(p2)

        p3 = Path(tmp) / "tech_spec.txt"
        p3.write_text("Печь для бани Атмосфера М\nКабель ВВГнг-LS 3х2.5, м", encoding="utf-8")

        # Test multiple files extraction
        multi_items = extract_items_from_files_or_text([p1, p2])
        assert len(multi_items) == 3
        assert ("Кислород газообразный", "м3") in multi_items
        assert ("Сталь листовая 10мм", "т") in multi_items
        assert ("Подшипник 13630", "шт") in multi_items

        # Test text file fallback
        text_items = extract_items_from_files_or_text([p3])
        assert len(text_items) == 2
        assert ("Печь для бани Атмосфера М", "шт") in text_items
        assert ("Кабель ВВГнг-LS 3х2.5", "м") in text_items



