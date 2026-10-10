from __future__ import annotations

from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook

from app.report_builder import (
    QUOTE_REQUEST_INTRO,
    write_quote_request_docx,
    write_supplier_xlsx,
)
from app.supplier_search import (
    ProcurementItem,
    ProcurementProfile,
    _accepted_supplier_results,
)


class MultiItemSuppliersTests(unittest.IsolatedAsyncioTestCase):
    def test_accepted_supplier_results_multi_item_round_robin_balance(self) -> None:
        profile = ProcurementProfile(
            summary="Комплексная закупка оборудования",
            items=(
                ProcurementItem(id="item-1", name="Установка плазменной резки"),
                ProcurementItem(id="item-2", name="Винтовой компрессор"),
                ProcurementItem(id="item-3", name="Фильтрационная вытяжка"),
            ),
        )

        reviewed = []
        # Item 1: 15 verified candidates
        for i in range(15):
            reviewed.append({
                "company_name": f"Плазморез Завод {i}",
                "site": f"https://plasma{i}.ru",
                "phone": "+7 999 000 00 01",
                "email": f"sales@plasma{i}.ru",
                "evidence_status": "verified",
                "quality_score": 90 - i,
                "procurement_item_id": "item-1",
                "procurement_item": "Установка плазменной резки",
                "product_fit": "exact",
            })
        # Item 2: 10 verified candidates
        for i in range(10):
            reviewed.append({
                "company_name": f"Компрессор Пром {i}",
                "site": f"https://compressor{i}.ru",
                "phone": "+7 999 000 00 02",
                "email": f"sales@compressor{i}.ru",
                "evidence_status": "verified",
                "quality_score": 88 - i,
                "procurement_item_id": "item-2",
                "procurement_item": "Винтовой компрессор",
                "product_fit": "exact",
            })
        # Item 3: 10 verified candidates
        for i in range(10):
            reviewed.append({
                "company_name": f"Вытяжка Вент {i}",
                "site": f"https://filter{i}.ru",
                "phone": "+7 999 000 00 03",
                "email": f"sales@filter{i}.ru",
                "evidence_status": "verified",
                "quality_score": 87 - i,
                "procurement_item_id": "item-3",
                "procurement_item": "Фильтрационная вытяжка",
                "product_fit": "exact",
            })

        # Target 15 -> fair quota is 5 per item
        accepted = _accepted_supplier_results(reviewed, target=15, profile=profile, limit_to_target=True)
        self.assertEqual(len(accepted), 15)

        counts = {"item-1": 0, "item-2": 0, "item-3": 0}
        for item in accepted:
            p_id = item.get("procurement_item_id")
            if p_id in counts:
                counts[p_id] += 1

        self.assertEqual(counts["item-1"], 5)
        self.assertEqual(counts["item-2"], 5)
        self.assertEqual(counts["item-3"], 5)

    def test_accepted_supplier_results_multi_item_waterfall_overflow(self) -> None:
        profile = ProcurementProfile(
            summary="Комплексная закупка оборудования",
            items=(
                ProcurementItem(id="item-1", name="Установка плазменной резки"),
                ProcurementItem(id="item-2", name="Винтовой компрессор"),
                ProcurementItem(id="item-3", name="Адсорбционный осушитель"),
            ),
        )

        reviewed = []
        # Item 1: 20 verified candidates
        for i in range(20):
            reviewed.append({
                "company_name": f"Плазма {i}",
                "site": f"https://plasma{i}.ru",
                "phone": "+7 999 000 00 01",
                "email": f"sales@plasma{i}.ru",
                "evidence_status": "verified",
                "quality_score": 90 - i,
                "procurement_item_id": "item-1",
                "procurement_item": "Установка плазменной резки",
                "product_fit": "exact",
            })
        # Item 2: only 2 verified candidates
        for i in range(2):
            reviewed.append({
                "company_name": f"Компрессор {i}",
                "site": f"https://comp{i}.ru",
                "phone": "+7 999 000 00 02",
                "email": f"sales@comp{i}.ru",
                "evidence_status": "verified",
                "quality_score": 85 - i,
                "procurement_item_id": "item-2",
                "procurement_item": "Винтовой компрессор",
                "product_fit": "exact",
            })
        # Item 3: only 1 verified candidate
        reviewed.append({
            "company_name": "Осушитель Завод 0",
            "site": "https://dryer0.ru",
            "phone": "+7 999 000 00 03",
            "email": "sales@dryer0.ru",
            "evidence_status": "verified",
            "quality_score": 84,
            "procurement_item_id": "item-3",
            "procurement_item": "Адсорбционный осушитель",
            "product_fit": "exact",
        })

        # Target 15 -> quota per item is 5.
        # Item 2 has only 2, Item 3 has only 1.
        # Waterfall fills the remaining 15 - 5 - 2 - 1 = 7 slots from Item 1.
        # Total Item 1 = 5 + 7 = 12.
        accepted = _accepted_supplier_results(reviewed, target=15, profile=profile, limit_to_target=True)
        self.assertEqual(len(accepted), 15)

        counts = {"item-1": 0, "item-2": 0, "item-3": 0}
        for item in accepted:
            p_id = item.get("procurement_item_id")
            if p_id in counts:
                counts[p_id] += 1

        self.assertEqual(counts["item-2"], 2)
        self.assertEqual(counts["item-3"], 1)
        self.assertEqual(counts["item-1"], 12)

    def test_write_supplier_xlsx_multi_item_creates_tabs_and_grouped_comments(self) -> None:
        rows = [
            {
                "company_name": "ООО СтанкоМаш",
                "site": "https://stankomash.ru",
                "phone": "+7 495 111 22 33",
                "email": "sales@stankomash.ru",
                "product_fit": "exact",
                "product": "Установка плазменной резки с ЧПУ",
                "procurement_item": "Установка плазменной резки",
                "quality_score": 90,
            },
            {
                "company_name": "ООО ПневмоСнаб",
                "site": "https://pnevmosnab.ru",
                "phone": "+7 495 222 33 44",
                "email": "info@pnevmosnab.ru",
                "product_fit": "exact",
                "product": "Винтовой компрессор 10 бар",
                "procurement_item": "Винтовой компрессор",
                "quality_score": 85,
            },
        ]

        with tempfile.TemporaryDirectory() as tmp:
            xlsx_path = Path(tmp) / "suppliers.xlsx"
            write_supplier_xlsx(
                xlsx_path,
                rows,
                title="Комплексная закупка оборудования",
                target=2,
            )

            wb = load_workbook(xlsx_path)
            sheet_names = wb.sheetnames
            # Main sheet plus 2 dedicated item sheets
            self.assertEqual(len(sheet_names), 3)
            self.assertEqual(sheet_names[0], "Сводный реестр")
            self.assertIn("Установка плазменной резки", sheet_names[1])
            self.assertIn("Винтовой компрессор", sheet_names[2])

            # Check comments on main sheet contain item prefix
            ws_main = wb["Сводный реестр"]
            comment_row1 = ws_main["E6"].value
            self.assertIn("Позиция: Установка плазменной резки", comment_row1)

            # Check dedicated item sheet has only its rows
            ws_item1 = wb[sheet_names[1]]
            item1_company = ws_item1["A6"].value
            self.assertEqual(item1_company, "ООО СтанкоМаш")
            wb.close()

    def test_write_supplier_xlsx_single_item_strictly_single_sheet(self) -> None:
        rows = [
            {
                "company_name": "ООО СтанкоМаш",
                "site": "https://stankomash.ru",
                "phone": "+7 495 111 22 33",
                "email": "sales@stankomash.ru",
                "product_fit": "exact",
                "product": "Установка плазменной резки с ЧПУ",
                "procurement_item": "Установка плазменной резки",
                "quality_score": 90,
            },
            {
                "company_name": "ООО ПлазмаПром",
                "site": "https://plasmaprom.ru",
                "phone": "+7 495 555 66 77",
                "email": "zakaz@plasmaprom.ru",
                "product_fit": "exact",
                "product": "Установка плазменной резки",
                "procurement_item": "Установка плазменной резки",
                "quality_score": 88,
            },
        ]

        with tempfile.TemporaryDirectory() as tmp:
            xlsx_path = Path(tmp) / "suppliers.xlsx"
            write_supplier_xlsx(
                xlsx_path,
                rows,
                title="Закупка станка",
                target=2,
            )

            wb = load_workbook(xlsx_path)
            # Strictly 1 sheet for single item
            self.assertEqual(wb.sheetnames, ["Поставщики"])
            wb.close()

    def test_quote_request_intro_includes_partial_supply_disclaimer(self) -> None:
        self.assertIn("попозиционная поставка", QUOTE_REQUEST_INTRO.lower())

        from docx import Document

        markdown = f"""ЗАПРОС КП

{QUOTE_REQUEST_INTRO}

| № | Наименование | Характеристики | Ед.изм. | Кол-во |
|---|---|---|---|---|
| 1 | Установка плазменной резки | ЧПУ | шт | 1 |
| 2 | Винтовой компрессор | 10 бар | шт | 1 |

### Условия поставки
- **Срок поставки:** 30 дней
"""
        with tempfile.TemporaryDirectory() as tmp:
            doc_path = Path(tmp) / "quote.docx"
            write_quote_request_docx(doc_path, markdown, title="Запрос КП")
            doc = Document(doc_path)
            full_text = "\n".join(p.text for p in doc.paragraphs)
            table_text = "\n".join(cell.text for row in doc.tables[0].rows for cell in row.cells)

        self.assertIn("попозиционная поставка", full_text.lower())
        self.assertIn("Установка плазменной резки", table_text)


    async def test_e2e_multi_item_full_pipeline(self) -> None:
        import json
        from types import SimpleNamespace
        import app.supplier_search as supplier_search
        from app.supplier_search import Candidate

        original_call_llm = supplier_search.call_llm
        original_discover = supplier_search.discover_candidates
        original_collect = supplier_search.collect_pages
        original_dns = supplier_search.candidate_domain_resolves_fast

        async def fake_dns(_domain: str) -> bool:
            return True

        async def fake_call_llm(_settings, prompt: str, *args, routing_key: str = "", **kwargs) -> str:
            if routing_key == "supplier_procurement_profile":
                return json.dumps(
                    {
                        "summary": "Поставка станков лазерной резки и компрессорного оборудования",
                        "items": [
                            {"id": "item-1", "name": "Установка лазерной резки", "aliases": ["лазерный станок"]},
                            {"id": "item-2", "name": "Винтовой компрессор", "aliases": ["компрессорная станция"]},
                        ],
                        "excluded_terms": ["аренда", "б/у"],
                    },
                    ensure_ascii=False,
                )
            if routing_key == "supplier_query_generation":
                return json.dumps(
                    {
                        "queries": [
                            "лазерный станок оптом производитель официальный сайт",
                            "винтовой компрессор дилер поставщик контакты",
                        ]
                    },
                    ensure_ascii=False,
                )
            if routing_key == "minprom_registry_requirement":
                return json.dumps({"required": False, "reason": "Нет требований реестра"}, ensure_ascii=False)
            if routing_key == "supplier_candidate_reranker":
                return json.dumps(
                    {
                        "ranked": [
                            {"id": "0", "keep": True, "confidence": 0.95, "procurement_item_id": "item-1"},
                            {"id": "1", "keep": True, "confidence": 0.94, "procurement_item_id": "item-1"},
                            {"id": "2", "keep": True, "confidence": 0.93, "procurement_item_id": "item-2"},
                            {"id": "3", "keep": True, "confidence": 0.92, "procurement_item_id": "item-2"},
                        ]
                    },
                    ensure_ascii=False,
                )
            if routing_key == "supplier_candidate_verifier":
                if "laser1.ru" in prompt:
                    return json.dumps({
                        "action": "accept", "confidence": 0.95, "site_type": "supplier", "product_fit": "exact",
                        "procurement_item_id": "item-1", "procurement_item_name": "Установка лазерной резки",
                        "company_name": "ООО Лазер Пром", "product": "Установка лазерной резки",
                        "email": "sales@laser1.ru", "phone": "+7 495 101 01 01",
                        "evidence_url": "https://laser1.ru/catalog", "comments": "Производитель лазерных комплексов",
                    }, ensure_ascii=False)
                if "laser2.ru" in prompt:
                    return json.dumps({
                        "action": "accept", "confidence": 0.93, "site_type": "supplier", "product_fit": "exact",
                        "procurement_item_id": "item-1", "procurement_item_name": "Установка лазерной резки",
                        "company_name": "ООО СтанкоЛазер", "product": "Установка лазерной резки",
                        "email": "sales@laser2.ru", "phone": "+7 495 102 02 02",
                        "evidence_url": "https://laser2.ru/catalog", "comments": "Дистрибьютор промышленного оборудования",
                    }, ensure_ascii=False)
                if "comp1.ru" in prompt:
                    return json.dumps({
                        "action": "accept", "confidence": 0.92, "site_type": "supplier", "product_fit": "exact",
                        "procurement_item_id": "item-2", "procurement_item_name": "Винтовой компрессор",
                        "company_name": "ООО Компрессор Завод", "product": "Винтовой компрессор",
                        "email": "sales@comp1.ru", "phone": "+7 495 201 01 01",
                        "evidence_url": "https://comp1.ru/catalog", "comments": "Официальный завод компрессоров",
                    }, ensure_ascii=False)
                if "comp2.ru" in prompt:
                    return json.dumps({
                        "action": "accept", "confidence": 0.91, "site_type": "supplier", "product_fit": "exact",
                        "procurement_item_id": "item-2", "procurement_item_name": "Винтовой компрессор",
                        "company_name": "ООО ПневмоСервис", "product": "Винтовой компрессор",
                        "email": "sales@comp2.ru", "phone": "+7 495 202 02 02",
                        "evidence_url": "https://comp2.ru/catalog", "comments": "Складской дистрибьютор",
                    }, ensure_ascii=False)
                return json.dumps({"action": "reject", "reason": "не подходит"}, ensure_ascii=False)
            raise AssertionError(f"unexpected routing key {routing_key}")

        async def fake_discover(_settings, queries: list[str], max_results: int, **kwargs):
            return (
                [
                    Candidate(url="https://laser1.ru/catalog", domain="laser1.ru", title="Лазерные станки", snippet="производитель официальный сайт", source="test", query=queries[0]),
                    Candidate(url="https://laser2.ru/catalog", domain="laser2.ru", title="СтанкоЛазер", snippet="поставщик лазерного оборудования", source="test", query=queries[0]),
                    Candidate(url="https://comp1.ru/catalog", domain="comp1.ru", title="Винтовые компрессоры", snippet="завод компрессоров", source="test", query=queries[1] if len(queries) > 1 else queries[0]),
                    Candidate(url="https://comp2.ru/catalog", domain="comp2.ru", title="ПневмоСервис", snippet="компрессоры в наличии", source="test", query=queries[1] if len(queries) > 1 else queries[0]),
                ],
                {"provider_order": ["test"], "reports": [{"provider": "test", "status": "ok", "returned": 4}]},
            )

        async def fake_collect(url: str) -> list[dict]:
            domain = url.split("/")[2]
            return [{"url": url, "text": f"Компания {domain} поставщик официальный сайт каталог sales@{domain} +7 495 000 00 00"}]

        supplier_search.call_llm = fake_call_llm
        supplier_search.discover_candidates = fake_discover
        supplier_search.collect_pages = fake_collect
        supplier_search.candidate_domain_resolves_fast = fake_dns

        try:
            context = "Спецификация к закупке:\n1. Установка лазерной резки\n2. Винтовой компрессор"
            accepted, evidence = await supplier_search.discover_suppliers(
                SimpleNamespace(has_active_ai_provider=True),
                context,
                target=4,
            )
            # 1. Pipeline quota check: must have accepted all 4 with balanced allocation 2 and 2
            self.assertEqual(len(accepted), 4)
            item1_count = sum(1 for a in accepted if a.get("procurement_item_id") == "item-1")
            item2_count = sum(1 for a in accepted if a.get("procurement_item_id") == "item-2")
            self.assertEqual(item1_count, 2)
            self.assertEqual(item2_count, 2)

            # 2. Report generation check: multi-sheet Excel
            with tempfile.TemporaryDirectory() as tmp:
                xlsx_path = Path(tmp) / "report.xlsx"
                write_supplier_xlsx(xlsx_path, accepted, title="Лазер и компрессор", target=4)
                wb = load_workbook(xlsx_path)
                self.assertEqual(len(wb.sheetnames), 3)
                self.assertEqual(wb.sheetnames[0], "Сводный реестр")
                self.assertIn("Установка лазерной резки", wb.sheetnames[1])
                self.assertIn("Винтовой компрессор", wb.sheetnames[2])

                # Check main sheet comments have item prefix
                ws_main = wb["Сводный реестр"]
                self.assertIn("Позиция: Установка лазерной резки", str(ws_main["E6"].value))
                wb.close()

                # 3. RFQ docx check
                from docx import Document
                docx_path = Path(tmp) / "rfq.docx"
                rfq_md = f"ЗАПРОС КП\n\n{QUOTE_REQUEST_INTRO}\n\n| № | Товар | Кол-во |\n|---|---|---|\n| 1 | Установка лазерной резки | 1 |\n| 2 | Винтовой компрессор | 1 |"
                write_quote_request_docx(docx_path, rfq_md, title="Запрос КП")
                doc = Document(docx_path)
                full_text = "\n".join(p.text for p in doc.paragraphs)
                self.assertIn("попозиционная поставка", full_text.lower())
        finally:
            supplier_search.call_llm = original_call_llm
            supplier_search.discover_candidates = original_discover
            supplier_search.collect_pages = original_collect
            supplier_search.candidate_domain_resolves_fast = original_dns

    def test_accepted_supplier_results_per_item_respects_custom_client_target_without_cap(self) -> None:
        profile = ProcurementProfile(
            summary="Строительные материалы и оборудование",
            items=[
                ProcurementItem(id="item-1", name="Металлические профили"),
                ProcurementItem(id="item-2", name="Пиломатериалы"),
                ProcurementItem(id="item-3", name="Изоляционные материалы"),
                ProcurementItem(id="item-4", name="Сухие смеси"),
            ],
        )
        reviewed = []
        for item_idx in range(1, 5):
            for i in range(75):
                reviewed.append({
                    "site": f"https://supplier-cat{item_idx}-{i}.ru",
                    "company_name": f"Поставщик {item_idx} #{i}",
                    "evidence_status": "verified",
                    "quality_score": 85,
                    "procurement_item_id": f"item-{item_idx}",
                    "procurement_item": profile.items[item_idx - 1].name,
                    "product_fit": "exact",
                })

        accepted = _accepted_supplier_results(
            reviewed,
            target=70,
            profile=profile,
            multi_item_mode="per_item",
            limit_to_target=True,
        )
        self.assertEqual(len(accepted), 280)
        for item_idx in range(1, 5):
            item_count = sum(1 for r in accepted if r.get("procurement_item_id") == f"item-{item_idx}")
            self.assertEqual(item_count, 70)


if __name__ == "__main__":
    unittest.main()
