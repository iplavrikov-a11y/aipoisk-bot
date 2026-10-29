from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

import openpyxl
import pytest

from app.price_monitoring.models import ItemPriceResult
from app.price_monitoring.pipeline import run_price_monitoring_pipeline, process_single_item, _save_checkpoint, _load_checkpoint


@pytest.mark.asyncio
async def test_selected_source_numbers_and_distinct_ids_survive_report(tmp_path):
    rows = [{"id": "a2", "position_no": 2, "name": "Клапан", "unit": "шт", "source_file": "first.xlsx", "source_row": 3, "quantity": "5", "requirements": []},
            {"id": "b9", "position_no": 9, "name": "Клапан", "unit": "шт", "source_file": "second.xlsx", "source_row": 10, "quantity": "2", "requirements": [{"param_name": "Диаметр", "tz_requirement": "25 мм"}]}]
    async def process(item, *args, **kwargs):
        return ItemPriceResult(item[0], item[1], item[1], item[2], status="not_found")
    with patch("app.price_monitoring.pipeline.process_single_item", side_effect=process), patch("app.price_monitoring.pipeline.YandexSearchEngine") as engine:
        engine.return_value.total_requests = 0
        engine.return_value.close = AsyncMock()
        result = await run_price_monitoring_pipeline(b"", SimpleNamespace(yandex_search_folder_id="", yandex_search_api_key=""), str(tmp_path / "report.xlsx"), items=[("Клапан", "шт"), ("Клапан", "шт")], source_positions=rows)
    assert [item.row_number for item in result.items] == [1, 2]
    assert [item.source_position_no for item in result.items] == [2, 9]
    assert [item.position_id for item in result.items] == ["a2", "b9"]
    assert result.items[1].requirements == rows[1]["requirements"]
    workbook = openpyxl.load_workbook(tmp_path / "report.xlsx")
    assert [workbook["Детализированный реестр"].cell(i, 1).value for i in [2, 3]] == [2, 9]
    assert "second.xlsx" in workbook["Пример отчета"].cell(6, 1).comment.text
    assert "25 мм" in workbook["Детализированный реестр"].cell(3, 2).comment.text


@pytest.mark.asyncio
async def test_mismatched_inventory_length_fails_before_search(tmp_path):
    with pytest.raises(ValueError, match="позиц"):
        await run_price_monitoring_pipeline(b"", SimpleNamespace(), str(tmp_path / "report.xlsx"), items=[("Клапан", "шт")], source_positions=[])


def test_checkpoint_keeps_duplicate_original_numbers_and_separate_source_ids(tmp_path):
    path = tmp_path / "checkpoint.json"
    rows = [ItemPriceResult(index, "Клапан", "Клапан", "шт", position_id=f"id-{index}", source_position_no=9, source_file=f"{index}.xlsx", requirements=[{"param_name": "Диаметр", "tz_requirement": f"{index} мм"}], quantity=str(index)) for index in [1, 2]]
    _save_checkpoint(str(path), rows, 2)
    restored, _ = _load_checkpoint(str(path))
    assert [row.row_number for row in restored] == [1, 2]
    assert [row.position_id for row in restored] == ["id-1", "id-2"]
    assert [row.source_position_no for row in restored] == [9, 9]
    assert restored[1].requirements == rows[1].requirements
    assert restored[1].quantity == "2"


@pytest.mark.asyncio
async def test_matcher_receives_only_own_requirements_and_original_name_is_retained():
    import asyncio
    from app.exact_product.yandex_search import YandexSearchResult
    source = {"id": "p9", "position_no": 9, "name": "Клапан", "requirements": [{"param_name": "Диаметр", "tz_requirement": "25 мм"}]}
    engine = SimpleNamespace(search=AsyncMock(return_value=[YandexSearchResult("https://factory.ru/valve", "factory.ru", "Клапан 25 мм 1000 руб")]))
    with patch("app.price_monitoring.pipeline.fetch_page_content_and_meta", new_callable=AsyncMock, return_value=("Клапан диаметр 25 мм. Цена 1000 руб. Производитель промышленной продукции.", None, None, None)), patch("app.price_monitoring.pipeline.extract_supplier_quote", new_callable=AsyncMock, return_value=None) as extract:
        result = await process_single_item((1, "Клапан", "шт"), engine, SimpleNamespace(has_active_ai_provider=False), None, asyncio.Semaphore(1), source_position=source)
    assert extract.await_count >= 1
    assert "Диаметр: 25 мм" in extract.call_args.kwargs["target_item_name"]
    assert result.raw_name == "Клапан"
    assert result.position_id == "p9"


@pytest.mark.asyncio
async def test_checkpoint_resume_does_not_collapse_same_source_number_in_two_files(tmp_path):
    path = tmp_path / "checkpoint.json"
    sources = [{"id": f"p{index}", "position_no": 9, "name": "Клапан", "unit": "шт", "source_file": f"{index}.xlsx", "requirements": []} for index in [1, 2]]
    completed = [ItemPriceResult(index, "Клапан", "Клапан", "шт", status="not_found", position_id=f"p{index}", source_position_no=9, source_file=f"{index}.xlsx") for index in [1, 2]]
    _save_checkpoint(str(path), completed, 2)
    with patch("app.price_monitoring.pipeline.process_single_item", new_callable=AsyncMock) as process, patch("app.price_monitoring.pipeline.YandexSearchEngine") as engine:
        engine.return_value.total_requests = 0
        engine.return_value.close = AsyncMock()
        result = await run_price_monitoring_pipeline(b"", SimpleNamespace(yandex_search_folder_id="", yandex_search_api_key=""), str(tmp_path / "resume.xlsx"), items=[("Клапан", "шт"), ("Клапан", "шт")], source_positions=sources, checkpoint_path=str(path))
    process.assert_not_called()
    assert [row.position_id for row in result.items] == ["p1", "p2"]
    assert [row.source_position_no for row in result.items] == [9, 9]


@pytest.mark.asyncio
async def test_checkpoint_does_not_reuse_outcome_for_different_source_uuid(tmp_path):
    path = tmp_path / "checkpoint.json"
    _save_checkpoint(str(path), [ItemPriceResult(1, "Клапан", "Клапан", "шт", status="not_found", position_id="old-id", source_position_no=9)], 1)
    async def process(item, *args, **kwargs):
        return ItemPriceResult(1, "Клапан", "Клапан", "шт", status="not_found")
    with patch("app.price_monitoring.pipeline.process_single_item", side_effect=process) as process_mock, patch("app.price_monitoring.pipeline.YandexSearchEngine") as engine:
        engine.return_value.total_requests = 0
        engine.return_value.close = AsyncMock()
        result = await run_price_monitoring_pipeline(b"", SimpleNamespace(yandex_search_folder_id="", yandex_search_api_key=""), str(tmp_path / "changed.xlsx"), items=[("Клапан", "шт")], source_positions=[{"id": "new-id", "position_no": 9, "name": "Клапан", "requirements": []}], checkpoint_path=str(path))
    assert process_mock.call_count == 1
    assert result.items[0].position_id == "new-id"
