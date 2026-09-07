from __future__ import annotations

import logging
from pathlib import Path
from typing import TYPE_CHECKING

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

if TYPE_CHECKING:
    from .models import ExactProductReport

logger = logging.getLogger("tenderlex.exact_product.xlsx")


def write_exact_product_xlsx(
    path: str | Path,
    report: ExactProductReport,
    *,
    title: str = "Подбор товара и аналоги",
) -> Path:
    """Генерирует просторный, высококонтрастный и легко читаемый файл Excel во всю ширину экрана (legacy)."""
    target_path = Path(path)
    target_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Подбор товара и аналоги"

    # Палитра
    navy_dark = PatternFill(start_color="064E3B", end_color="064E3B", fill_type="solid")
    navy_pos = PatternFill(start_color="064E3B", end_color="064E3B", fill_type="solid")
    slate_hdr = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
    alt_hdr = PatternFill(start_color="134E4A", end_color="134E4A", fill_type="solid")
    card_bg = PatternFill(start_color="F4FBF7", end_color="F4FBF7", fill_type="solid")
    hint_bg = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")
    zebra_bg = PatternFill(start_color="F4FBF7", end_color="F4FBF7", fill_type="solid")
    white_bg = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    match_bg = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    clarify_bg = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    mismatch_bg = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")

    # Шрифты
    title_font = Font(name="Calibri", size=14, bold=True, color="047857")
    sub_font = Font(name="Calibri", size=11, bold=True, color="064E3B")
    sec_font = Font(name="Calibri", size=12, bold=True, color="047857")
    white_bold_11 = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    white_bold_10 = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    bold_10 = Font(name="Calibri", size=10, bold=True, color="0F172A")
    reg_10 = Font(name="Calibri", size=10, color="1E293B")
    hint_font = Font(name="Calibri", size=10, color="064E3B")
    green_bold = Font(name="Calibri", size=10, bold=True, color="15803D")
    amber_bold = Font(name="Calibri", size=10, bold=True, color="B45309")
    red_bold = Font(name="Calibri", size=10, bold=True, color="B91C1C")

    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    def _style_range(row: int, start_col: int, end_col: int, fill=None, font=None, align=None, border=None):
        if end_col > start_col:
            ws.merge_cells(start_row=row, start_column=start_col, end_row=row, end_column=end_col)
        for c in range(start_col, end_col + 1):
            cell = ws.cell(row=row, column=c)
            if fill:
                cell.fill = fill
            if font:
                cell.font = font
            if align:
                cell.alignment = align
            if border:
                cell.border = border

    def _calc_merged_h(text: str, total_w: int = 169, line_h: int = 19, min_h: int = 26) -> int:
        cleaned = str(text or "").strip()
        if not cleaned:
            return min_h
        chars_per_line = max(20, int(total_w * 0.65))
        lines = 0
        for part in cleaned.split("\n"):
            lines += max(1, (len(part) + chars_per_line - 1) // chars_per_line)
        return max(min_h, lines * line_h + 10)

    def _calc_row_h(cells: list[tuple[str, int]], line_h: int = 16, min_h: int = 26) -> int:
        max_lines = 1
        for val, col_w in cells:
            cleaned = str(val or "").strip()
            if not cleaned:
                continue
            chars_per_line = max(6, int(col_w * 0.85))
            for part in cleaned.split("\n"):
                lines = max(1, (len(part) + chars_per_line - 1) // chars_per_line)
                if lines > max_lines:
                    max_lines = lines
        return max(min_h, max_lines * line_h + 12)

    # 1. Шапка документа
    ws.append(["TenderLex | ПОДБОР ТОВАРА, ХАРАКТЕРИСТИКИ И АНАЛОГИ ПО ТЗ"])
    r1 = ws.max_row
    _style_range(r1, 1, 8, font=title_font, align=Alignment(vertical="center"))
    ws.row_dimensions[r1].height = 26

    ws.append([f"Закупка: {report.procurement_title} | Всего позиций в спецификации: {report.total_positions}"])
    r2 = ws.max_row
    _style_range(r2, 1, 8, font=sub_font, align=Alignment(vertical="center"))
    ws.row_dimensions[r2].height = 22

    # 2. Навигация и Резюме
    hint_msg = "💡 КАК РАБОТАТЬ С ОТЧЁТОМ: Вверху — сводная ведомость по всей закупке. Ниже — построчные характеристики для первой части заявки (Форма 2) и альтернативные заводы РФ по каждой позиции."
    ws.append([hint_msg])
    r3 = ws.max_row
    _style_range(r3, 1, 8, fill=hint_bg, font=hint_font, align=Alignment(vertical="top", wrap_text=True))
    ws.row_dimensions[r3].height = _calc_merged_h(hint_msg, 169, 19, 26)

    if report.summary:
        sum_text = f"Экспертное резюме: {report.summary}"
        ws.append([sum_text])
        r4 = ws.max_row
        _style_range(r4, 1, 8, fill=card_bg, font=reg_10, align=Alignment(vertical="top", wrap_text=True))
        ws.row_dimensions[r4].height = _calc_merged_h(sum_text, 169, 19, 32)

    # Разделитель
    ws.append([None] * 8)
    r_sep1 = ws.max_row
    ws.row_dimensions[r_sep1].height = 14

    # 3. Раздел 1: СВОДНАЯ ВЕДОМОСТЬ
    ws.append(["1. СВОДНАЯ ВЕДОМОСТЬ ПО ВСЕМ ПОЗИЦИЯМ СПЕЦИФИКАЦИИ"])
    r_s1 = ws.max_row
    _style_range(r_s1, 1, 8, font=sec_font, align=Alignment(vertical="center"))
    ws.row_dimensions[r_s1].height = 26

    headers1 = [
        "№",
        "Позиция из ТЗ",
        "Выявленный бренд",
        "Точная модель / артикул",
        "Завод-изготовитель",
        "Реестр ГИСП",
        "Соответствие",
        "Основной аналог (РФ)",
    ]
    ws.append(headers1)
    r_h1 = ws.max_row
    ws.row_dimensions[r_h1].height = 26
    for col_idx in range(1, 9):
        cell = ws.cell(row=r_h1, column=col_idx)
        cell.fill = navy_dark
        cell.font = white_bold_10
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    col_widths_s1 = [6, 30, 24, 24, 24, 16, 15, 28]

    for pos_idx, pos in enumerate(report.positions, start=1):
        gisp_text = (
            f"№ {pos.gisp_match.registry_number}"
            if (pos.gisp_match and pos.gisp_match.registry_number)
            else "Не в реестре"
        )
        main_alt = pos.alternative_brands[0] if pos.alternative_brands else None
        alt_str = f"{main_alt.brand} {main_alt.model} ({main_alt.manufacturer})" if main_alt else "—"

        conf_pct = int(round(pos.confidence * 100))
        conf_label = f"{conf_pct}%" if conf_pct >= 60 else f"{conf_pct}% (Откл.)"

        row_vals = [
            pos.position_no or pos_idx,
            pos.name_in_tz,
            pos.identified_brand,
            pos.identified_model,
            pos.manufacturer,
            gisp_text,
            conf_label,
            alt_str,
        ]
        ws.append(row_vals)
        curr_r = ws.max_row
        fill_row = zebra_bg if pos_idx % 2 == 1 else white_bg

        for col_idx in range(1, 9):
            c = ws.cell(row=curr_r, column=col_idx)
            c.border = thin_border
            c.fill = fill_row
            if col_idx == 1:
                c.font = bold_10
                c.alignment = Alignment(horizontal="center", vertical="top")
            elif col_idx == 6:
                c.font = green_bold if (pos.gisp_match and pos.gisp_match.registry_number) else reg_10
                c.alignment = Alignment(horizontal="center", vertical="top")
            elif col_idx == 7:
                if conf_pct >= 85:
                    c.font = green_bold
                elif conf_pct >= 60:
                    c.font = amber_bold
                else:
                    c.font = red_bold
                c.alignment = Alignment(horizontal="center", vertical="top")
            elif col_idx in (3, 4):
                c.font = bold_10
                c.alignment = Alignment(vertical="top", wrap_text=True)
            else:
                c.font = reg_10
                c.alignment = Alignment(vertical="top", wrap_text=True)

        cells_with_w = list(zip([str(v) for v in row_vals], col_widths_s1))
        ws.row_dimensions[curr_r].height = _calc_row_h(cells_with_w, line_h=16, min_h=26)

    # Разделитель
    ws.append([None] * 8)
    r_sep2 = ws.max_row
    ws.row_dimensions[r_sep2].height = 16

    # 4. Раздел 2: ПОДРОБНЫЕ ХАРАКТЕРИСТИКИ
    ws.append(["2. ПОДРОБНЫЕ ХАРАКТЕРИСТИКИ ДЛЯ ЗАЯВКИ И АНАЛОГИ ПО КАЖДОЙ ПОЗИЦИИ"])
    r_s2 = ws.max_row
    _style_range(r_s2, 1, 8, font=sec_font, align=Alignment(vertical="center"))
    ws.row_dimensions[r_s2].height = 28

    for pos_idx, pos in enumerate(report.positions, start=1):
        ws.append([None] * 8)
        r_pos_sep = ws.max_row
        ws.row_dimensions[r_pos_sep].height = 10

        gisp_text = (
            f"ГИСП № {pos.gisp_match.registry_number}"
            if (pos.gisp_match and pos.gisp_match.registry_number)
            else "ГИСП: Не в реестре"
        )
        conf_pct = int(round(pos.confidence * 100))
        status_prefix = "" if conf_pct >= 60 else f"ВНИМАНИЕ — ОТКЛОНЕНИЕ ОТ ТЗ ({conf_pct}%): "
        banner_text = (
            f"ПОЗИЦИЯ №{pos.position_no}: {pos.name_in_tz}   |   "
            f"{status_prefix}Товар: {pos.identified_brand} {pos.identified_model} ({pos.manufacturer})   |   {gisp_text}"
        )

        banner_fill = navy_pos if conf_pct >= 60 else PatternFill(start_color="991B1B", end_color="991B1B", fill_type="solid")
        ws.append([banner_text])
        r_ban = ws.max_row
        _style_range(r_ban, 1, 8, fill=banner_fill, font=white_bold_11, align=Alignment(vertical="center", wrap_text=True))
        ws.row_dimensions[r_ban].height = _calc_merged_h(banner_text, 169, 20, 28)

        if pos.reasoning:
            reason_text = f"Обоснование соответствия ТЗ: {pos.reasoning}"
            ws.append([reason_text])
            r_rsn = ws.max_row
            _style_range(r_rsn, 1, 8, fill=card_bg, font=reg_10, align=Alignment(vertical="top", wrap_text=True), border=thin_border)
            ws.row_dimensions[r_rsn].height = _calc_merged_h(reason_text, 169, 18, 26)

        # Таблица характеристик (Форма 2)
        if pos.specs_breakdown:
            ws.append(["№", "Наименование параметра", "Требование заказчика (ТЗ)", "Конкретный показатель товара", "", "Соответствие", "Примечание и обоснование показателя", ""])
            r_sp_h = ws.max_row
            ws.row_dimensions[r_sp_h].height = 24
            _style_range(r_sp_h, 1, 1, fill=slate_hdr, font=white_bold_10, align=Alignment(horizontal="center", vertical="center"))
            _style_range(r_sp_h, 2, 2, fill=slate_hdr, font=white_bold_10, align=Alignment(horizontal="center", vertical="center"))
            _style_range(r_sp_h, 3, 3, fill=slate_hdr, font=white_bold_10, align=Alignment(horizontal="center", vertical="center"))
            _style_range(r_sp_h, 4, 5, fill=slate_hdr, font=white_bold_10, align=Alignment(horizontal="center", vertical="center"))
            _style_range(r_sp_h, 6, 6, fill=slate_hdr, font=white_bold_10, align=Alignment(horizontal="center", vertical="center"))
            _style_range(r_sp_h, 7, 8, fill=slate_hdr, font=white_bold_10, align=Alignment(horizontal="center", vertical="center"))

            for s_idx, spec in enumerate(pos.specs_breakdown, start=1):
                status_str = "Подходит" if spec.status == "match" else "Уточнить (по паспорту)" if spec.status == "clarify" else "Отклонение"
                ws.append([
                    s_idx,
                    spec.param_name,
                    spec.tz_requirement,
                    spec.product_fact,
                    "",
                    status_str,
                    spec.comment,
                    "",
                ])
                r_sp = ws.max_row
                fill_r = zebra_bg if s_idx % 2 == 1 else white_bg

                _style_range(r_sp, 1, 1, fill=fill_r, font=bold_10, align=Alignment(horizontal="center", vertical="top"), border=thin_border)
                _style_range(r_sp, 2, 2, fill=fill_r, font=reg_10, align=Alignment(vertical="top", wrap_text=True), border=thin_border)
                _style_range(r_sp, 3, 3, fill=fill_r, font=reg_10, align=Alignment(vertical="top", wrap_text=True), border=thin_border)
                _style_range(r_sp, 4, 5, fill=fill_r, font=bold_10, align=Alignment(vertical="top", wrap_text=True), border=thin_border)

                status_fill = match_bg if spec.status == "match" else clarify_bg if spec.status == "clarify" else mismatch_bg
                status_font = green_bold if spec.status == "match" else amber_bold if spec.status == "clarify" else red_bold
                _style_range(r_sp, 6, 6, fill=status_fill, font=status_font, align=Alignment(horizontal="center", vertical="top"), border=thin_border)

                _style_range(r_sp, 7, 8, fill=fill_r, font=reg_10, align=Alignment(vertical="top", wrap_text=True), border=thin_border)

                sp_cells = [
                    (str(s_idx), 6),
                    (spec.param_name, 30),
                    (spec.tz_requirement, 24),
                    (spec.product_fact, 44),
                    (status_str, 15),
                    (spec.comment, 43),
                ]
                ws.row_dimensions[r_sp].height = _calc_row_h(sp_cells, line_h=16, min_h=24)

        # Таблица аналогов позиции
        if pos.alternative_brands:
            ws.append(["№", "Аналог (Бренд / Модель)", "Завод-изготовитель", "Страна", "Реестр РФ", "Совместимость", "Особенности, отличия и обоснование замены", ""])
            r_alt_h = ws.max_row
            ws.row_dimensions[r_alt_h].height = 24
            _style_range(r_alt_h, 1, 1, fill=alt_hdr, font=white_bold_10, align=Alignment(horizontal="center", vertical="center"))
            _style_range(r_alt_h, 2, 2, fill=alt_hdr, font=white_bold_10, align=Alignment(horizontal="center", vertical="center"))
            _style_range(r_alt_h, 3, 3, fill=alt_hdr, font=white_bold_10, align=Alignment(horizontal="center", vertical="center"))
            _style_range(r_alt_h, 4, 4, fill=alt_hdr, font=white_bold_10, align=Alignment(horizontal="center", vertical="center"))
            _style_range(r_alt_h, 5, 5, fill=alt_hdr, font=white_bold_10, align=Alignment(horizontal="center", vertical="center"))
            _style_range(r_alt_h, 6, 6, fill=alt_hdr, font=white_bold_10, align=Alignment(horizontal="center", vertical="center"))
            _style_range(r_alt_h, 7, 8, fill=alt_hdr, font=white_bold_10, align=Alignment(horizontal="center", vertical="center"))

            for a_idx, alt in enumerate(pos.alternative_brands, start=1):
                notes_text = alt.notes or "Взаимозаменяемый промышленный аналог"
                alt_conf_pct = int(round(alt.confidence * 100))
                alt_conf_label = f"{alt_conf_pct}%" if alt_conf_pct >= 60 else f"{alt_conf_pct}% (Откл.)"
                ws.append([
                    a_idx,
                    f"{alt.brand} {alt.model}",
                    alt.manufacturer,
                    "Россия",
                    "Реестр РФ",
                    alt_conf_label,
                    notes_text,
                    "",
                ])
                r_alt = ws.max_row
                fill_a = zebra_bg if a_idx % 2 == 1 else white_bg
                alt_conf_font = green_bold if alt_conf_pct >= 85 else amber_bold if alt_conf_pct >= 60 else red_bold

                _style_range(r_alt, 1, 1, fill=fill_a, font=bold_10, align=Alignment(horizontal="center", vertical="top"), border=thin_border)
                _style_range(r_alt, 2, 2, fill=fill_a, font=bold_10, align=Alignment(vertical="top", wrap_text=True), border=thin_border)
                _style_range(r_alt, 3, 3, fill=fill_a, font=reg_10, align=Alignment(vertical="top", wrap_text=True), border=thin_border)
                _style_range(r_alt, 4, 4, fill=fill_a, font=reg_10, align=Alignment(horizontal="center", vertical="top"), border=thin_border)
                _style_range(r_alt, 5, 5, fill=fill_a, font=green_bold, align=Alignment(horizontal="center", vertical="top"), border=thin_border)
                _style_range(r_alt, 6, 6, fill=fill_a, font=alt_conf_font, align=Alignment(horizontal="center", vertical="top"), border=thin_border)
                _style_range(r_alt, 7, 8, fill=fill_a, font=reg_10, align=Alignment(vertical="top", wrap_text=True), border=thin_border)

                alt_cells = [
                    (str(a_idx), 6),
                    (f"{alt.brand} {alt.model}", 30),
                    (alt.manufacturer, 24),
                    ("Россия", 20),
                    ("Реестр РФ", 16),
                    (f"{int(alt.confidence * 100)}%", 15),
                    (notes_text, 43),
                ]
                ws.row_dimensions[r_alt].height = _calc_row_h(alt_cells, line_h=16, min_h=24)

            # Детальные показатели для Формы 2 по каждому аналогу в Excel
            for a_idx, alt in enumerate(pos.alternative_brands, start=1):
                if not alt.specs_breakdown:
                    continue
                ws.append([None] * 8)
                r_alt_sp_sep = ws.max_row
                ws.row_dimensions[r_alt_sp_sep].height = 8

                gisp_alt_tag = (
                    f"ГИСП № {alt.gisp_match.registry_number}"
                    if (alt.gisp_match and alt.gisp_match.registry_number)
                    else "ГИСП: Реестр РФ"
                )
                alt_hdr_text = (
                    f"Показатели аналога №{a_idx} для заявки — {alt.brand} {alt.model} ({alt.manufacturer}) | "
                    f"{gisp_alt_tag} | Совместимость: {int(alt.confidence * 100)}%"
                )
                ws.append([alt_hdr_text])
                r_alt_ban = ws.max_row
                _style_range(r_alt_ban, 1, 8, fill=card_bg, font=bold_10, align=Alignment(vertical="center", wrap_text=True), border=thin_border)
                ws.row_dimensions[r_alt_ban].height = _calc_merged_h(alt_hdr_text, 169, 18, 26)

                ws.append(["№", "Наименование параметра", "Требование заказчика (ТЗ)", f"Конкретный показатель ({alt.brand})", "", "Соответствие", "Примечание и обоснование показателя", ""])
                r_asp_h = ws.max_row
                ws.row_dimensions[r_asp_h].height = 24
                _style_range(r_asp_h, 1, 1, fill=alt_hdr, font=white_bold_10, align=Alignment(horizontal="center", vertical="center"))
                _style_range(r_asp_h, 2, 2, fill=alt_hdr, font=white_bold_10, align=Alignment(horizontal="center", vertical="center"))
                _style_range(r_asp_h, 3, 3, fill=alt_hdr, font=white_bold_10, align=Alignment(horizontal="center", vertical="center"))
                _style_range(r_asp_h, 4, 5, fill=alt_hdr, font=white_bold_10, align=Alignment(horizontal="center", vertical="center"))
                _style_range(r_asp_h, 6, 6, fill=alt_hdr, font=white_bold_10, align=Alignment(horizontal="center", vertical="center"))
                _style_range(r_asp_h, 7, 8, fill=alt_hdr, font=white_bold_10, align=Alignment(horizontal="center", vertical="center"))

                for s_idx, spec in enumerate(alt.specs_breakdown, start=1):
                    status_str = "Подходит" if spec.status == "match" else "Уточнить (по паспорту)" if spec.status == "clarify" else "Отклонение"
                    ws.append([
                        s_idx,
                        spec.param_name,
                        spec.tz_requirement,
                        spec.product_fact,
                        "",
                        status_str,
                        spec.comment,
                        "",
                    ])
                    r_asp = ws.max_row
                    fill_r = zebra_bg if s_idx % 2 == 1 else white_bg

                    _style_range(r_asp, 1, 1, fill=fill_r, font=bold_10, align=Alignment(horizontal="center", vertical="top"), border=thin_border)
                    _style_range(r_asp, 2, 2, fill=fill_r, font=reg_10, align=Alignment(vertical="top", wrap_text=True), border=thin_border)
                    _style_range(r_asp, 3, 3, fill=fill_r, font=reg_10, align=Alignment(vertical="top", wrap_text=True), border=thin_border)
                    _style_range(r_asp, 4, 5, fill=fill_r, font=bold_10, align=Alignment(vertical="top", wrap_text=True), border=thin_border)

                    status_fill = match_bg if spec.status == "match" else clarify_bg if spec.status == "clarify" else mismatch_bg
                    status_font = green_bold if spec.status == "match" else amber_bold if spec.status == "clarify" else red_bold
                    _style_range(r_asp, 6, 6, fill=status_fill, font=status_font, align=Alignment(horizontal="center", vertical="top"), border=thin_border)

                    _style_range(r_asp, 7, 8, fill=fill_r, font=reg_10, align=Alignment(vertical="top", wrap_text=True), border=thin_border)

                    sp_cells = [
                        (str(s_idx), 6),
                        (spec.param_name, 30),
                        (spec.tz_requirement, 24),
                        (spec.product_fact, 44),
                        (status_str, 15),
                        (spec.comment, 43),
                    ]
                    ws.row_dimensions[r_asp].height = _calc_row_h(sp_cells, line_h=16, min_h=24)

    # 5. Дисклеймер
    ws.append([None] * 8)
    r_disc_sep = ws.max_row
    ws.row_dimensions[r_disc_sep].height = 10

    disc_text = f"Примечание: {report.disclaimer}"
    ws.append([disc_text])
    r_disc = ws.max_row
    _style_range(r_disc, 1, 8, font=Font(name="Calibri", size=9, italic=True, color="64748B"), align=Alignment(vertical="center", wrap_text=True))
    ws.row_dimensions[r_disc].height = _calc_merged_h(disc_text, 169, 16, 26)

    # 6. Ширина колонок
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 24
    ws.column_dimensions["D"].width = 22
    ws.column_dimensions["E"].width = 22
    ws.column_dimensions["F"].width = 16
    ws.column_dimensions["G"].width = 15
    ws.column_dimensions["H"].width = 28

    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    wb.save(str(target_path))
    return target_path
